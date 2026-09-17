"""Audit the three public spreadsheets linked by the CLARITY data note.

Run with the project's Python environment. This reads local files only and emits
aggregate counts; candidate time pairs are not verified image pairs.
"""

import hashlib
import json
import numbers
import re
import statistics
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
FILES = {
    "MU-Glioma-Post_ClinicalData-July2025.xlsx":
        "https://www.cancerimagingarchive.net/wp-content/uploads/MU-Glioma-Post_ClinicalData-July2025.xlsx",
    "MU-Glioma-Post_Segmentation_Volumes.xlsx":
        "https://www.cancerimagingarchive.net/wp-content/uploads/MU-Glioma-Post_Segmentation_Volumes.xlsx",
    "ISPY2-Imaging-Cohort-1-Clinical-Data.xlsx":
        "https://www.cancerimagingarchive.net/wp-content/uploads/ISPY2-Imaging-Cohort-1-Clinical-Data.xlsx",
}


def audit():
    result = {"retrieved_date": "2026-09-16", "files": {}}
    for name, url in FILES.items():
        path = ROOT / name
        result["files"][name] = {
            "source_url": url,
            "bytes": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        }

    wb = openpyxl.load_workbook(ROOT / next(iter(FILES)), read_only=True, data_only=True)
    rows = list(wb["MU Glioma Post"].values)
    headers = rows[0]
    patients = [dict(zip(headers, row)) for row in rows[1:]
                if row[0] and str(row[0]).startswith("PatientID_")]
    time_columns = [h for h in headers if h and "MRI (Timepoint_" in h]
    times = [[p[h] for h in time_columns if isinstance(p[h], numbers.Real)]
             for p in patients]
    ordered_times = [sorted(set(t)) for t in times]
    gaps = [b - a for t in ordered_times for a, b in zip(t, t[1:]) if b > a]
    result["mu_clinical"] = {
        "patient_rows": len(patients),
        "unique_patients": len({p["Patient_ID"] for p in patients}),
        "numeric_mri_time_entries": sum(map(len, times)),
        "numeric_timepoints_per_patient": dict(sorted(Counter(map(len, times)).items())),
        "patients_with_at_least_two_times": sum(len(t) >= 2 for t in ordered_times),
        "candidate_adjacent_pairs_sorted_by_day": len(gaps),
        "gap_days_min_median_max": [min(gaps), statistics.median(gaps), max(gaps)],
        "patients_with_nonascending_timepoint_columns": sum(t != sorted(t) for t in times),
        "patients_with_duplicate_numeric_days": sum(len(t) != len(set(t)) for t in times),
        "death_flag_counts": dict(Counter(p["Overall Survival (Death)"] for p in patients)),
        "numeric_diagnosis_to_death_durations": sum(isinstance(
            p["Number of days from Diagnosis to death (Days)"], numbers.Real) for p in patients),
        "columns_matching_survival_censoring": [h for h in headers if h and any(
            token in h.lower() for token in ("death", "survival", "contact", "follow", "censor"))],
        "caveat": "Counts use numeric spreadsheet cells only; no MRI archive was downloaded or matched.",
    }
    wb.close()

    result["mu_volume_sheets"] = {}
    wb = openpyxl.load_workbook(ROOT / "MU-Glioma-Post_Segmentation_Volumes.xlsx",
                               read_only=True, data_only=True)
    for sheet in wb:
        rows = list(sheet.values)
        ids = [str(row[0]).strip() for row in rows[1:] if row[0]]
        matches = [re.search(r"PatientID_0*(\d+)", pid) for pid in ids]
        result["mu_volume_sheets"][sheet.title] = {
            "nonempty_rows": len(ids),
            "unique_raw_ids": len(set(ids)),
            "unique_normalized_patient_ids": len({m.group(1) for m in matches if m}),
            "rows_without_timepoint_suffix": sum("Post-treatment_" not in pid for pid in ids),
            "columns": [h for h in rows[0] if h],
        }
    wb.close()

    wb = openpyxl.load_workbook(ROOT / "ISPY2-Imaging-Cohort-1-Clinical-Data.xlsx",
                               read_only=True, data_only=True)
    rows = list(wb["ISPY2_n985_TCIA_clinical"].values)
    ids = [row[0] for row in rows[1:] if row[0] is not None]
    result["ispy2_clinical"] = {"rows": len(ids), "unique_patients": len(set(ids)),
                                "columns": list(rows[0])}
    wb.close()
    return result


if __name__ == "__main__":
    output = ROOT / "metadata_audit.json"
    output.write_text(json.dumps(audit(), ensure_ascii=False, indent=2) + "\n")
    print(output)
