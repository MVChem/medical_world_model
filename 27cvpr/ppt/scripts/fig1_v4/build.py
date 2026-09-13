"""Recreate fig1_v4 with editable PowerPoint shapes and tight vector exports."""
from io import BytesIO
from pathlib import Path
import subprocess
import tempfile

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "ppt/fig1_v4.pptx"
SOURCE = ROOT / "source/fig1_v4.png"
W, H, Y0, PPI = 1731, 640, 140, 120
INK, WHITE = "171717", "FFFFFF"
BLUE, PINK, GREEN, BEIGE = "9BCAE9", "E7BCE7", "C7E6A9", "EBD9C6"
PLUM, OLIVE, BROWN = "68256C", "557C2F", "8C6639"
CREAM, PALE_BLUE, SLOT, GRAY = "FEFAD1", "D0E7FA", "E89EBA", "BEBEBE"
FONT = "Comic Sans MS"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W / PPI), Inches(H / PPI)
prs.core_properties.title = "MedWorld-JEPA — fig1_v4"
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)


def u(px):
    return Inches(px / PPI)


def style(shape, fill, stroke=INK, width=1.0, name=""):
    shape.name = name or shape.name
    shape.shadow.inherit = False
    effect = shape._element.find("./" + qn("p:style") + "/" + qn("a:effectRef"))
    if effect is not None:
        effect.set("idx", "0")
    if hasattr(shape, "fill"):
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shape.fill.background()
    if stroke:
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    return shape


def box(x, y, w, h, fill=WHITE, stroke=INK, radius=0.08, name="", width=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, u(x), u(y - Y0), u(w), u(h))
    if radius:
        s.adjustments[0] = radius
    return style(s, fill, stroke, width, name)


def text(x, y, w, h, value, size=20, bold=False, font=FONT,
         color=INK, align=PP_ALIGN.CENTER, italic=False, rotation=0):
    s = slide.shapes.add_textbox(u(x), u(y - Y0), u(w), u(h))
    s.name = "Label: " + value.replace("\n", " / ")
    s.rotation = rotation
    tf = s.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, label in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.04
        r = p.add_run()
        r.text = label
        r.font.name = font
        r.font.size = Pt(size * 72 / PPI)
        r.font.bold, r.font.italic = bold, italic
        r.font.color.rgb = RGBColor.from_string(color)
    return s


def line(points, arrow=False, color=INK, width=1.5, name="Connection"):
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      u(x1), u(y1 - Y0), u(x2), u(y2 - Y0))
        style(s, None, color, width, name)
        if arrow and i == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "sm", "len": "sm"}.items():
                end.set(key, value)
            s._element.spPr.get_or_add_ln().append(end)


def polygon(points, fill, stroke=INK, name="Freeform"):
    points = [(x, y - Y0) for x, y in points]
    builder = slide.shapes.build_freeform(*points[0], scale=u(1))
    builder.add_line_segments(points[1:], close=True)
    return style(builder.convert_to_shape(), fill, stroke, 1.0, name)


def trapezoid(x, y, w, h, fill, name, taper=0.24):
    # A native trapezoid rotated so the output edge is the short edge.
    s = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID,
                              u(x + (w - h) / 2), u(y - Y0 + (h - w) / 2),
                              u(h), u(w))
    s.rotation = 90
    s.adjustments[0] = taper
    return style(s, fill, INK, 1.1, name)


def port(x, y, label_x=None, label_y=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y - Y0), u(43), u(43))
    style(s, BLUE, INK, 1.1, "Shared image features D_t")
    text(x, y + 1, 43, 39, "Dₜ", 24, True, italic=True)
    if label_x is not None:
        text(label_x, label_y, 184, 30, "Shared image features", 17, True,
             align=PP_ALIGN.LEFT)


def token(x, y, color, w=27, h=28, name="Token"):
    return box(x, y, w, h, color, "787878", 0.06, name, 0.6)


def flame(x, y):
    points = [(11, 0), (15, 8), (13, 15), (19, 10), (22, 19),
              (21, 26), (16, 31), (8, 32), (2, 28), (0, 22),
              (4, 12), (5, 20), (10, 12)]
    polygon([(x + a, y + b) for a, b in points], "FF812C", None, "Trainable flame")
    polygon([(x + 12, y + 13), (x + 17, y + 23), (x + 15, y + 29),
             (x + 8, y + 29), (x + 6, y + 24)], "FFDD48", None, "Flame center")


def upright_block(x, y, w, h, label):
    box(x, y, w, h, CREAM, INK, 0.10, label)
    text(x + (w - h) / 2, y + (h - w) / 2, h, w,
         label, 18.5, True, font="Arial", rotation=270)


def thumbnail(bounds, x, y, w, h, name):
    data = BytesIO()
    with Image.open(SOURCE) as im:
        im.crop(bounds).save(data, format="PNG")
    data.seek(0)
    s = slide.shapes.add_picture(data, u(x), u(y - Y0), u(w), u(h))
    s.name = name


# Module backgrounds.
box(12, 450, 272, 236, "FFF6E8", INK, 0.07, "Report / history / timing card", 0.9)
box(577, 254, 397, 330, PINK, PLUM, 0.07, "VLM — no transformer blocks", 1.3)
box(595, 363, 360, 116, "FFF9F0", None, 0.12, "Integrated VLM input strip")
box(1148, 159, 257, 206, GREEN, OLIVE, 0.10, "Task decoder", 1.0)
box(1148, 451, 257, 248, BEIGE, BROWN, 0.09, "World model", 1.0)
box(1516, 159, 194, 243, "EDF6FB", "B7BEC2", 0.06, "Task outputs", 0.6)

# Short local connections, including one shared patient-state branch.
line([(251, 317), (285, 317)], True, name="Studies to V-JEPA")
line([(388, 324), (446, 324)], True, name="Encoder to resampler")
line([(416, 324), (416, 346)], name="Encoder feature port")
line([(561, 325), (568, 325), (568, 394), (601, 394)], True,
     name="Visual input to integrated VLM strip")
line([(284, 560), (347, 560)], True, name="Report to tokenizer")
line([(500, 560), (558, 560), (558, 446), (601, 446)], True,
     name="Text input to integrated VLM strip")
line([(974, 420), (994, 420)], True, name="VLM to patient state")
line([(1034, 420), (1050, 420)], name="Shared patient state")
line([(1050, 420), (1050, 206), (1148, 206)], True, name="Patient state to task decoder")
line([(1050, 420), (1050, 480), (1148, 480)], True, name="Patient state to world model")
line([(1115, 252), (1148, 252)], True, name="Task query to decoder")
line([(1218, 371), (1218, 342)], True, name="Shared features to task decoder")
line([(1405, 279), (1423, 279)], True, name="Decoder to task head")
line([(1491, 279), (1516, 279)], True, name="Task head to outputs")
line([(1100, 533), (1148, 533)], True, name="Forecast query to world model")
line([(1100, 613), (1148, 613)], True, name="Horizon to world model")
line([(1207, 709), (1207, 670)], True, name="Shared features to world model")
line([(1405, 603), (1428, 603)], True, name="World model to predicted tokens")
line([(1467, 602), (1497, 602)], True, color="666666", width=1.1,
     name="Predicted tokens to prediction loss")
line([(1642, 602), (1610, 602)], True, color="666666", width=1.1,
     name="Future targets to prediction loss")
line([(1553, 657), (1553, 637)], True, width=1.0, name="Future loss note")

# Current evidence and native document icon.
text(18, 172, 248, 62, "Current / prior studies\n(at or before t)", 23, True)
thumbnail((23, 252, 131, 383), 22, 251, 109, 132, "Current chest radiograph")
thumbnail((144, 252, 250, 383), 142, 251, 109, 132, "Prior chest CT")
trapezoid(285, 234, 103, 169, "BFDFF9", "V-JEPA encoder trapezoid", 0.46)
text(294, 287, 88, 65, "V-JEPA\nencoder", 22)
port(395, 341)
text(382, 384, 140, 49, "Shared\nimage features", 18, True)
box(446, 280, 115, 89, CREAM, INK, 0.12, "Resampler + projector")
text(450, 286, 107, 77, "Resampler\n+ projector", 21, True)
text(21, 459, 253, 82, "Report / history /\ntiming / clinical context\n(at or before t)", 21, True)
polygon([(39, 571), (77, 571), (91, 585), (91, 644), (39, 644)], WHITE,
        name="Native report document")
line([(77, 571), (77, 585), (91, 585)], width=0.8, name="Document fold")
for yy in range(589, 635, 6):
    line([(46, yy), (83 if yy < 631 else 72, yy)], width=0.9, name="Document text")
text(114, 559, 153, 104, "Findings: …\nHistory: …\nTime: …\nIndication: …", 20,
     font="Arial", align=PP_ALIGN.LEFT)
box(347, 515, 153, 91, CREAM, INK, 0.12, "Tokenizer + embedding")
text(353, 526, 141, 65, "Tokenizer\n+ embedding", 22, True)

# VLM: one input strip, no decorative transformer blocks.
text(671, 275, 116, 52, "VLM", 34, True)
text(785, 282, 113, 40, "(LoRA)", 25, True)
for i in range(3):
    token(611 + i * 35, 382, BLUE, name=f"Visual input token {i + 1}")
    token(718 + i * 32, 382, GRAY, name=f"Text input token {i + 1}")
text(646, 380, 27, 28, "…", 22, True)
text(750, 380, 27, 28, "…", 22, True)
for i in range(8):
    token(823 + i * 16, 382, SLOT, 13, 28, f"Learned input slot {i + 1}")
text(608, 417, 100, 51, "Visual\ntokens", 19)
text(708, 417, 104, 51, "Text\ntokens", 19)
text(814, 417, 140, 51, "8 learned\nslots", 19)
text(976, 221, 75, 51, "Patient\nstate", 19)
text(989, 273, 50, 26, "Sₜ", 25, True, italic=True)
box(994, 300, 40, 278, "FFF9FC", "626262", 0.18, "Eight patient-state tokens", 0.8)
for i in range(8):
    token(1001, 307 + i * 34, SLOT, 26, 26, f"Patient-state token {i + 1}")

# Task branch and the requested simple native trapezoid head.
token(1086, 237, BLUE, 29, 29, "Task query token")
text(1062, 275, 76, 51, "Task\nquery", 20)
text(1168, 177, 195, 39, "Task decoder", 26, True)
flame(1365, 177)
upright_block(1179, 223, 70, 118, "Transformer\nblock")
upright_block(1302, 223, 71, 118, "Transformer\nblock")
text(1251, 252, 48, 43, "…", 30, True)
port(1197, 372, 1250, 378)
trapezoid(1423, 186, 68, 184, PALE_BLUE, "Task head trapezoid")
text(1425, 246, 64, 65, "Task\nhead", 22, True)
text(1524, 163, 179, 29, "Clinical prediction", 20, True)
line([(1530, 201), (1530, 247), (1584, 247)], width=1.0, name="Clinical score axes")
for x, yy, hh in [(1539, 227, 16), (1553, 217, 26), (1569, 211, 32)]:
    box(x, yy, 10, hh, BLUE, None, 0, "Clinical score bar")
text(1596, 199, 108, 50, "Risk: 0.72\n(positive)", 20, font="Arial", align=PP_ALIGN.LEFT)
text(1525, 265, 162, 29, "Segmentation", 20, True, align=PP_ALIGN.LEFT)
thumbnail((1528, 300, 1623, 393), 1528, 300, 95, 93, "Segmentation thumbnail")
text(1630, 319, 73, 55, "Task\nloss", 20, True)

# World model and its local future supervision.
text(1168, 463, 195, 39, "World model", 26, True)
text(1178, 504, 193, 28, "Latent predictor", 22, True)
flame(1365, 466)
upright_block(1175, 543, 71, 127, "Predictor\nblock")
upright_block(1302, 543, 71, 127, "Predictor\nblock")
text(1251, 576, 48, 42, "…", 30, True)
token(1071, 518, BLUE, 29, 29, "Forecast query token")
text(1037, 550, 96, 49, "Forecast\nquery", 19)
token(1071, 599, GRAY, 29, 28, "Horizon token")
text(1037, 629, 97, 47, "Horizon\nh", 19)
port(1186, 710, 1238, 716)
text(1403, 469, 129, 53, "Predicted\nfuture tokens", 20)
text(1590, 468, 137, 54, "Future\ntarget tokens", 20, True)
for xx, prefix in [(1428, "Predicted future"), (1642, "Future target")]:
    box(xx, 529, 39, 155, "FAFAFA", "5D5D5D", 0.16, f"{prefix} token column", 0.9)
    for i, yy in enumerate([536, 574, 649]):
        token(xx + 6, yy, GRAY, 27, 27, f"{prefix} token {i + 1}")
    text(xx + 2, 610, 35, 33, "⋮", 27, font="Arial")
box(1497, 554, 113, 83, "FAD9CD", INK, 0.13, "Prediction loss")
text(1501, 563, 105, 63, "Prediction\nloss", 20)
text(1476, 659, 157, 25, "Future loss", 20, True)
text(1480, 685, 140, 48, "shapes patient\nstate", 17)
text(1610, 695, 119, 44, "Real follow-up;\nEMA, stop-grad", 15)

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
with tempfile.TemporaryDirectory(prefix="fig1_v4_lo_") as profile:
    subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}",
                    "--headless", "--convert-to", "pdf:impress_pdf_Export",
                    "--outdir", str(OUT.parent), str(OUT)], check=True, capture_output=True)
with fitz.open(OUT.with_suffix(".pdf")) as pdf:
    page = pdf[0]
    # Outlined SVG text preserves the reference font on other machines.
    OUT.with_suffix(".svg").write_text(page.get_svg_image(text_as_path=True))
    page.get_pixmap(matrix=fitz.Matrix(3, 3)).save(OUT.with_suffix(".png"))
print(OUT)
