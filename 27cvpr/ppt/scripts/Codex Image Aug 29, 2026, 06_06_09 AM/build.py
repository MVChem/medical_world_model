from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
NAME = "Codex Image Aug 29, 2026, 06_06_09 AM"
SOURCE = ROOT / "source" / f"{NAME}.png"
OUTPUT = ROOT / "ppt" / f"{NAME}.pptx"

W, H, PX_PER_INCH = 1869, 842, 100
BG, WHITE, INK = "FBF9F7", "FEFDFC", "171419"
GRAY, LIGHT = "626466", "A6A6A4"
BLUE, BLUE_MID, BLUE_PALE = "315D84", "8EA5B9", "DCE5EC"
PURPLE, PURPLE_MID, PURPLE_PALE = "60416F", "9A87A4", "E6DFE8"
ORANGE, ORANGE_MID, ORANGE_PALE = "B55B3E", "D89B86", "EED4CA"
GREEN, GREEN_MID, GREEN_PALE = "426F50", "91AD97", "E4ECE5"
FONT, MATH = "Arial", "Times New Roman"


def u(px):
    return Inches(px / PX_PER_INCH)


def rgb(value):
    return RGBColor.from_string(value)


def clean(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def dash(shape, value="dash"):
    ln = shape._element.spPr.get_or_add_ln()
    for child in list(ln):
        if child.tag.endswith("prstDash"):
            ln.remove(child)
    node = OxmlElement("a:prstDash")
    node.set("val", value)
    ln.append(node)


def arrow(shape):
    ln = shape._element.spPr.get_or_add_ln()
    node = OxmlElement("a:tailEnd")
    node.set("type", "triangle")
    node.set("w", "sm")
    node.set("len", "sm")
    ln.append(node)


def box(slide, x, y, w, h, fill=WHITE, line=GRAY, width=1, radius=True, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, u(x), u(y), u(w), u(h))
    s.name = name or "Box"
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(width)
    if radius:
        s.adjustments[0] = 0.08
    return clean(s)


def text(slide, x, y, w, h, value, size=12, color=INK, bold=False, italic=False,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, font=FONT, name=None):
    s = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    s.name = name or f"Text: {value[:24]}"
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = u(1)
    tf.margin_top = tf.margin_bottom = u(1)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return clean(s)


def rich(slide, x, y, w, h, runs, size=12, align=PP_ALIGN.CENTER, name=None):
    s = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    s.name = name or "Rich text"
    tf = s.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = u(1)
    tf.margin_top = tf.margin_bottom = u(1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    for item in runs:
        value, font, bold, italic, color = item[:5]
        r = p.add_run()
        r.text = value
        r.font.name = font
        r.font.size = Pt(size if len(item) == 5 else size * 0.72)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)
        if len(item) > 5:
            r._r.get_or_add_rPr().set("baseline", str(item[5]))
    return clean(s)


def line(slide, x1, y1, x2, y2, color=GRAY, width=1.15, end=False, dashed=False, name=None):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2))
    s.name = name or "Connector"
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if end:
        arrow(s)
    if dashed:
        dash(s)
    return clean(s)


def path(slide, points, color=GRAY, width=1.15, end=True, dashed=False, name="Path"):
    for i, (a, b) in enumerate(zip(points, points[1:])):
        line(slide, *a, *b, color, width, end and i == len(points) - 2, dashed, f"{name} {i + 1}")


def poly(slide, points, fill, edge, width=0.8, name="Polygon"):
    x0, y0 = min(x for x, _ in points), min(y for _, y in points)
    local = [(x - x0, y - y0) for x, y in points]
    b = slide.shapes.build_freeform(local[0][0], local[0][1], scale=Inches(1) / PX_PER_INCH)
    b.add_line_segments(local[1:], close=True)
    s = b.convert_to_shape(u(x0), u(y0))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(edge)
    s.line.width = Pt(width)
    return clean(s)


def oval(slide, x, y, w, h, fill=None, edge=GRAY, width=1, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(w), u(h))
    s.name = name or "Oval"
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(edge)
    s.line.width = Pt(width)
    return clean(s)


def scan_blob(image, crop):
    stream = BytesIO()
    image.crop(crop).save(stream, format="PNG")
    stream.seek(0)
    return stream


def picture(slide, blob, x, y, w, h, name):
    blob.seek(0)
    s = slide.shapes.add_picture(blob, u(x), u(y), u(w), u(h))
    s.name = name
    return clean(s)


def grid(slide, x, y, w, h, cols, rows, color, fill, name):
    box(slide, x, y, w, h, fill, color, 0.8, False, name)
    for i in range(1, cols):
        line(slide, x + w * i / cols, y, x + w * i / cols, y + h, color, 0.45, name=f"{name} v{i}")
    for i in range(1, rows):
        line(slide, x, y + h * i / rows, x + w, y + h * i / rows, color, 0.45, name=f"{name} h{i}")


def mesh_plane(slide, x, y, w, h, slant, fill, edge, name, cols=5, rows=3):
    pts = [(x + slant, y), (x + w, y), (x + w - slant, y + h), (x, y + h)]
    poly(slide, pts, fill, edge, 0.65, name)
    for i in range(1, cols):
        xt = x + slant + (w - slant) * i / cols
        xb = x + (w - slant) * i / cols
        line(slide, xt, y, xb, y + h, edge, 0.35, name=f"{name} grid v{i}")
    for i in range(1, rows):
        yy = y + h * i / rows
        line(slide, x + slant * (1 - i / rows), yy,
             x + w - slant * i / rows, yy, edge, 0.35, name=f"{name} grid h{i}")


def query_pill(slide, x, y, color, name):
    box(slide, x, y, 79, 21, "F1EEF1", color, 0.7, True, name)
    line(slide, x + 21, y + 8, x + 41, y + 8, color, 0.7, name=f"{name} line 1")
    line(slide, x + 18, y + 13, x + 49, y + 13, color, 0.7, name=f"{name} line 2")


def globe(slide, x, y, name):
    oval(slide, x, y, 58, 58, None, PURPLE, 1.1, name)
    oval(slide, x + 13, y + 5, 32, 48, None, PURPLE, 0.75, f"{name} meridians")
    line(slide, x + 29, y + 4, x + 29, y + 54, PURPLE, 0.65, name=f"{name} axis")
    line(slide, x + 6, y + 22, x + 52, y + 22, PURPLE, 0.65, name=f"{name} latitude 1")
    line(slide, x + 5, y + 35, x + 53, y + 35, PURPLE, 0.65, name=f"{name} latitude 2")


def lungs(slide, x, y, name):
    grid(slide, x, y, 57, 58, 3, 3, PURPLE, "F6F1F6", name)
    line(slide, x + 28, y + 11, x + 28, y + 27, PURPLE, 1.1, name=f"{name} trachea")
    left = [(x + 26, y + 26), (x + 20, y + 23), (x + 17, y + 30), (x + 12, y + 42),
            (x + 18, y + 49), (x + 27, y + 45), (x + 29, y + 36)]
    right = [(x + 30, y + 26), (x + 36, y + 23), (x + 40, y + 30), (x + 45, y + 42),
             (x + 39, y + 49), (x + 30, y + 45), (x + 28, y + 36)]
    poly(slide, left, "A999AD", PURPLE, 0.8, f"{name} left lung")
    poly(slide, right, "A999AD", PURPLE, 0.8, f"{name} right lung")


def clock(slide, x, y, name):
    oval(slide, x, y, 50, 50, BG, INK, 1.2, name)
    cx, cy = x + 25, y + 25
    line(slide, cx, cy, cx, y + 10, INK, 1.1, name=f"{name} hour")
    line(slide, cx, cy, x + 39, y + 34, INK, 1.1, name=f"{name} minute")
    oval(slide, cx - 2, cy - 2, 4, 4, INK, INK, 0.3, name=f"{name} pin")
    for dx, dy in [(25, 5), (45, 25), (25, 45), (5, 25)]:
        oval(slide, x + dx - 1, y + dy - 1, 2, 2, INK, INK, 0.2, name=f"{name} tick")


def compare_icon(slide, x, y, name):
    oval(slide, x, y, 55, 55, BG, ORANGE, 1.15, name)
    nodes = [(13, 36), (22, 16), (34, 28), (42, 10), (44, 40)]
    for i, j in [(0, 1), (1, 2), (2, 3), (2, 4)]:
        line(slide, x + nodes[i][0], y + nodes[i][1], x + nodes[j][0], y + nodes[j][1], ORANGE, 0.7,
             dashed=True, name=f"{name} edge {i}-{j}")
    for i, (dx, dy) in enumerate(nodes):
        oval(slide, x + dx - 3, y + dy - 3, 6, 6, WHITE, ORANGE, 0.8, name=f"{name} node {i + 1}")


def draw_routes(slide):
    line(slide, 12, 255, 324, 255, GRAY, 1.1, True, name="Evidence route")
    path(slide, [(397, 198), (397, 126), (817, 126)], BLUE, 1.35, True, name="Encoder to dense memory")
    line(slide, 970, 126, 1260, 126, BLUE, 1.35, True, name="Dense memory to future study")
    path(slide, [(1105, 126), (1105, 246), (1161, 246), (1161, 299)], BLUE, 1.25, True,
         name="Dense memory to representation")
    path(slide, [(504, 302), (522, 302), (522, 397), (539, 397)], GRAY, 1.1, True,
         name="Encoder to resampler")
    line(slide, 627, 400, 650, 400, GRAY, 1.15, True, name="Resampler to visual tokens")
    line(slide, 758, 401, 803, 401, PURPLE, 1.15, True, name="Visual tokens to slots")
    line(slide, 1092, 401, 1148, 401, PURPLE, 1.15, True, name="Slots to representation")
    line(slide, 313, 453, 650, 453, GRAY, 1.1, True, name="Text evidence route")
    line(slide, 759, 453, 803, 453, GRAY, 1.1, True, name="Text tokens to slots")
    line(slide, 16, 487, 803, 487, GRAY, 0.85, name="Evidence baseline")
    line(slide, 803, 487, 1100, 487, PURPLE, 1.05, name="State baseline")
    line(slide, 724, 487, 724, 613, BLUE, 1.2, True, name="Dense input to predictor")
    path(slide, [(1016, 487), (1016, 536), (823, 536), (823, 613)], PURPLE, 1.15, True,
         name="Slots input to predictor")
    line(slide, 922, 554, 1290, 554, ORANGE, 1.25, True, name="Horizon route")
    line(slide, 922, 554, 922, 613, ORANGE, 1.1, True, name="Transition input to predictor")
    line(slide, 1020, 554, 1020, 613, ORANGE, 1.1, True, name="Horizon input to predictor")


def draw_evidence(slide, scans):
    text(slide, 65, 141, 190, 34, "Evidence ≤ t", 15.5, name="Evidence heading")
    frames = [(27, 189, 80, 91, 0), (117, 189, 84, 91, 1), (211, 189, 80, 91, 2)]
    interiors = [(32, 194, 69, 81), (123, 194, 72, 81), (217, 194, 68, 81)]
    for (x, y, w, h, idx), (ix, iy, iw, ih) in zip(frames, interiors):
        box(slide, x, y, w, h, WHITE, GRAY, 1, True, f"Evidence scan frame {idx + 1}")
        picture(slide, scans[idx], ix, iy, iw, ih, f"Evidence scan {idx + 1}")

    box(slide, 332, 198, 172, 200, None, GRAY, 1, True, "Medical encoder")
    text(slide, 340, 210, 156, 49, "Medical 2D / 3D\nencoder", 13.2, name="Medical encoder label")
    for row in range(3):
        for col in range(3):
            fill = BLUE_MID if col > 0 else "C9CDD0"
            box(slide, 347 + col * 15, 270 + row * 23, 12, 19, fill, "89919A", 0.5, False,
                f"2D token {row + 1}-{col + 1}")
    box(slide, 361, 330, 11, 8, "C9CDD0", "89919A", 0.45, False, "2D token small")
    grid(slide, 431, 278, 45, 59, 3, 3, BLUE, "9EB3C5", "Encoder cube front")
    poly(slide, [(431, 278), (443, 268), (488, 268), (476, 278)], "BDCAD5", BLUE, 0.65, "Encoder cube top")
    poly(slide, [(476, 278), (488, 268), (488, 327), (476, 337)], "7893AA", BLUE, 0.65,
         "Encoder cube side")
    for i in range(1, 3):
        line(slide, 443 + i * 15, 268, 431 + i * 15, 278, BLUE, 0.35, name=f"Cube top grid {i}")
        line(slide, 476, 278 + i * 19.7, 488, 268 + i * 19.7, BLUE, 0.35, name=f"Cube side grid {i}")
    path(slide, [(381, 346), (381, 357), (387, 363), (449, 363), (456, 357), (456, 346)], GRAY, 0.9,
         False, name="Encoder token brace")
    line(slide, 418, 363, 418, 398, GRAY, 0.85, name="Encoder output stem")

    box(slide, 539, 365, 88, 70, None, GRAY, 1, True, "Resampler and projector")
    text(slide, 543, 371, 80, 57, "Resampler\n+ projector", 10.5, name="Resampler label")
    for i in range(3):
        box(slide, 656 + i * 35, 386, 27, 30, "86A1B6", BLUE, 0.8, True, f"Visual token {i + 1}")
    text(slide, 761, 384, 31, 31, "···", 14, BLUE, True, name="Visual token ellipsis")

    # Symbol-only report, chat, and compact text evidence.
    box(slide, 19, 406, 47, 62, None, GRAY, 0.9, True, "Report icon")
    box(slide, 29, 416, 14, 16, "D6D7D5", GRAY, 0.55, False, "Report thumbnail")
    for i, length in enumerate((24, 30, 25, 20)):
        line(slide, 29, 440 + i * 7, 29 + length, 440 + i * 7, GRAY, 0.65, name=f"Report line {i + 1}")
    bubble = box(slide, 78, 414, 40, 37, None, GRAY, 0.9, True, "Chat icon")
    poly(slide, [(89, 449), (93, 449), (89, 456)], BG, GRAY, 0.7, "Chat tail")
    line(slide, 89, 426, 108, 426, GRAY, 0.65, name="Chat line 1")
    line(slide, 89, 434, 104, 434, GRAY, 0.65, name="Chat line 2")
    for i, x in enumerate((131, 201)):
        box(slide, x, 443, 59, 19, "E4E4E2", GRAY, 0.7, True, f"Text token {i + 1}")
        line(slide, x + 25, 452, x + 39, 452, GRAY, 0.55, name=f"Text token line {i + 1}")
    text(slide, 271, 439, 32, 24, "···", 14, INK, True, name="Text token ellipsis")


def draw_dense_and_slots(slide):
    rich(slide, 775, 16, 245, 36,
         [("Dense memory  ", FONT, False, False, INK), ("Dₜ¹:ᴸ", MATH, False, True, INK)],
         15.5, name="Dense memory heading")
    planes = [
        (807, 143, 139, 24, 19, "A5B4C1"), (819, 126, 136, 22, 18, "C0CAD2"),
        (828, 108, 126, 21, 17, "D0D7DC"), (834, 92, 119, 20, 17, "BAC6CF"),
        (840, 76, 109, 18, 16, "CED6DB"), (842, 61, 96, 28, 25, "AAB8C5"),
    ]
    for i, args in enumerate(planes):
        mesh_plane(slide, *args, BLUE, f"Dense memory layer {i + 1}", cols=5, rows=3)

    text(slide, 770, 255, 175, 37, "VLM + LoRA", 16, name="VLM label")
    rich(slide, 862, 302, 205, 37,
         [("8 generic slots  ", FONT, False, False, INK), ("Sₜ", MATH, False, True, INK)],
         13.2, name="Slot heading")
    path(slide, [(809, 355), (809, 351), (815, 345), (943, 345), (950, 350),
                 (957, 345), (1087, 345), (1093, 351), (1093, 355)], PURPLE, 1.05, False,
         name="Slot overbrace")
    for i, x in enumerate((811, 847, 883, 919, 955, 991, 1027, 1063)):
        box(slide, x, 367, 26, 101, PURPLE_MID, PURPLE, 1, True, f"Generic slot {i + 1}")
        top = box(slide, x + 1.1, 368, 23.8, 69, "F8F5F8", None, 0, True, f"Generic slot {i + 1} top")
        top.adjustments[0] = 0.35
        box(slide, x + 1.2, 430, 23.6, 10, PURPLE_MID, None, 0, False, f"Generic slot {i + 1} seam fill")
        line(slide, x + 4, 432, x + 22, 432, PURPLE, 0.65, name=f"Generic slot {i + 1} seam")

    path(slide, [(1156, 302), (1162, 302), (1167, 308), (1167, 365), (1171, 377),
                 (1176, 383), (1171, 389), (1167, 401), (1167, 477), (1162, 489),
                 (1156, 490)], INK, 1.05, False, name="Representation brace")
    rich(slide, 1172, 344, 149, 48,
         [("Rₜ = (Dₜ¹:ᴸ, Sₜ)", MATH, False, True, INK)], 14.5, name="Representation formula")
    box(slide, 1178, 394, 130, 53, "EEE9EF", PURPLE, 0.9, True, "Representation property")
    text(slide, 1182, 399, 122, 43, "task- &\nhorizon-independent", 9.3, name="Representation property label")


def draw_future_target(slide, future_scan):
    rich(slide, 1268, 16, 201, 38,
         [("Future study  ", FONT, False, False, INK), ("Xₜ⁺", MATH, False, True, INK)],
         15.5, name="Future study heading")
    box(slide, 1283, 65, 99, 118, None, GRAY, 0.9, True, "Future study frame")
    picture(slide, future_scan, 1288, 70, 89, 108, "Future study scan")
    line(slide, 1392, 148, 1422, 148, GREEN, 1.15, True, name="Future scan to EMA")
    text(slide, 1423, 67, 126, 48, "EMA target ·\nstop-gradient", 11.7, align=PP_ALIGN.LEFT,
         name="EMA target label")
    ema = box(slide, 1428, 121, 87, 61, "EEF3EE", GREEN, 1, True, "EMA target module")
    dash(ema)
    text(slide, 1448, 128, 47, 45, "↻", 29, GREEN, font="Arial Unicode MS", name="EMA refresh glyph")
    line(slide, 1516, 148, 1546, 148, GREEN, 1.1, True, name="EMA to target tokens")
    box(slide, 1557, 119, 52, 58, "B9C8BA", GREEN, 0.85, True, "Target embedding")
    grid(slide, 1634, 119, 54, 58, 4, 4, GREEN, "C7D2C8", "Target grid coarse")
    grid(slide, 1714, 122, 50, 53, 7, 7, GREEN, "D6DED6", "Target grid medium")
    grid(slide, 1781, 130, 38, 38, 9, 9, GREEN, "DEE5DE", "Target grid fine")
    text(slide, 1822, 128, 38, 39, "···", 15, INK, True, name="Target ellipsis")
    line(slide, 1555, 92, 1840, 92, GREEN, 1, dashed=True, name="Target lane top")
    line(slide, 1281, 200, 1840, 200, GREEN, 1, dashed=True, name="Target lane bottom")
    line(slide, 1782, 200, 1782, 454, GREEN, 1.15, True, name="Target to comparison")


def draw_queries(slide):
    rich(slide, 1344, 252, 282, 43,
         [("Read-only queries  ", FONT, False, False, INK), ("Q", MATH, False, True, INK)],
         15, name="Read-only queries heading")
    box(slide, 1325, 298, 409, 211, None, GRAY, 0.95, True, "Read-only query panel")
    line(slide, 1333, 405, 1726, 405, LIGHT, 0.85, name="Query panel divider")

    rich(slide, 1344, 313, 160, 33,
         [("task ", FONT, False, False, INK), ("q → e", MATH, False, True, INK),
          ("q", MATH, False, True, INK, -25000)], 12.2,
         align=PP_ALIGN.LEFT, name="Task query label")
    query_pill(slide, 1345, 355, PURPLE, "Task query token")
    line(slide, 1427, 365, 1460, 365, PURPLE, 1, True, name="Task token to embedding")
    box(slide, 1470, 348, 29, 30, "A99BB1", PURPLE, 0.8, True, "Task embedding")
    path(slide, [(1506, 365), (1553, 365), (1553, 354), (1579, 342)], PURPLE, 1, True,
         name="Task embedding to globe")
    globe(slide, 1582, 317, "Global task icon")
    line(slide, 1641, 346, 1664, 346, PURPLE, 1, True, name="Global task to dense task")
    lungs(slide, 1663, 318, "Dense task icon")

    rich(slide, 1344, 418, 284, 35,
         [("transition ", FONT, False, False, INK), ("q + e(h) → Aₜ,ₕ", MATH, False, True, INK)],
         12.2, align=PP_ALIGN.LEFT, name="Transition query label")
    query_pill(slide, 1345, 455, ORANGE, "Transition query token")
    line(slide, 1427, 466, 1468, 466, ORANGE, 1, True, name="Transition token to action")
    box(slide, 1478, 454, 30, 31, ORANGE_MID, ORANGE, 0.85, True, "Transition embedding")
    path(slide, [(1514, 469), (1597, 469), (1597, 509), (1597, 625)], ORANGE, 1.05, True,
         name="Transition embedding output")

    clock(slide, 1274, 518, "Horizon clock")
    text(slide, 1270, 567, 59, 27, "e(h)", 13, italic=True, font=MATH, name="Horizon clock label")
    path(slide, [(1299, 518), (1299, 472), (1337, 472)], ORANGE, 1.05, True,
         name="Horizon to transition query")
    path(slide, [(1169, 554), (1169, 494)], ORANGE, 1.1, True, name="Feedback to representation")


def draw_predictor(slide):
    text(slide, 716, 557, 70, 34, "Dₜ¹:ᴸ", 13, italic=True, font=MATH, name="Dense predictor input label")
    text(slide, 811, 557, 46, 34, "Sₜ", 13, italic=True, font=MATH, name="Slot predictor input label")
    text(slide, 906, 557, 63, 34, "Aₜ,ₕ", 13, italic=True, font=MATH, name="Action predictor input label")
    text(slide, 1000, 557, 57, 34, "e(h)", 13, italic=True, font=MATH, name="Horizon predictor input label")
    box(slide, 669, 613, 378, 154, None, ORANGE, 1, True, "Latent predictor")
    text(slide, 734, 627, 247, 34, "Latent predictor", 15.5, name="Latent predictor label")

    inputs = [(738, 675), (738, 706), (738, 738)]
    h1 = [(812, 693), (812, 721)]
    h2 = [(847, 676), (847, 707), (847, 739)]
    out = (900, 705)
    for i, a in enumerate(inputs):
        for j, b in enumerate(h1):
            if i != 1 or j == 1:
                line(slide, *a, *b, ORANGE if i != 1 else GRAY, 0.65, name=f"Predictor edge i{i}h{j}")
    for i, a in enumerate(h1):
        for j, b in enumerate(h2):
            line(slide, *a, *b, GRAY, 0.65, name=f"Predictor edge h{i}m{j}")
    for i, a in enumerate(h2):
        line(slide, *a, *out, GRAY, 0.65, name=f"Predictor edge m{i}o")
    for i, (x, y) in enumerate(inputs):
        oval(slide, x - 8, y - 8, 16, 16, BG, ORANGE, 0.9, name=f"Predictor input node {i + 1}")
    for i, (x, y) in enumerate(h1 + h2 + [out]):
        oval(slide, x - 9, y - 9, 18, 18, BG, GRAY, 0.9, name=f"Predictor node {i + 1}")
    line(slide, 922, 705, 978, 705, GRAY, 1.1, True, name="Predictor internal output")
    line(slide, 1047, 694, 1128, 694, ORANGE, 1.15, True, name="Predictor output")

    line(slide, 1119, 625, 1701, 625, ORANGE, 1, dashed=True, name="Prediction lane top")
    line(slide, 1119, 772, 1701, 772, ORANGE, 1, dashed=True, name="Prediction lane bottom")
    box(slide, 1148, 663, 68, 75, "D89C87", ORANGE, 0.95, True, "Predicted embedding")
    rich(slide, 1264, 631, 250, 41,
         [("Prediction  ", FONT, False, False, INK), ("Ŷₜ,ₕ", MATH, False, True, INK)],
         15.2, name="Prediction heading")
    grid(slide, 1271, 678, 76, 77, 4, 4, ORANGE, "EED8CF", "Prediction grid coarse")
    grid(slide, 1388, 682, 64, 72, 7, 8, ORANGE, "E8C8BC", "Prediction grid medium")
    grid(slide, 1490, 691, 52, 53, 10, 10, ORANGE, "E6C0B2", "Prediction grid fine")
    text(slide, 1567, 687, 45, 41, "···", 16, INK, True, name="Prediction ellipsis")

    text(slide, 1570, 516, 73, 31, "Aₜ,ₕ", 13, italic=True, font=MATH, name="Action output label")
    compare_icon(slide, 1755, 457, "Future comparison")
    path(slide, [(1782, 512), (1782, 710), (1654, 710)], ORANGE, 1.15, True, name="Future loss route")
    rich(slide, 1648, 719, 211, 42,
         [("Future loss shapes  ", FONT, False, False, INK), ("Rₜ", MATH, False, True, INK)],
         14.2, name="Future loss label")


def build():
    with Image.open(SOURCE) as image:
        scans = [
            scan_blob(image, (32, 194, 101, 275)),
            scan_blob(image, (123, 194, 195, 275)),
            scan_blob(image, (217, 194, 285, 275)),
        ]
        future_scan = scan_blob(image, (1288, 70, 1377, 178))

    prs = Presentation()
    prs.slide_width = u(W)
    prs.slide_height = u(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    draw_routes(slide)
    draw_evidence(slide, scans)
    draw_dense_and_slots(slide)
    draw_future_target(slide, future_scan)
    draw_queries(slide)
    draw_predictor(slide)

    prs.core_properties.title = "MedWorld-JEPA editable architecture figure"
    prs.core_properties.subject = f"Editable recreation of {SOURCE.name}"
    prs.core_properties.author = "Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
