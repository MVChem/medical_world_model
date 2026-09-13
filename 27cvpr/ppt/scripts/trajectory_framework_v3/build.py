"""Make the existing v2 PowerPoint easier to edit, preserving its artwork.

The input is a snapshot of the actual v2 deck. No artwork is regenerated from
the old drawing script. Text and arrows become independently selectable;
only small component groups and decorative effects remain grouped.
"""
from __future__ import annotations

import argparse
from copy import deepcopy
import json
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

from lxml import etree as ET
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "assets/source_v2.pptx"
OUT = HERE.parents[1] / "ppt" / f"{HERE.name}.pptx"
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
PX = 7620  # 120 design pixels per inch.


def node(tag, **attributes):
    value = OxmlElement(tag)
    for key, item in attributes.items():
        value.set(key, str(item))
    return value


def u(value):
    return round(value * PX)


def shape_name(element):
    return element.find(".//p:cNvPr", NS).get("name")


def rename(element, value):
    element.find(".//p:cNvPr", NS).set("name", value)
    return element


def children(element):
    return [child for child in element if child.tag in {
        f"{{{NS['p']}}}sp", f"{{{NS['p']}}}grpSp", f"{{{NS['p']}}}cxnSp"}]


def find_named(element, name):
    for shape in children(element):
        if shape_name(shape) == name:
            return shape
        if shape.tag.endswith("}grpSp"):
            found = find_named(shape, name)
            if found is not None:
                return found
    return None


def check_group_transforms(root):
    # This source uses absolute slide coordinates inside all groups. Check this
    # before lifting children, so a transformed source is never silently moved.
    for transform in root.findall(".//p:grpSpPr/a:xfrm", NS):
        values = [transform.find("a:"+tag, NS) for tag in ["off", "ext", "chOff", "chExt"]]
        assert dict(values[0].attrib) == dict(values[2].attrib)
        assert dict(values[1].attrib) == dict(values[3].attrib)
        assert not any(transform.get(k) not in (None, "0", "false")
                       for k in ["rot", "flipH", "flipV"])


def add_component_group(slide, name, elements):
    """A single-level group for one card, one point, or one visual effect."""
    group = slide.shapes.add_group_shape()
    group.name = name
    for index, element in enumerate(elements, 1):
        copied = deepcopy(element)
        rename(copied, f"{name} / {index:02} {shape_name(copied)}")
        group.shapes._spTree.append(copied)
    group.shapes._recalculate_extents()
    return group._element


def add_native_arrow(slide, old, name):
    """Replace a 7-vertex arrow outline with a native adjustable rightArrow."""
    props = old.find("p:spPr", NS)
    transform = props.find("a:xfrm", NS)
    off, extent = transform.find("a:off", NS), transform.find("a:ext", NS)
    path = props.find("a:custGeom/a:pathLst/a:path", NS)
    points = [
        (int(off.get("x")) + int(pt.get("x"))*int(extent.get("cx"))/int(path.get("w")),
         int(off.get("y")) + int(pt.get("y"))*int(extent.get("cy"))/int(path.get("h")))
        for pt in path.findall(".//a:pt", NS)]
    assert len(points) == 7
    start = tuple((a+b)/2 for a, b in zip(points[0], points[6]))
    end = points[3]
    shaft = math.dist(points[0], points[6])
    head = math.dist(points[2], points[4])
    head_base = tuple((a+b)/2 for a, b in zip(points[2], points[4]))
    head_length = math.dist(head_base, end)
    length = math.dist(start, end)
    angle = math.degrees(math.atan2(end[1]-start[1], end[0]-start[0])) % 360
    cx, cy = ((a+b)/2 for a, b in zip(start, end))
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
        round(cx-length/2), round(cy-head/2), round(length), round(head))
    arrow.name = name
    arrow.rotation = angle
    arrow.adjustments[0] = shaft/head
    arrow.adjustments[1] = head_length/min(length, head)
    arrow.fill.background()
    gradient = deepcopy(props.find("a:gradFill", NS))
    gradient.find("a:lin", NS).set("ang", "0")
    gradient.set("rotWithShape", "1")
    new_props = arrow._element.spPr
    new_props.remove(new_props.find("a:noFill", NS))
    new_props.insert(2, gradient)
    arrow.line.fill.background()
    arrow.shadow.inherit = False
    for ref in arrow._element.xpath("./p:style/a:effectRef"):
        ref.set("idx", "0")
    return arrow._element


def velocity_text(slide):
    """Replace the old outlined v glyph with actual editable text."""
    shape = slide.shapes.add_textbox(u(493), u(313), u(72), u(106))
    shape.name = "TEXT 01.3 | 速度符号 v"
    frame = shape.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_top = frame.margin_bottom = 0
    frame.margin_left = frame.margin_right = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1
    run = paragraph.add_run()
    run.text = "v"
    run.font.name = "Arial"
    run.font.bold = run.font.italic = True
    run.font.size = Pt(102*72/120)
    run.font.color.rgb = RGBColor.from_string("2A82FA")
    return shape._element


def make_components(slide, panel, number, deferred_text):
    tree = slide.shapes._spTree

    def lift(element, name):
        result = rename(deepcopy(element), name)
        tree.append(result)
        return result

    for component in children(panel):
        name = shape_name(component)
        texts = component.findall("./p:txBody//a:t", NS)
        if texts:
            line = 1 if "line 1" in name else 2
            label = "".join(t.text or "" for t in texts)
            deferred_text.append(rename(deepcopy(component), f"TEXT {number:02}.{line} | {label}"))
        elif "pale " in name:
            add_component_group(slide, f"FX {number:02} | 背景柔光", children(component))
        elif "frame " in name and number == 1:
            index = int(name.split("frame ")[1][:2])
            add_component_group(slide, f"CARD 01.{index} | 透视卡片", children(component))
        elif "forward generative trajectory" in name:
            elements = children(component)
            light = find_named(component, "Trajectory soft white light")
            rim = find_named(component, "Trajectory white rim")
            add_component_group(slide, "FX 01.1 | 轨迹柔光及白色衬线", children(light)+[rim])
            lift(find_named(component, "Trajectory blue Bezier"), "ARROW 01.1 | 蓝色轨迹杆（编辑顶点/线宽）")
            lift(find_named(component, "Trajectory arrow head"), "ARROW 01.2 | 蓝色轨迹箭头头部")
            for index in range(1, 4):
                point = [e for e in elements if shape_name(e).startswith(f"Time point {index}:")]
                add_component_group(slide, f"POINT 01.{index} | 轨迹时间点", point)
        elif "blue velocity mark" in name:
            for index, element in enumerate(children(component)[:3], 1):
                lift(element, f"FX 01.2.{index} | 速度短线")
            deferred_text.append(velocity_text(slide))
        elif "cycle arrows" in name:
            for index, element in enumerate(children(component), 1):
                direction = "上方循环箭头" if index == 1 else "下方循环箭头"
                lift(element, f"ARROW 02.{index} | {direction}（编辑顶点）")
        elif "image card" in name:
            index = {"left": 1, "middle": 2, "right": 3}[name.split()[1]]
            add_component_group(slide, f"CARD 02.{index} | 图像卡片", children(component))
        elif "consistency link" in name:
            lift(component, f"ARROW 02.{int(name[-1])+2} | 原生双向箭头（黄色手柄）")
        elif "outward coral arrows" in name or "inward green arrows" in name:
            inward = "inward" in name
            for index, element in enumerate(children(component), 1):
                label = "绿色向内" if inward else "红色向外"
                add_native_arrow(slide, element, f"ARROW 03.{'I' if inward else 'O'}{index} | {label}箭头（黄色手柄）")
        elif "concentric target" in name:
            for index, element in enumerate(children(component), 1):
                lift(element, f"TARGET 03.{index} | 独立同心圆")
        else:
            raise ValueError(f"Unclassified component: {name}")


def build():
    prs = Presentation(SOURCE)
    slide = prs.slides[0]
    tree = slide.shapes._spTree
    check_group_transforms(tree)
    panels = [deepcopy(group) for group in children(tree)]
    assert len(panels) == 3
    for element in children(tree):
        tree.remove(element)
    text = []
    for index, panel in enumerate(panels, 1):
        make_components(slide, panel, index, text)
    # Keep every text box top-level and above graphical elements in z-order.
    for element in text:
        tree.append(element)
    for index, prop in enumerate(tree.findall(".//p:cNvPr", NS), 2):
        prop.set("id", str(index))
    # No grouped, locked or outlined text remains.
    for locks in tree.xpath(".//a:spLocks | .//a:grpSpLocks"):
        for key in list(locks.attrib):
            if key != "noGrp":
                del locks.attrib[key]
    prs.core_properties.title = "Trajectory Framework v3 — Independent Editable Components"
    slide.notes_slide.notes_text_frame.text = (
        "All seven text boxes are independent top-level objects, including the velocity v. "
        "Each arrow is independently selectable. Twelve radial arrows and two double-headed "
        "arrows are native preset shapes with adjustment handles. The trajectory shaft and "
        "head are separate; cycle arrows are native freeforms editable using Edit Points. "
        "Small card, time-point and decorative-light groups have only one nesting level. "
        "Use the Selection Pane and FX-prefixed names to select or hide decorations. "
        "Decorative glow does not automatically follow subsequent edits to a curve. "
        "The velocity symbol is now editable Arial Bold Italic text and visually approximates "
        "the former outlined glyph. All other source artwork comes directly from source_v2.pptx.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    validate()
    inventory(prs)


def validate():
    with ZipFile(OUT) as archive:
        root = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
        tree = root.find("p:cSld/p:spTree", NS)
        assert not root.findall(".//p:pic", NS)
        assert not [n for n in archive.namelist() if n.startswith("ppt/media/")]
        assert not root.findall(".//p:grpSp/p:grpSp", NS), "Keep only one level of grouping"
        assert not root.findall(".//p:grpSp//a:t", NS), "Text must never live inside a group"
        labels = [s for s in tree.findall("p:sp", NS) if s.findall("p:txBody//a:t", NS)]
        assert len(labels) == 7
        assert sorted("".join(s.xpath(".//a:t/text()", namespaces=NS)) for s in labels) == sorted([
            "v", "1)  Generative Trajectory", "Representation", "2)  Trajectory Consistency",
            "Evolution", "3)  Real-Centered Trajectory", "Optimization"])
        assert len(tree.findall('p:sp/p:spPr/a:prstGeom[@prst="rightArrow"]', NS)) == 12
        assert len(tree.findall('p:sp/p:spPr/a:prstGeom[@prst="leftRightArrow"]', NS)) == 2
        names = [prop.get("name") for prop in root.findall(".//p:cNvPr", NS)]
        assert len(names) == len(set(names)), "Selection Pane names should be unique"
        ids = [prop.get("id") for prop in root.findall(".//p:cNvPr", NS)]
        assert len(ids) == len(set(ids))
        print(f"Validated {len(children(tree))} top-level components; 7 independent text boxes; "
              "14 adjustable native arrows; no embedded images or nested component groups.")


def inventory(prs):
    data = []
    for shape in prs.slides[0].shapes:
        item = {"name": shape.name, "type": str(shape.shape_type),
                "bounds_px": [round(n/PX, 3) for n in [shape.left, shape.top, shape.width, shape.height]]}
        if shape.has_text_frame and shape.text:
            item["text"] = shape.text
        if shape.shape_type == 6:
            item["children"] = [s.name for s in shape.shapes]
        data.append(item)
    (HERE / "components.json").write_text(json.dumps(data, ensure_ascii=False, indent=2)+"\n")


def render():
    import fitz
    command = shutil.which("libreoffice") or shutil.which("soffice")
    if not command:
        raise RuntimeError("LibreOffice is required to render the existing PPTX.")
    with tempfile.TemporaryDirectory(prefix="trajectory_v3_") as folder:
        folder = Path(folder)
        result = subprocess.run([
            command, f"-env:UserInstallation={(folder/'profile').as_uri()}",
            "--headless", "--convert-to", "pdf:impress_pdf_Export",
            "--outdir", str(folder), str(OUT)
        ], capture_output=True, text=True, timeout=120)
        pdf = folder / OUT.with_suffix(".pdf").name
        if result.returncode or not pdf.exists():
            raise RuntimeError(result.stdout+result.stderr)
        with fitz.open(pdf) as doc:
            assert len(doc) == 1
            page = doc[0]
            assert abs(page.rect.width/page.rect.height-3) < .001
            assert "Optimization" in page.get_text()
            pix = page.get_pixmap(matrix=fitz.Matrix((3072-.001)/page.rect.width,
                                                     (1024-.001)/page.rect.height), alpha=False)
            assert (pix.width, pix.height) == (3072, 1024)
            pix.save(OUT.with_suffix(".png"))
    print(OUT.with_suffix(".png"))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument("--render-only", action="store_true")
    parser.add_argument("--no-render", action="store_true")
    args = parser.parse_args()
    if args.render_only:
        render()
    else:
        if OUT.exists() and not args.overwrite:
            parser.error("PPT already exists; use --overwrite to rebuild, or --render-only to preserve edits.")
        build()
        if not args.no_render:
            render()
    print(OUT)
