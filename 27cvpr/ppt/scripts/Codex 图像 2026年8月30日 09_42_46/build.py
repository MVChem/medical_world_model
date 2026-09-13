from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


NAME = "Codex 图像 2026年8月30日 09_42_46"
PPT_ROOT = Path(__file__).resolve().parents[2]
ASSETS = PPT_ROOT / "source/mimic_appendix_cases"
OUTPUT = PPT_ROOT / "ppt" / f"{NAME}.pptx"

W, H, PPI = 1491, 1055, 100
FONT = "Arial"
FONT_NARROW = "Roboto Condensed"
BG = "FAF8F2"
INK = "0A2944"
BODY = "242629"
MUTED = "4E5153"
BLUE = "0B4A76"
BLUE_HEADER = "3F6589"
BLUE_PALE = "EDF3F6"
TEAL = "0B626A"
TEAL_HEADER = "5F9A9B"
TEAL_PALE = "EFF6F3"
GOLD_HEADER = "C6A75D"
GOLD_PALE = "FBF5E9"
LINE_BLUE = "AFC5CE"
LINE_TEAL = "A8C5C1"
LINE_GOLD = "E2D1AA"
WHITE = "FFFFFF"


def u(px):
    return Inches(px / PPI)


def rgb(value):
    return RGBColor.from_string(value)


def clean(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def box(slide, x, y, w, h, fill=WHITE, line=None, width=0.8, rounded=True, name="Box"):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, u(x), u(y), u(w), u(h))
    shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width)
    if rounded:
        shape.adjustments[0] = 0.08
    return clean(shape)


def oval(slide, x, y, w, h, fill=None, line=INK, width=1, name="Oval"):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(w), u(h))
    shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width)
    return clean(shape)


def text(slide, x, y, w, h, value, size=10, color=BODY, bold=False,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, name=None,
         margin=1, spacing=0.94, font=FONT):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    shape.name = name or f"Text: {value[:24]}"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = u(margin)
    frame.margin_top = frame.margin_bottom = u(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = spacing
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return clean(shape)


def line(slide, x1, y1, x2, y2, color=INK, width=1, dashed=False,
         arrow=False, name="Line"):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2)
    )
    shape.name = name
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    xml_line = shape._element.spPr.get_or_add_ln()
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        xml_line.append(dash)
    if arrow:
        end = OxmlElement("a:tailEnd")
        end.set("type", "triangle")
        end.set("w", "sm")
        end.set("len", "sm")
        xml_line.append(end)
    return clean(shape)


def add_cropped_picture(slide, path, x, y, w, h, focus_x=0.5, focus_y=0.5, name="Image"):
    picture = slide.shapes.add_picture(str(path), u(x), u(y), u(w), u(h))
    picture.name = name
    with Image.open(path) as image:
        image_ratio = image.width / image.height
    box_ratio = w / h
    if image_ratio < box_ratio:
        keep = image_ratio / box_ratio
        top = max(0, min(1 - keep, focus_y - keep / 2))
        picture.crop_top = top
        picture.crop_bottom = 1 - keep - top
    elif image_ratio > box_ratio:
        keep = box_ratio / image_ratio
        left = max(0, min(1 - keep, focus_x - keep / 2))
        picture.crop_left = left
        picture.crop_right = 1 - keep - left
    box(slide, x, y, w, h, None, "65747E", 0.65, False, f"{name} border")
    return clean(picture)


def add_bullets(slide, x, y, w, h, items, size=9.6):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    shape.name = "Follow-up findings"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = u(2)
    frame.margin_top = frame.margin_bottom = u(1)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(items):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = 0.92
        run = paragraph.add_run()
        run.text = f"•  {item}"
        run.font.name = FONT_NARROW
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(BODY)
    return clean(shape)


def check_icon(slide, x, y, scale=1, color=BLUE, name="Check"):
    oval(slide, x, y, 20 * scale, 20 * scale, None, color, 1.1, f"{name} circle")
    line(slide, x + 5 * scale, y + 10 * scale, x + 8.5 * scale,
         y + 14 * scale, color, 1.2, name=f"{name} tick 1")
    line(slide, x + 8.5 * scale, y + 14 * scale, x + 15 * scale,
         y + 6 * scale, color, 1.2, name=f"{name} tick 2")


def hospital_icon(slide, x, y, color=BLUE):
    box(slide, x + 6, y + 5, 27, 31, None, color, 1, False, "Hospital outline")
    box(slide, x, y + 13, 7, 23, None, color, 1, False, "Hospital left wing")
    box(slide, x + 32, y + 13, 7, 23, None, color, 1, False, "Hospital right wing")
    line(slide, x + 16, y + 7, x + 23, y + 7, color, 1.4, name="Hospital cross h")
    line(slide, x + 19.5, y + 3.5, x + 19.5, y + 11, color, 1.4, name="Hospital cross v")
    for dx in (10, 24):
        for dy in (15, 23):
            box(slide, x + dx, y + dy, 4, 4, color, color, 0.3, False, "Hospital window")
    box(slide, x + 17, y + 28, 5, 8, None, color, 0.8, False, "Hospital door")


def eye_icon(slide, x, y, color=BODY):
    arc = slide.shapes.add_shape(MSO_SHAPE.ARC, u(x), u(y), u(20), u(12))
    arc.name = "Audit eye upper"
    arc.fill.background()
    arc.line.color.rgb = rgb(color)
    arc.line.width = Pt(0.9)
    arc.rotation = 0
    clean(arc)
    lower = slide.shapes.add_shape(MSO_SHAPE.ARC, u(x), u(y), u(20), u(12))
    lower.name = "Audit eye lower"
    lower.fill.background()
    lower.line.color.rgb = rgb(color)
    lower.line.width = Pt(0.9)
    lower.rotation = 180
    clean(lower)
    oval(slide, x + 8, y + 4, 4, 4, color, color, 0.5, "Audit eye iris")


def monitor_icon(slide, x, y, color=INK):
    box(slide, x, y, 31, 22, None, color, 1, True, "Ventilation monitor")
    line(slide, x + 4, y + 12, x + 9, y + 12, color, 0.9, name="Monitor trace 1")
    line(slide, x + 9, y + 12, x + 12, y + 7, color, 0.9, name="Monitor trace 2")
    line(slide, x + 12, y + 7, x + 16, y + 16, color, 0.9, name="Monitor trace 3")
    line(slide, x + 16, y + 16, x + 20, y + 10, color, 0.9, name="Monitor trace 4")
    line(slide, x + 20, y + 10, x + 27, y + 10, color, 0.9, name="Monitor trace 5")
    line(slide, x + 15.5, y + 22, x + 15.5, y + 28, color, 1, name="Monitor stand")
    line(slide, x + 7, y + 28, x + 24, y + 28, color, 1, name="Monitor base")


def imaging_icon(slide, x, y, color=INK):
    oval(slide, x + 4, y, 25, 25, None, color, 1, "Imaging gantry")
    oval(slide, x + 9, y + 5, 15, 15, None, color, 0.8, "Imaging aperture")
    line(slide, x, y + 29, x + 28, y + 29, color, 1, name="Imaging bed")
    line(slide, x + 5, y + 25, x + 1, y + 29, color, 1, name="Imaging bed incline")


def culture_icon(slide, x, y, color=INK):
    oval(slide, x + 2, y + 1, 28, 28, None, color, 1, "Culture dish")
    for dx, dy, size in ((8, 8, 4), (18, 6, 3), (14, 17, 4), (22, 19, 3)):
        oval(slide, x + dx, y + dy, size, size, TEAL_PALE, color, 0.7, "Culture colony")
    line(slide, x + 7, y + 24, x + 25, y + 6, color, 0.8, name="Culture swab")


def iv_icon(slide, x, y, color=INK):
    line(slide, x + 15, y, x + 15, y + 30, color, 1, name="IV pole")
    line(slide, x + 10, y + 1, x + 20, y + 1, color, 1, name="IV hook")
    box(slide, x + 5, y + 6, 15, 18, None, color, 1, True, "IV bag")
    line(slide, x + 8, y + 11, x + 17, y + 11, color, 0.7, name="IV fluid")
    line(slide, x + 12.5, y + 24, x + 12.5, y + 30, color, 0.8, name="IV tube")
    oval(slide, x + 11, y + 29, 3, 3, None, color, 0.7, "IV connector")


def vial_icon(slide, x, y, color=INK):
    box(slide, x + 8, y, 17, 5, None, color, 0.9, False, "Medication cap")
    box(slide, x + 6, y + 5, 21, 26, None, color, 1, True, "Medication vial")
    box(slide, x + 10, y + 15, 13, 8, TEAL, TEAL, 0.5, False, "Medication label")


def device_icon(slide, x, y, color=BODY):
    oval(slide, x + 7, y, 13, 13, None, color, 0.9, "Support device bulb")
    box(slide, x + 3, y + 12, 21, 14, None, color, 0.9, True, "Support device body")
    line(slide, x + 13.5, y + 26, x + 13.5, y + 34, color, 0.9, name="Support device tube")
    line(slide, x + 13.5, y + 34, x + 21, y + 38, color, 0.9, name="Support device lead")


def balance_icon(slide, x, y, color="9A8251"):
    line(slide, x + 18, y + 2, x + 18, y + 25, color, 1, name="Balance stem")
    line(slide, x + 7, y + 7, x + 29, y + 7, color, 1, name="Balance beam")
    line(slide, x + 8, y + 7, x + 4, y + 18, color, 0.8, name="Balance left cord")
    line(slide, x + 8, y + 7, x + 12, y + 18, color, 0.8, name="Balance left cord 2")
    line(slide, x + 28, y + 7, x + 24, y + 18, color, 0.8, name="Balance right cord")
    line(slide, x + 28, y + 7, x + 32, y + 18, color, 0.8, name="Balance right cord 2")
    line(slide, x + 3, y + 18, x + 13, y + 18, color, 1, name="Balance left pan")
    line(slide, x + 23, y + 18, x + 33, y + 18, color, 1, name="Balance right pan")
    line(slide, x + 10, y + 25, x + 26, y + 25, color, 1, name="Balance base")


def header(slide, x, y, w, number, label, fill):
    box(slide, x, y, w, 39, fill, None, rounded=True, name=f"Header {number}")
    number_x = {1: x + 74, 2: x + 13, 3: x + 88}[number]
    oval(slide, number_x, y + 7, 24, 24, WHITE, None, name=f"Header number {number}")
    text(slide, number_x, y + 6, 24, 25, str(number), 11.5, INK, True,
         name=f"Header number text {number}")
    text(slide, number_x + 30, y + 4, x + w - number_x - 34, 31, label, 12.0,
         WHITE, True, align=PP_ALIGN.LEFT, name=f"Header label {number}",
         font=FONT_NARROW)


def context_item(slide, x, y, w, icon, label, detail, compact=False):
    icon_x = x + w / 2 - 16
    icon(slide, icon_x, y, INK)
    text(slide, x + 2, y + 31, w - 4, 17, label, 8.8 if compact else 9.2,
         TEAL, True, name=f"Context label {label}")
    text(slide, x + 3, y + 47, w - 6, 31 if compact else 35, detail,
         7.7 if compact else 8.1, BODY, name=f"Context detail {label}", spacing=0.9)


def context_card(slide, data, x, y, w, h):
    box(slide, x, y, w, h, TEAL_PALE, LINE_TEAL, 0.8, True, "Retrospective context card")
    text(slide, x + 14, y + 8, w - 28, 22,
         "Linked retrospective context (MIMIC-IV) — not model input",
         9.7, TEAL, True, name="Retrospective context heading")
    items = data["context"]
    if len(items) == 2:
        divider_x = x + w / 2
        line(slide, divider_x, y + 38, divider_x, y + h - 12, "9DB6B6", 0.7,
             dashed=True, name="Context divider")
        for i, (icon, label, detail) in enumerate(items):
            item_x = x + i * w / 2
            icon(slide, item_x + 68, y + 40, INK)
            text(slide, item_x + 108, y + 36, 105, 20, label, 9.2, TEAL, True,
                 align=PP_ALIGN.LEFT, name=f"Context label {label}")
            text(slide, item_x + 108, y + 55, 106, 40, detail, 8.0, BODY,
                 align=PP_ALIGN.LEFT, name=f"Context detail {label}", spacing=0.9)
        return
    compact = len(items) == 5
    inner_x = x + 10
    item_w = (w - 20) / len(items)
    icon_y = y + 39 if h >= 130 else y + 37
    for i, (icon, label, detail) in enumerate(items):
        ix = inner_x + i * item_w
        if i:
            line(slide, ix, icon_y + 2, ix, y + h - 13, "9DB6B6", 0.7,
                 dashed=True, name="Context divider")
        context_item(slide, ix, icon_y, item_w, icon, label, detail, compact)


def metadata(slide, data, y):
    box(slide, 568, y, 184, 26, "EDF1F5", "7297B9", 0.8, True, "Horizon bin label")
    text(slide, 570, y + 1, 180, 23, "Horizon bin (model context)", 9.6, INK,
         name="Horizon bin label text")
    text(slide, 599, y + 25, 121, 26, data["horizon"], 14.7, BLUE, True,
         name="Horizon bin value")
    text(slide, 788, y - 1, 174, 20, "Exact elapsed time", 9.5, BODY,
         name="Elapsed time label")
    text(slide, 810, y + 19, 130, 24, data["elapsed"], 13.2, BODY,
         name="Elapsed time value")
    text(slide, 790, y + 42, 170, 18, "(observed, not input)", 8.2, BODY,
         name="Elapsed time note")


def left_case(slide, data, y, h):
    left_y = y + data.get("label_offset", 0)
    box(slide, 16, left_y, 126, 37, BLUE, None, rounded=True, name=f"{data['case']} label")
    text(slide, 19, left_y + 2, 120, 32, data["case"], 15.5, WHITE, True,
         align=PP_ALIGN.LEFT, margin=12, name=f"{data['case']} label text")
    box(slide, 16, left_y + 37, 126, 36, TEAL_PALE, LINE_TEAL, 0.65, True,
        f"{data['case']} status")
    text(slide, 27, left_y + 39, 110, 31, data["status"], 10.2, INK, True,
         align=PP_ALIGN.LEFT, name=f"{data['case']} status text")

    care_y = y + 89
    hospital_icon(slide, 25, care_y)
    text(slide, 69, care_y + 4, 92, 37, f"Care setting:\n{data['care']}", 8.8,
         BODY, align=PP_ALIGN.LEFT, name=f"{data['case']} care setting", spacing=0.9)

    report_y = y + (154 if h >= 270 else 137)
    report_h = y + h - report_y - 12
    box(slide, 16, report_y, 146, report_h, BLUE_PALE, LINE_BLUE, 0.7, True,
        f"{data['case']} current finding card")
    check_icon(slide, 26, report_y + 20, 1, BLUE, f"{data['case']} current finding")
    text(slide, 56, report_y + 10, 97, report_h - 18, data["current"], 9.5,
         INK, True, align=PP_ALIGN.LEFT, name=f"{data['case']} current finding text",
         spacing=0.91, font=FONT_NARROW)


def right_findings(slide, data, y, h):
    top = y + data.get("right_offset", 0)
    outer_h = h - data.get("right_offset", 0) - 12
    box(slide, 1320, top, 157, outer_h, GOLD_PALE, LINE_GOLD, 0.7, True,
        f"{data['case']} follow-up card")
    chip_h = 72 if h >= 270 else 65
    chip_y = top + outer_h - chip_h - 5
    add_bullets(slide, 1329, top + 13, 139, chip_y - top - 20, data["findings"],
                9.8 if data["case"] != "Case C" else 9.2)
    box(slide, 1325, chip_y, 147, chip_h, "FCF6EA", LINE_GOLD, 0.7, True,
        f"{data['case']} follow-up note")
    if data["note_kind"] == "stable":
        check_icon(slide, 1336, chip_y + 18, 1.05, BODY, "Stable follow-up")
    else:
        device_icon(slide, 1335, chip_y + 14, BODY)
    text(slide, 1371, chip_y + 8, 93, chip_h - 16, data["note"], 9.5, BODY,
         True, align=PP_ALIGN.LEFT, name=f"{data['case']} follow-up note text",
         spacing=0.9, font=FONT_NARROW)


def add_case(slide, data):
    y, h = data["y"], data["h"]
    line(slide, 521, y + 9, 1000, y + 9, BLUE, 1.5, arrow=True,
         name=f"{data['case']} trajectory arrow")
    add_cropped_picture(slide, ASSETS / data["current_file"], 179, y, 337, h,
                        focus_y=data["current_focus"], name=f"{data['case']} current radiograph")
    add_cropped_picture(slide, ASSETS / data["target_file"], 1008, y, 306, h,
                        focus_y=data["target_focus"], name=f"{data['case']} observed follow-up")
    left_case(slide, data, y, h)
    right_findings(slide, data, y, h)
    metadata(slide, data, y + 22)

    card_y = y + 102
    context_card(slide, data, 536, card_y, 450, data["context_h"])
    audit_y = card_y + data["context_h"] + 4
    audit_h = 30 if data["case"] != "Case C" else 25
    box(slide, 536, audit_y, 450, audit_h, "FBF7EE", LINE_GOLD, 0.7, True,
        f"{data['case']} audit note")
    eye_icon(slide, 546, audit_y + (8 if audit_h == 30 else 6), BODY)
    text(slide, 572, audit_y + 3, 400, audit_h - 6,
         "These records are retrospective audit context for the interval, not causal.",
         7.7, BODY, align=PP_ALIGN.LEFT, name=f"{data['case']} audit note text")


CASES = [
    {
        "case": "Case A", "status": "clear change", "care": "SICU",
        "y": 166, "h": 277, "label_offset": 8, "right_offset": 10,
        "horizon": "0–24 h", "elapsed": "20.1 h", "context_h": 153,
        "current_file": "case_a_current_cxr.jpg", "target_file": "case_a_observed_followup_cxr.jpg",
        "current_focus": 0.421, "target_focus": 0.509,
        "current": "no edema /\neffusion /\npneumonia\nreported",
        "findings": ["marked bilateral\nopacity / edema", "small right pleural\neffusion",
                     "infectious-appearing\ndiffuse process"],
        "note_kind": "device", "note": "Support-device\ndifference\npresent",
        "context": [
            (monitor_icon, "Ventilation", "mechanical\nventilation"),
            (imaging_icon, "Imaging", "prior CXRs,\nCT chest"),
            (iv_icon, "Lines", "central line,\narterial line"),
            (vial_icon, "IV medications", "diuretics,\nantibiotics"),
        ],
    },
    {
        "case": "Case B", "status": "stable follow-up", "care": "SICU",
        "y": 468, "h": 253, "right_offset": 7,
        "horizon": "0–24 h", "elapsed": "≈24.0 h", "context_h": 111,
        "current_file": "case_b_current_cxr.jpg", "target_file": "case_b_observed_followup_cxr.jpg",
        "current_focus": 0.429, "target_focus": 0.5,
        "current": "bibasal\nopacification\npleural\neffusions\ncompressive\natelectasis",
        "findings": ["bibasal\nopacification", "pleural\neffusions", "compressive\natelectasis"],
        "note_kind": "stable", "note": "stable\nfollow-up",
        "context": [
            (iv_icon, "Lines", "central line,\nPICC"),
            (vial_icon, "IV medications", "vasopressors,\nantimicrobials"),
        ],
    },
    {
        "case": "Case C", "status": "clear change", "care": "trauma SICU",
        "y": 741, "h": 255,
        "horizon": "24–72 h", "elapsed": "50.5 h", "context_h": 137,
        "current_file": "case_c_current_cxr.jpg", "target_file": "case_c_observed_followup_cxr.jpg",
        "current_focus": 0.400, "target_focus": 0.382,
        "current": "no focal opacity\nor pleural\neffusion",
        "findings": ["new bilateral\nparenchymal\nopacities", "right upper lung\nmost severe",
                     "probable\nmultifocal\npneumonia", "possible pleural\neffusion"],
        "note_kind": "device", "note": "Support-device\ndifference\npresent",
        "context": [
            (monitor_icon, "Ventilation", "mechanical\nventilation"),
            (imaging_icon, "Imaging", "serial CXRs,\nCT scans"),
            (culture_icon, "Cultures", "blood, sputum,\nrespiratory"),
            (iv_icon, "Lines", "central line,\nPICC"),
            (vial_icon, "IV medications", "broad-spectrum\nantibiotics"),
        ],
    },
]


def build():
    needed = [ASSETS / data[key] for data in CASES for key in ("current_file", "target_file")]
    missing = [str(path) for path in needed if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing radiograph assets:\n" + "\n".join(missing))

    presentation = Presentation()
    presentation.slide_width = u(W)
    presentation.slide_height = u(H)
    presentation.core_properties.title = "Representative longitudinal MIMIC trajectories"
    presentation.core_properties.subject = "Editable appendix figure for MedWorld-JEPA"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    text(slide, 185, 6, 1220, 51,
         "Representative longitudinal MIMIC trajectories used by MedWorld-JEPA",
         30, INK, True, name="Figure title", font=FONT_NARROW)
    text(slide, 385, 57, 770, 49,
         "Current chest radiograph evidence is paired with one observed future study;\n"
         "linked MIMIC-IV records provide retrospective audit context for the interval.",
         15.5, MUTED, name="Figure subtitle", spacing=0.95)

    header(slide, 179, 116, 337, 1, "Current evidence at t", BLUE_HEADER)
    header(slide, 553, 116, 417, 2, "Linked retrospective context — not model input", TEAL_HEADER)
    header(slide, 1008, 116, 460, 3, "Observed follow-up target t+", GOLD_HEADER)

    for data in CASES:
        add_case(slide, data)

    line(slide, 17, 454, 1475, 454, "B8C9CC", 0.8, dashed=True, name="Case A separator")
    line(slide, 17, 729, 1475, 729, "B8C9CC", 0.8, dashed=True, name="Case B separator")

    box(slide, 65, 1008, 1374, 39, "FBF7EC", LINE_GOLD, 0.7, True, "Scientific boundary footer")
    balance_icon(slide, 80, 1014)
    text(slide, 123, 1013, 1290, 28,
         "Each follow-up is one observed future, not the only possible future. Report-derived findings are weak evaluation labels. Post-current clinical events are retrospective and non-causal.",
         9.8, BODY, align=PP_ALIGN.LEFT, name="Scientific boundary footer text")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
