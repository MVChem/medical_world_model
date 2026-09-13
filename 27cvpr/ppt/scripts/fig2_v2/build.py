"""Build and render the compact, editable, single-column Figure 2."""

import argparse
from pathlib import Path
import shutil
import subprocess
import tempfile

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt


ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / "ppt" / "fig2_v2.pptx"
PAPER = ROOT.parent / "imgs" / "fig2_v2.pdf"
WIDTH, HEIGHT = 236.25, 98.0  # The manuscript's 3.28125-inch column width.
INK = "101010"
FONT = "Comic Sans MS"


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(WIDTH), Pt(HEIGHT)
    prs.core_properties.title = "Figure 2 — Eight-slot state construction"
    prs.core_properties.subject = "Compact single-column method detail"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    def no_effects(shape):
        shape.shadow.inherit = False
        for effect in shape._element.xpath("./p:style/a:effectRef"):
            effect.set("idx", "0")

    def box(name, x, y, w, h, fill, stroke, radius=.12, line_width=.55):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h)
        )
        shape.name = name
        shape.adjustments[0] = radius
        no_effects(shape)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(line_width)
        return shape

    def label(name, value, x, y, w, h, size=7.5, bold=False,
              font=FONT, align=PP_ALIGN.CENTER, italic=False):
        shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
        shape.name = name
        tf = shape.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(value.split("\n")):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.alignment = align
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = line
            run.font.name = font
            run.font.size = Pt(size)
            run.font.bold, run.font.italic = bold, italic
            run.font.color.rgb = RGBColor.from_string(INK)
        return shape

    def route(name, points, arrow=True):
        if len(points) == 2:
            (x1, y1), (x2, y2) = points
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, Pt(x1), Pt(y1), Pt(x2), Pt(y2)
            )
        else:
            path = slide.shapes.build_freeform(*points[0], scale=Pt(1))
            path.add_line_segments(points[1:], close=False)
            shape = path.convert_to_shape()
            shape.fill.background()
        shape.name = name
        no_effects(shape)
        shape.line.color.rgb = RGBColor.from_string(INK)
        shape.line.width = Pt(.65)
        if arrow:
            end = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "med", "len": "med"}.items():
                end.set(key, value)
            shape._element.spPr.get_or_add_ln().append(end)
        return shape

    # Both sources remain inside the VLM. Four colored blocks select depths;
    # the short gray blocks only suggest omitted layers, not exact layer counts.
    for branch, panel_y, token_y, panel_fill, panel_stroke, fill, stroke in [
        ("vision-language", 18, 34, "FCEFFB", "8F329B", "EFA9D2", "CD45A7"),
        ("vision encoder", 52, 68, "E6F3FF", "0872CA", "9FD5F5", "05A6E1"),
    ]:
        box(f"Panel: {branch}", 2, panel_y, 62, 29, panel_fill, panel_stroke)
        label(f"Label: {branch}", "Vision-language" if branch == "vision-language"
              else "Vision encoder", 5, panel_y + 1.4, 56, 11, size=7.3)
        blocks = slide.shapes.add_group_shape()
        blocks.name = f"Depth samples: {branch}"
        cursor = 8.9
        for i in range(7):
            selected = i % 2 == 0
            width = 5.3 if selected else 3.0
            token = box(f"Layer: {branch} {'selected' if selected else 'omitted'} {i + 1}",
                        cursor, token_y, width, 8.5,
                        fill if selected else "D1D1D1",
                        stroke if selected else "A3A3A3", line_width=.4)
            blocks.shapes._spTree.insert_element_before(token._element, "p:extLst")
            cursor += width + 3.0
        blocks.shapes._recalculate_extents()

    label("Label: VLM", "VLM", 18, 2, 30, 13, size=9.5, bold=True)
    box("Module: lightweight adapter", 77, 39, 43, 33, "FFFADB", "99882C")
    label("Label: lightweight adapter", "Lightweight\nadapter", 78, 44, 41, 23,
          size=7.2)
    box("Module: task decoder", 178, 39, 35, 33, "DDEFD5", "287F31")
    label("Label: task decoder", "Task\ndecoder", 179, 44, 33, 23, size=7.8)

    for branch, y, fill, stroke in [
        ("vision-language", 42, "EFA9D2", "CD45A7"),
        ("vision encoder", 59, "9FD5F5", "05A6E1"),
    ]:
        row = slide.shapes.add_group_shape()
        row.name = f"Slot row: {branch}"
        for i in range(4):
            slot = box(f"Slot: {branch} {i + 1}", 132 + 8 * i, y, 5.5, 10,
                       fill, stroke, line_width=.5)
            row.shapes._spTree.insert_element_before(slot._element, "p:extLst")
        row.shapes._recalculate_extents()

    route("Arrow: fusion layers to adapter", [(64, 37), (70, 37), (70, 47), (76, 47)])
    route("Arrow: visual layers to adapter", [(64, 71), (70, 71), (70, 64), (76, 64)])
    route("Arrow: adapter to fusion slots", [(120, 47), (130.5, 47)])
    route("Arrow: adapter to visual slots", [(120, 64), (130.5, 64)])
    route("Bracket: eight-slot state", [(163, 47), (170, 47), (170, 64), (163, 64)],
          arrow=False)
    route("Arrow: state to decoder", [(170, 55.5), (176.5, 55.5)])
    route("Arrow: decoder to loss", [(213, 55.5), (219, 55.5)])

    label("Label: eight slots", "8 slots", 129, 27, 36, 12, size=8)
    label("Equation: state dimensions", "8 × d", 132, 73, 30, 12,
          size=8.5, font="Times New Roman", italic=True)
    label("Label: joint optimization", "Joint optimization", 72, 86, 93, 11, size=7.5)
    loss = label("Equation: task loss", "L", 220, 47, 15, 15,
                 size=10, font="Times New Roman", italic=True, align=PP_ALIGN.LEFT)
    run = loss.text_frame.paragraphs[0].add_run()
    run.text = "task"
    run.font.name = "Times New Roman"
    run.font.size = Pt(8.5)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(INK)
    run._r.get_or_add_rPr().set("baseline", "-23000")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


def render(install=False):
    """Render the saved PPTX, including any manual edits, using LibreOffice."""
    with tempfile.TemporaryDirectory(prefix="fig2_v2_render_") as directory:
        work = Path(directory)
        result = subprocess.run([
            "libreoffice", f"-env:UserInstallation={(work / 'profile').as_uri()}",
            "--headless", "--convert-to", "pdf:impress_pdf_Export",
            "--outdir", str(work), str(OUTPUT),
        ], check=True, capture_output=True, text=True, timeout=120)
        rendered = work / OUTPUT.with_suffix(".pdf").name
        if not rendered.is_file():
            raise RuntimeError(result.stdout + result.stderr)
        shutil.copy2(rendered, OUTPUT.with_suffix(".pdf"))
    with fitz.open(OUTPUT.with_suffix(".pdf")) as pdf:
        assert len(pdf) == 1
        page = pdf[0]
        page.get_pixmap(matrix=fitz.Matrix(6, 6)).save(OUTPUT.with_suffix(".png"))
        # The slide already has the exact publication width; keep its canvas.
        for block in page.get_text("blocks"):
            rect = fitz.Rect(block[:4])
            if not (page.rect + (-.5, -.5, .5, .5)).contains(rect):
                raise RuntimeError(f"Text exceeds slide bounds: {block[4]!r}")
    if install:
        PAPER.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(OUTPUT.with_suffix(".pdf"), PAPER)
        print(PAPER)
    print(OUTPUT.with_suffix(".png"))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render-only", action="store_true",
                        help="Export the existing PPTX, preserving manual edits")
    parser.add_argument("--install", action="store_true", help="Update the manuscript PDF asset")
    args = parser.parse_args()
    if not args.render_only:
        build()
    render(install=args.install)
