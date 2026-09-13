"""Rebuild the supplied three-panel trajectory figure as native PPT shapes.

Coordinates below refer to a 2048 x (2048/3) design canvas. The source has
the same 3:1 aspect ratio. No bitmap or SVG is embedded in the slide.
"""
from __future__ import annotations

import argparse
import math
from pathlib import Path
import shutil
import subprocess
import tempfile
from zipfile import ZipFile

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
NAME = HERE.name
OUT = ROOT / "ppt" / f"{NAME}.pptx"
W, H, PPI = 2048, 2048 / 3, 120
INK = "1C2941"


def u(value):
    return int(round(value * 914400 / PPI))


def xml(tag, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(key, str(value))
    return element


def rgba(parent, color, opacity=1):
    node = xml("a:srgbClr", val=color)
    if opacity < 1:
        node.append(xml("a:alpha", val=round(opacity * 100000)))
    parent.append(node)
    return node


def style(shape, fill=None, line=None, width=1, opacity=1):
    # Avoid theme-dependent shadow, fill or line effects.
    shape.shadow.inherit = False
    # LibreOffice still resolves effectRef even when effectLst is empty.
    # Set its index to zero as well to suppress the template's default shadow.
    for ref in shape._element.xpath("./p:style/a:effectRef"):
        ref.set("idx", "0")
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        if opacity < 1:
            shape.fill._xPr.solidFill[0].append(
                xml("a:alpha", val=round(opacity * 100000)))
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = RGBColor.from_string(line)
        shape.line.width = Pt(width * 72 / PPI)
        shape.line._get_or_add_ln().set("cap", "rnd")
        shape.line._get_or_add_ln().append(xml("a:round"))
    else:
        shape.line.fill.background()
    return shape


def gradient(shape, stops, angle=45):
    """Native DrawingML linear fill; each stop is (position, RGB, alpha)."""
    shape.fill.gradient()
    node = shape.fill._xPr.gradFill
    for child in list(node):
        node.remove(child)
    node.set("rotWithShape", "1")
    listing = xml("a:gsLst")
    for position, color, opacity in stops:
        stop = xml("a:gs", pos=round(position * 100000))
        rgba(stop, color, opacity)
        listing.append(stop)
    node.append(listing)
    node.append(xml("a:lin", ang=round(angle * 60000), scaled="1"))
    return shape


def group(parent, name):
    result = parent.add_group_shape()
    result.name = name
    return result.shapes


def oval(parent, name, x, y, w, h, fill=None, line=None, width=1, opacity=1):
    result = parent.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(w), u(h))
    result.name = name
    return style(result, fill, line, width, opacity)


def path(parent, name, commands, fill=None, line=None, width=1, opacity=1):
    """Editable freeform with exact cubic Beziers instead of sampled points."""
    coords = [xy for cmd in commands for xy in cmd[1:]]
    xs, ys = coords[::2], coords[1::2]
    x0, y0, x1, y1 = min(xs), min(ys), max(xs), max(ys)
    dx, dy = max(x1 - x0, 0.1), max(y1 - y0, 0.1)
    result = parent.add_shape(MSO_SHAPE.RECTANGLE, u(x0), u(y0), u(dx), u(dy))
    result.name = name
    props = result._element.spPr
    props.remove(props.prstGeom)
    geometry = xml("a:custGeom")
    for tag in ["a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"]:
        geometry.append(xml(tag))
    geometry.append(xml("a:rect", l="0", t="0", r="r", b="b"))
    listing = xml("a:pathLst")
    drawing = xml("a:path", w=u(dx), h=u(dy))
    if fill is None:
        drawing.set("fill", "none")
    for command in commands:
        kind, points = command[0], command[1:]
        segment = xml({"M": "a:moveTo", "L": "a:lnTo",
                       "C": "a:cubicBezTo", "Z": "a:close"}[kind])
        for i in range(0, len(points), 2):
            segment.append(xml("a:pt", x=u(points[i] - x0), y=u(points[i + 1] - y0)))
        drawing.append(segment)
    listing.append(drawing)
    geometry.append(listing)
    props.insert(1, geometry)
    return style(result, fill, line, width, opacity)


def polygon(parent, name, points, **kwargs):
    return path(parent, name, [("M", *points[0]),
                               *[("L", *p) for p in points[1:]], ("Z",)], **kwargs)


def rounded_quad(parent, name, points, radius=0.065, **kwargs):
    """Round all four corners of a perspective card with cubic segments."""
    commands = []
    for i, point in enumerate(points):
        previous, following = points[i - 1], points[(i + 1) % len(points)]
        before = tuple(p + (v - p) * radius for p, v in zip(point, previous))
        after = tuple(p + (v - p) * radius for p, v in zip(point, following))
        commands.append(("M" if i == 0 else "L", *before))
        commands.append(("C", *point, *point, *after))
    commands.append(("Z",))
    return path(parent, name, commands, **kwargs)


def caption(parent, name, value, cx, y, width, size=44):
    shape = parent.add_textbox(u(cx - width / 2), u(y), u(width), u(58))
    shape.name = name
    frame = shape.text_frame
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = 1
    run = paragraph.add_run()
    run.text = value
    run.font.name = "Arial"
    run.font.size = Pt(size * 72 / PPI)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(INK)
    return shape


def soft_ellipse(parent, name, cx, cy, rx, ry, color, opacity=.008, layers=16):
    # A faint halo made entirely from editable translucent concentric ovals.
    # It intentionally approximates the reference's raster blur without pixels.
    shapes = group(parent, name)
    for i in range(layers):
        ratio = 1 - i / layers * .58
        oval(shapes, f"{name}: halo {i + 1:02}", cx-rx*ratio, cy-ry*ratio,
             rx*2*ratio, ry*2*ratio, color, opacity=opacity)


def landscape(parent, name, corners, blue, opacity=1):
    """Map the simple mountains and sun into a perspective quadrilateral."""
    def project(x, y):
        tl, tr, br, bl = corners
        return tuple((1-y)*((1-x)*a+x*b)+y*((1-x)*d+x*c)
                     for a, b, c, d in zip(tl, tr, br, bl))
    mountains = [(.075, .81), (.30, .385), (.34, .368), (.38, .40),
                 (.555, .68), (.635, .535), (.675, .52), (.708, .55), (.91, .80)]
    result = polygon(parent, name + ": mountains", [project(x, y) for x, y in mountains],
                     fill=blue, opacity=opacity)
    gradient(result, [(0, "89BDFA", opacity), (1, blue, opacity)], 55)
    center = project(.73, .285)
    radius = 10.5
    result = oval(parent, name + ": sun", center[0]-radius, center[1]-radius,
                  radius*2, radius*2.2, blue, opacity=opacity)
    gradient(result, [(0, "9DCBFE", opacity), (1, blue, opacity)], 65)


def arrow(parent, name, start, end, shaft=15, head=43, head_length=32,
          color="FA7584", tail_opacity=.25):
    dx, dy = end[0]-start[0], end[1]-start[1]
    length = math.hypot(dx, dy)
    tx, ty = dx/length, dy/length
    nx, ny = -ty, tx
    def xy(along, across):
        return (start[0]+along*tx+across*nx, start[1]+along*ty+across*ny)
    points = [xy(0, -shaft/2), xy(length-head_length, -shaft/2),
              xy(length-head_length, -head/2), xy(length, 0),
              xy(length-head_length, head/2), xy(length-head_length, shaft/2),
              xy(0, shaft/2)]
    result = polygon(parent, name, points, fill=color)
    gradient(result, [(0, color, tail_opacity), (.5, color, .8), (1, color, 1)],
             math.degrees(math.atan2(dy, dx)) % 360)
    return result


def panel_one(slide):
    panel = group(slide.shapes, "01 — Generative Trajectory Representation")
    soft_ellipse(panel, "01: pale blue light", 355, 270, 319, 190, "7EB5FF", .005)
    # Seven translucent image planes, ordered from the most distant to nearest.
    cards = [
        ([(457, 91), (620, 120), (620, 259), (457, 220)], "D8E8FD", .69),
        ([(393, 114), (564, 143), (563, 282), (393, 242)], "C4DDFB", .67),
        ([(333, 137), (490, 163), (489, 301), (333, 263)], "B1D4FB", .65),
        ([(271, 158), (418, 185), (417, 340), (271, 300)], "9BC6F9", .65),
        ([(211, 181), (360, 209), (358, 365), (211, 324)], "85B9F7", .63),
        ([(150, 203), (302, 232), (301, 392), (150, 350)], "72AAF1", .65),
        ([(90, 226), (241, 258), (241, 419), (90, 376)], "66A6EF", .78),
    ]
    for i, (points, color, opacity) in enumerate(cards):
        card = group(panel, f"01: frame {7-i:02} (rear to front)")
        shape = rounded_quad(card, "Translucent image plane", points,
                             fill=color, line="FFFFFF" if i < 3 else None, width=1)
        gradient(shape, [(0, color, opacity), (1, "C9E3FF", opacity*.7)], 65)
        if i == 6:
            frame = rounded_quad(card, "Front image border", points,
                                 fill=None, line="659EEE", width=6)
            front = [(98, 234), (234, 265), (234, 406), (98, 371)]
            # The front plane is the only fully visible landscape thumbnail.
            mountain = polygon(card, "Front image mountains", [
                (103, 351), (135, 300), (140, 296), (145, 298),
                (175, 348), (183, 331), (189, 326), (194, 330),
                (216, 371)], fill="5195E8")
            gradient(mountain, [(0, "4B91E7", .7), (1, "81B6F3", .75)], 90)
            oval(card, "Front image sun", 189, 277, 24, 26, "D1E8FF", opacity=.94)
            # Add a pale lower edge, following the glass plane perspective.
            path(card, "Front frame lower highlight", [("M", 98, 376),
                 ("L", 231, 412)], line="C8E4FF", width=2)

    trajectory = group(panel, "01: forward generative trajectory")
    curve = [("M", 142, 380), ("C", 201, 363, 282, 330, 352, 291),
             ("C", 423, 253, 504, 218, 570, 180)]
    path(trajectory, "Trajectory white rim", curve, line="FFFFFF", width=17)
    path(trajectory, "Trajectory blue Bezier", curve, line="2582FA", width=11)
    arrow(trajectory, "Trajectory arrow head", (554, 189), (594, 168),
          shaft=11, head=42, head_length=34, color="267FFC", tail_opacity=1)
    for i, (cx, cy) in enumerate([(239, 345), (352, 291), (456, 235)]):
        oval(trajectory, f"Time point {i+1}: white rim", cx-16, cy-16, 32, 32, "FFFFFF")
        oval(trajectory, f"Time point {i+1}: blue ring", cx-12, cy-12, 24, 24, "2582FA")
        oval(trajectory, f"Time point {i+1}: center", cx-4.3, cy-4.3, 8.6, 8.6, "FFFFFF")

    speed = group(panel, "01: blue velocity mark")
    for i, (x0, yy, x1) in enumerate([(409, 370, 488), (396, 388, 482), (382, 406, 483)]):
        shape = speed.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, u(x0), u(yy), u(x1-x0), u(9))
        shape.name = f"Velocity streak {i+1}"
        shape.adjustments[0] = .5
        style(shape, "408FFD")
        gradient(shape, [(0, "81BBFD", 0), (.9, "348AFF", .92), (1, "348AFF", .94)], 0)
    glyph = path(speed, "Velocity v: editable outline", [
        ("M", 501, 366), ("C", 506, 363, 511, 362, 517, 362),
        ("L", 522, 402), ("C", 532, 389, 538, 375, 539, 363),
        ("L", 552, 363), ("C", 548, 381, 537, 402, 523, 416),
        ("L", 511, 416), ("L", 506, 373), ("L", 501, 374), ("Z",)], fill="2A82FA")
    gradient(glyph, [(0, "3B90FF", 1), (1, "257DF7", 1)], 70)
    caption(panel, "01: title line 1", "1)  Generative Trajectory", 355, 501, 660)
    caption(panel, "01: title line 2", "Representation", 374, 556, 650)


def panel_two(slide):
    panel = group(slide.shapes, "02 — Trajectory Consistency Evolution")
    soft_ellipse(panel, "02: pale violet light", 1024, 275, 305, 200, "9DA8FA", .006)
    arcs = group(panel, "02: cycle arrows")
    upper = path(arcs, "Upper cycle ribbon: left to right", [
        ("M", 849, 172), ("C", 940, 103, 1090, 96, 1182, 152),
        ("L", 1187, 145), ("C", 1189, 143, 1191, 146, 1192, 149),
        ("L", 1199, 178), ("C", 1200, 181, 1198, 183, 1195, 182),
        ("L", 1162, 176), ("C", 1158, 175, 1156, 173, 1159, 170),
        ("L", 1168, 162), ("C", 1079, 111, 946, 121, 857, 181),
        ("C", 850, 187, 843, 180, 849, 172), ("Z",)], fill="7476EA")
    gradient(upper, [(0, "8AC4FA", 1), (.28, "438FFC", 1), (1, "8D66DF", 1)], 0)
    lower = path(arcs, "Lower cycle ribbon: right to left", [
        ("M", 1184, 365), ("C", 1192, 359, 1199, 368, 1193, 375),
        ("C", 1110, 451, 957, 452, 870, 394), ("L", 864, 402),
        ("C", 862, 405, 859, 403, 858, 400), ("L", 848, 371),
        ("C", 847, 367, 849, 364, 853, 365), ("L", 885, 371),
        ("C", 889, 372, 890, 375, 887, 377), ("L", 880, 383),
        ("C", 965, 438, 1109, 429, 1184, 365), ("Z",)], fill="7979E8")
    gradient(lower, [(0, "8C67E0", 1), (.75, "4E90F9", 1), (1, "A5CBF9", 1)], 0)

    cards = [
        ("left", [(754, 191), (888, 212), (888, 334), (754, 351)], "5697F3", "66A1F5"),
        ("middle", [(954, 208), (1088, 213), (1088, 334), (954, 338)], "7B71E9", "7675EA"),
        ("right", [(1157, 212), (1293, 191), (1293, 351), (1157, 334)], "757AEB", "6195EF"),
    ]
    for name, corners, border, mountain in cards:
        card = group(panel, f"02: {name} image card")
        frame = rounded_quad(card, "Image frame", corners, radius=.035,
                             fill="D5E6FE", line=border, width=7)
        gradient(frame, [(0, "DCEBFF", 1), (1, "BCD7FE", 1)], 70)
        landscape(card, f"{name} landscape", corners, mountain, .95)
    for index, xx in enumerate([899, 1101]):
        shape = panel.add_shape(MSO_SHAPE.LEFT_RIGHT_ARROW, u(xx), u(262), u(44), u(25))
        shape.name = f"02: adjacent consistency link {index+1}"
        shape.adjustments[0], shape.adjustments[1] = .42, .52
        style(shape, "9389E9")
        gradient(shape, [(0, "A0A6F1", 1), (1, "9184E8", 1)], 0)
    caption(panel, "02: title line 1", "2)  Trajectory Consistency", 1024, 501, 680)
    caption(panel, "02: title line 2", "Evolution", 1028, 560, 600)


def panel_three(slide):
    panel = group(slide.shapes, "03 — Real-Centered Trajectory Optimization")
    cx, cy = 1694, 280
    soft_ellipse(panel, "03: pale target light", cx, cy, 174, 174, "9CBFF0", .005)
    outwards = group(panel, "03: eight outward coral arrows")
    for i in range(8):
        theta = i*math.pi/4
        # Diagonal tips follow the original's almost square radial envelope.
        tip = {0: 235, 2: 205, 4: 235, 6: 216}.get(i, 239)
        start = (cx+129*math.cos(theta), cy+129*math.sin(theta))
        end = (cx+tip*math.cos(theta), cy+tip*math.sin(theta))
        arrow(outwards, f"Outward arrow {i+1:02}", start, end, shaft=17,
              head=44, head_length=34, color="FC7788", tail_opacity=.12)
    target = group(panel, "03: real-centered concentric target")
    for name, radius, stops in [
        ("Outer target disk", 148, [(0, "F1F7FD", 1), (1, "E8F1FC", 1)]),
        ("Middle target disk", 106, [(0, "D8E9FD", 1), (1, "CCE1FC", 1)]),
        ("Inner target disk", 67, [(0, "C6DFFE", 1), (1, "B9D6FE", 1)]),
        ("Blue center", 37, [(0, "4191FC", 1), (1, "2578F7", 1)]),
    ]:
        shape = oval(target, name, cx-radius, cy-radius, radius*2, radius*2, "D5E6FF")
        gradient(shape, stops, 65)
    inward = group(panel, "03: four inward green arrows")
    for i, theta in enumerate([math.pi/4, 3*math.pi/4, 5*math.pi/4, 7*math.pi/4]):
        start = (cx+134*math.cos(theta), cy+134*math.sin(theta))
        end = (cx+68*math.cos(theta), cy+68*math.sin(theta))
        arrow(inward, f"Inward arrow {i+1:02}", start, end, shaft=17,
              head=42, head_length=30, color="2FB69F", tail_opacity=.15)
    caption(panel, "03: title line 1", "3)  Real-Centered Trajectory", 1703, 501, 654, size=42)
    caption(panel, "03: title line 2", "Optimization", 1710, 560, 650)


def build():
    prs = Presentation()
    prs.slide_width, prs.slide_height = u(W), u(H)
    prs.core_properties.title = "Trajectory Framework — Editable Vector Reconstruction"
    prs.core_properties.subject = "Generative trajectory, consistency evolution, real-centered optimization"
    prs.core_properties.author = ""
    prs.core_properties.keywords = "native editable shapes; vector; reference reconstruction"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    panel_one(slide)
    panel_two(slide)
    panel_three(slide)
    slide.notes_slide.notes_text_frame.text = (
        "Reconstructed from assets/reference.png. All visible slide content is native "
        "PowerPoint geometry or editable text; no bitmap or SVG objects are embedded. "
        "Three named top-level groups correspond to the three panels. Ungroup or use "
        "the Selection Pane to edit parts. Glass transparency and light are vector "
        "approximations of the reference. Arial Bold matches the reference visually. "
        "The transparent reference is shown on a white slide background.")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    validate()


def validate():
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(OUT) as archive:
        media = [name for name in archive.namelist() if name.startswith("ppt/media/")]
        root = etree.fromstring(archive.read("ppt/slides/slide1.xml"))
        assert not media, f"Unexpected embedded media: {media}"
        assert not root.findall(".//p:pic", ns)
        assert len(root.findall("./p:cSld/p:spTree/p:grpSp", ns)) == 3
        labels = root.xpath(".//a:t/text()", namespaces=ns)
        assert labels == ["1)  Generative Trajectory", "Representation",
                          "2)  Trajectory Consistency", "Evolution",
                          "3)  Real-Centered Trajectory", "Optimization"], labels
        native = len(root.findall(".//p:sp", ns))
        freeforms = len(root.findall(".//a:custGeom", ns))
        print(f"Validated: {native} native shapes/text boxes, {freeforms} freeforms, 0 embedded images.")


def render():
    import fitz
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required to render a preview from the PPTX.")
    with tempfile.TemporaryDirectory(prefix="trajectory_framework_render_") as directory:
        work = Path(directory)
        result = subprocess.run([
            executable, f"-env:UserInstallation={(work / 'profile').as_uri()}",
            "--headless", "--convert-to", "pdf:impress_pdf_Export", "--outdir", str(work), str(OUT)
        ], capture_output=True, text=True, timeout=120)
        pdf = work / OUT.with_suffix(".pdf").name
        if result.returncode or not pdf.exists():
            raise RuntimeError(f"PPTX rendering failed: {result.stdout}\n{result.stderr}")
        with fitz.open(pdf) as document:
            assert len(document) == 1
            page = document[0]
            assert abs(page.rect.width/page.rect.height-3) < .001
            extracted = " ".join(page.get_text().split())
            for label in ["Generative Trajectory", "Representation", "Trajectory Consistency",
                          "Evolution", "Real-Centered Trajectory", "Optimization"]:
                assert label in extracted, f"Missing rendered text: {label}"
            # Account for PDF page-size rounding so raster bounds do not ceil
            # to an unwanted extra pixel. The aspect correction is negligible.
            matrix = fitz.Matrix((3072-.001)/page.rect.width, (1024-.001)/page.rect.height)
            preview = page.get_pixmap(matrix=matrix, alpha=False)
            assert (preview.width, preview.height) == (3072, 1024)
            preview.save(OUT.with_suffix(".png"))
            print(f"Rendered {OUT.with_suffix('.png')}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--overwrite", action="store_true", help="Rebuild an existing PPTX from this script.")
    parser.add_argument("--render-only", action="store_true", help="Render the existing PPTX, preserving manual edits.")
    parser.add_argument("--no-render", action="store_true", help="Build only the PPTX without a preview.")
    args = parser.parse_args()
    if args.render_only:
        render()
    else:
        if OUT.exists() and not args.overwrite:
            parser.error(f"{OUT} exists; use --overwrite to rebuild or --render-only to preserve edits.")
        build()
        if not args.no_render:
            render()
    print(OUT)
