from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "source/fig1_v2.png"
OUTPUT = ROOT / "ppt/fig1_v2.pptx"

SW, SH = 13.333333, 7.5
PX_W, PX_H = 1672, 941
XPI, YPI = PX_W / SW, PX_H / SH

WHITE, INK, GRAY = "FFFFFF", "111111", "AEB6C2"
BLUE, BLUE_FILL, BLUE_TOKEN = "1764C0", "EEF5FC", "C8DDF3"
PURPLE, PURPLE_FILL, PURPLE_TOKEN = "7030A0", "F5F0FA", "E2D4EC"
ORANGE, ORANGE_FILL = "E55700", "FFF4EC"
GREEN, GREEN_FILL = "3B7D23", "F2F8EF"
FONT = "Arial"


def ex(px):
    return Inches(px / XPI)


def ey(px):
    return Inches(px / YPI)


def rgb(value):
    return RGBColor.from_string(value)


def clean(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def end_arrow(shape):
    line = shape._element.spPr.get_or_add_ln()
    node = OxmlElement("a:tailEnd")
    node.set("type", "triangle")
    node.set("w", "sm")
    node.set("len", "sm")
    line.append(node)


def dash(shape, value="dash"):
    line = shape._element.spPr.get_or_add_ln()
    node = OxmlElement("a:prstDash")
    node.set("val", value)
    line.append(node)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = rgb(WHITE)


def box(x, y, w, h, fill=WHITE, line=GRAY, width=1.0, radius=.07, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, ex(x), ey(y), ex(w), ey(h))
    if name:
        shape.name = name
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return clean(shape)


def rect(x, y, w, h, fill=WHITE, line=GRAY, width=1.0, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ex(x), ey(y), ex(w), ey(h))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return clean(shape)


def oval(x, y, w, h, fill=WHITE, line=GRAY, width=1.0, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, ex(x), ey(y), ex(w), ey(h))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return clean(shape)


def text(x, y, w, h, value, size=10, color=INK, bold=False, italic=False,
         align=PP_ALIGN.CENTER, font=FONT, name=None):
    shape = slide.shapes.add_textbox(ex(x), ey(y), ex(w), ey(h))
    if name:
        shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = ex(1)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 0.92
    run = p.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return clean(shape)


def connector(x1, y1, x2, y2, color=INK, width=1.2, arrow=True, dashed=False, name=None):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, ex(x1), ey(y1), ex(x2), ey(y2))
    if name:
        shape.name = name
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        end_arrow(shape)
    if dashed:
        dash(shape)
    return clean(shape)


def path(points, color=INK, width=1.2, arrow=True, name="connector"):
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        connector(x1, y1, x2, y2, color, width, arrow and i == len(points) - 2,
                  name=f"{name} {i + 1}")


def label_box(x, y, w, h, value, color, fill, size=9, name=None):
    shape = box(x, y, w, h, fill, color, 1.0, name=name)
    text(x + 4, y + 3, w - 8, h - 6, value, size, INK, True, name=f"{name} label" if name else None)
    return shape


def token(x, y, w=18, h=20, color=BLUE, fill=BLUE_TOKEN, value="", size=7.2, name=None):
    rect(x, y, w, h, fill, color, .9, name)
    if value:
        text(x, y, w, h, value, size, INK, True, font="Cambria Math")


def token_row(x, y, count, color, fill, w=18, h=20, gap=5, labels=None, name="tokens"):
    for i in range(count):
        token(x + i * (w + gap), y, w, h, color, fill,
              labels[i] if labels else "", name=f"{name} {i + 1}")


def grid(x, y, w, h, color, rows=3, cols=5, fill=None, name="grid"):
    shape = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, ex(x), ey(y), ex(w), ey(h))
    shape.name = name
    shape.rotation = 180
    shape.fill.solid() if fill else shape.fill.background()
    if fill:
        shape.fill.fore_color.rgb = rgb(fill)
        shape.fill.transparency = 32
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(.8)
    clean(shape)
    inset = w * .13
    for i in range(1, cols):
        t = i / cols
        connector(x + inset + t * (w - 2 * inset), y,
                  x + t * w, y + h, color, .45, False, name=f"{name} col {i}")
    for i in range(1, rows):
        yy = y + i * h / rows
        shrink = inset * (1 - i / rows)
        connector(x + shrink, yy, x + w - shrink, yy, color, .45, False,
                  name=f"{name} row {i}")


def crop_blob(image, bounds):
    stream = BytesIO()
    image.crop(bounds).save(stream, format="PNG")
    return stream.getvalue()


def picture(blob, x, y, w, h, name):
    shape = slide.shapes.add_picture(BytesIO(blob), ex(x), ey(y), ex(w), ey(h))
    shape.name = name
    return clean(shape)


def chart_icon(x, y):
    box(x, y, 51, 61, WHITE, "555555", 1.0, name="Global output icon")
    connector(x + 12, y + 48, x + 12, y + 16, "555555", 1.0, False)
    connector(x + 12, y + 48, x + 40, y + 48, "555555", 1.0, False)
    for bx, by, bw, bh in [(18, 33, 5, 15), (27, 24, 5, 24), (36, 12, 5, 36)]:
        rect(x + bx, y + by, bw, bh, "555555", "555555", .4)


def document_icon(x, y):
    rect(x, y, 22, 33, WHITE, "555555", .9)
    for yy in (9, 15, 21, 27):
        connector(x + 5, y + yy, x + 17, y + yy, "555555", .55, False)


def clipboard_icon(x, y):
    rect(x, y + 3, 22, 30, WHITE, "555555", .9)
    box(x + 6, y, 10, 7, WHITE, "555555", .7, .15)
    for yy in (12, 19, 26):
        connector(x + 5, y + yy, x + 17, y + yy, "555555", .55, False)


def clock_icon(x, y):
    oval(x, y, 31, 31, WHITE, "555555", .9)
    connector(x + 15.5, y + 5, x + 15.5, y + 16, "555555", .8, False)
    connector(x + 15.5, y + 16, x + 23, y + 21, "555555", .8, False)


def people_icon(x, y):
    oval(x + 8, y, 11, 11, "777777", "777777", .5)
    oval(x, y + 6, 9, 9, "777777", "777777", .5)
    oval(x + 20, y + 6, 9, 9, "777777", "777777", .5)
    box(x + 5, y + 12, 18, 19, "F4F4F4", "777777", .8, .22)


source = Image.open(SOURCE)
scans = {
    "current_ct": crop_blob(source, (27, 95, 121, 179)),
    "current_mri": crop_blob(source, (141, 94, 223, 180)),
    "prior_xr": crop_blob(source, (22, 241, 82, 318)),
    "prior_brain": crop_blob(source, (96, 241, 157, 318)),
    "prior_bone": crop_blob(source, (165, 241, 229, 318)),
    "follow_ct": crop_blob(source, (270, 804, 361, 887)),
    "follow_mri": crop_blob(source, (384, 803, 457, 887)),
    "segment": crop_blob(source, (1482, 226, 1537, 291)),
}

# Connectors are drawn first so every module stays visually on top.
connector(270, 207, 306, 207, BLUE, 1.3, True, name="Studies to medical encoder")
path([(407, 235), (447, 235), (447, 158), (503, 158)], BLUE, 1.3, True, "Encoder to dense memory")
path([(349, 305), (349, 352), (394, 352)], BLUE, 1.3, True, "Encoder to resampler")
connector(503, 352, 529, 352, BLUE, 1.3, True, name="Resampler to visual tokens")
connector(664, 352, 683, 374, BLUE, 1.3, True, name="Visual tokens to VLM input")
connector(266, 500, 320, 500, PURPLE, 1.25, True, name="Metadata to tokenizer")
connector(425, 506, 462, 506, PURPLE, 1.25, True, name="Tokenizer to text tokens")
path([(611, 506), (661, 506), (661, 374), (683, 374)], PURPLE, 1.25, True, "Text tokens to VLM input")
connector(835, 392, 835, 471, PURPLE, 1.25, True, name="VLM input to VLM")
connector(835, 540, 835, 578, PURPLE, 1.25, True, name="VLM to patient slots")
path([(1023, 596), (1041, 596), (1041, 169), (1105, 169)], PURPLE, 1.35, True, "Slots to query decoder")
path([(678, 214), (1073, 214), (1073, 268), (1400, 268), (1400, 254)], BLUE, 1.25, True, "Dense bypass to task head")
connector(1165, 91, 1165, 124, PURPLE, 1.2, True, name="Task query to decoder")
connector(1244, 191, 1283, 191, PURPLE, 1.25, True, name="Decoder to task latent")
path([(1309, 191), (1327, 191), (1327, 217), (1349, 217)], PURPLE, 1.25, True, "Task latent to task head")
path([(1448, 216), (1460, 216), (1460, 154), (1484, 154)], PURPLE, 1.25, True, "Head to global output")
path([(1448, 216), (1460, 216), (1460, 260), (1482, 260)], PURPLE, 1.25, True, "Head to dense output")

# Forecasting branch.
path([(1041, 374), (1258, 374), (1258, 455)], PURPLE, 1.15, True, "Slots to conditional state")
path([(1174, 423), (1191, 423), (1191, 491), (1208, 491)], ORANGE, 1.2, True, "Transition query to conditional state")
connector(1174, 524, 1208, 524, ORANGE, 1.2, True, name="Horizon to conditional state")
connector(1309, 491, 1334, 491, ORANGE, 1.25, True, name="Conditional state to predictor")
connector(1174, 543, 1334, 543, ORANGE, 1.05, True, name="Horizon to predictor")
connector(1334, 269, 1334, 482, PURPLE, 1.15, True, name="Patient slots to predictor")
connector(1390, 269, 1390, 482, BLUE, 1.15, True, name="Dense memory to predictor")
connector(1433, 522, 1453, 522, ORANGE, 1.25, True, name="Predictor to future latents")

# Target path and loss connections.
connector(4, 724, 1664, 724, GREEN, 1.15, False, True, "Training-only divider")
connector(494, 834, 547, 834, GREEN, 1.25, True, name="Follow-up to target encoder")
path([(726, 842), (758, 842), (758, 773), (786, 773)], GREEN, 1.15, True, "Target encoder to global latent")
path([(726, 842), (758, 842), (758, 850), (787, 850)], GREEN, 1.15, True, "Target encoder to regional latents")
path([(957, 773), (1209, 773), (1209, 844), (1251, 844)], GREEN, 1.2, True, "Global target to loss")
connector(1094, 850, 1251, 850, GREEN, 1.2, True, name="Regional target to loss")
path([(1507, 686), (1507, 844), (1362, 844)], ORANGE, 1.25, True, "Predicted latents to loss")

# Left evidence cards.
text(27, 9, 360, 29, "Current evidence at or before t", 12.2, INK, True,
     align=PP_ALIGN.LEFT, name="Current evidence heading")
box(16, 52, 255, 286, WHITE, GRAY, .9, name="Current studies card")
text(50, 65, 185, 28, "Current study Xₜ", 10.5, BLUE, True)
picture(scans["current_ct"], 27, 95, 94, 84, "Current CT")
picture(scans["current_mri"], 141, 94, 82, 86, "Current MRI")
connector(27, 198, 255, 198, GRAY, .7, False, name="Current/prior divider")
text(48, 210, 190, 29, "Optional prior studies", 10.2, BLUE, True)
picture(scans["prior_xr"], 22, 241, 60, 77, "Prior X-ray")
picture(scans["prior_brain"], 96, 241, 61, 77, "Prior brain scan")
picture(scans["prior_bone"], 165, 241, 64, 77, "Prior bone scan")
text(237, 265, 25, 35, "…", 13, INK, True)

box(16, 355, 251, 266, WHITE, GRAY, .9, name="Clinical metadata card")
document_icon(34, 376)
clipboard_icon(34, 435)
clock_icon(31, 495)
people_icon(30, 554)
for yy, label, value, italic in [
    (376, "Report", "…mild\ncardiomegaly…", True),
    (436, "History", "HTN, DM2,\nformer smoker…", False),
    (496, "Timing", "Study time\n2024-05-10 09:14", False),
    (555, "Clinical\ncontext", "Age 62, Male\nInpatient, ICU…", True),
]:
    text(69, yy, 70, 36, label, 9.0, INK, True, align=PP_ALIGN.LEFT)
    text(146, yy, 111, 38, value, 8.0, INK, False, italic, PP_ALIGN.LEFT)

# Image representation path.
label_box(306, 159, 101, 146, "Medical\n2D / 3D\nencoder", BLUE, BLUE_FILL, 10.5, "Medical encoder")
text(368, 171, 29, 30, "❄", 16, "2E9EEA")
label_box(394, 322, 109, 60, "Resampler\n+ projector", BLUE, BLUE_FILL, 9.2, "Resampler")
text(494, 30, 180, 51, "Multiscale dense\nspatial memory", 9.8, INK, True)
text(537, 71, 105, 32, "Dₜ¹:ᴸ", 15.0, INK, True, True, font="Cambria Math")
grid(508, 112, 111, 40, BLUE, 3, 5, BLUE_FILL, "Dense map 1/4")
grid(512, 165, 105, 39, BLUE, 3, 5, BLUE_FILL, "Dense map 1/8")
grid(520, 216, 94, 34, BLUE, 3, 5, BLUE_FILL, "Dense map 1/16")
text(624, 119, 42, 27, "1/4", 9.2, INK, True)
text(624, 173, 42, 27, "1/8", 9.2, INK, True)
text(624, 224, 42, 27, "1/16", 9.2, INK, True)
text(549, 256, 30, 45, "⋮", 16, INK, True)
text(651, 101, 42, 207, "}", 65, BLUE, False, font="Cambria Math")

token_row(531, 343, 6, BLUE, BLUE_TOKEN, 18, 19, 5, name="Visual token")
text(543, 366, 108, 41, "Visual tokens\n(learned)", 8.5, INK, True)

# Text path and ordered sequence.
label_box(320, 465, 106, 84, "VLM native\ntokenizer /\nembedding", PURPLE, PURPLE_FILL, 8.7, "VLM tokenizer")
token_row(464, 495, 7, PURPLE, PURPLE_FILL, 18, 19, 5, name="Native text token")
text(487, 520, 105, 40, "Text tokens\n(native)", 8.3, INK, True)

text(714, 326, 300, 25, "VLM input (single ordered token row)", 8.5, INK, True)
box(683, 352, 359, 40, WHITE, GRAY, .8, name="Ordered VLM input")
token_row(694, 361, 4, BLUE, BLUE_TOKEN, 18, 20, 5, name="Input visual token")
connector(783, 358, 783, 385, GRAY, .6, False)
token_row(795, 361, 4, PURPLE, PURPLE_FILL, 18, 20, 5, name="Input text token")
connector(884, 358, 884, 385, GRAY, .6, False)
token_row(893, 361, 8, PURPLE, PURPLE_TOKEN, 16, 20, 2,
          [f"s{i}" for i in range(1, 9)], "Learned input slot")
text(685, 394, 99, 25, "visual tokens", 7.4)
text(787, 394, 96, 25, "text tokens", 7.4)
text(878, 391, 170, 59, "8 learned slot\nembeddings s₁ … s₈\n(input slots)", 7.4, INK, True)

label_box(687, 472, 337, 68, "VLM + trainable LoRA", PURPLE, PURPLE_FILL, 11.2, "VLM with LoRA")
text(979, 485, 28, 36, "♦", 18, PURPLE)
text(629, 574, 46, 42, "Sₜ", 15, INK, True, True, font="Cambria Math")
token_row(688, 579, 8, PURPLE, PURPLE_TOKEN, 34, 34, 9,
          [f"S{i}" for i in range(1, 9)], "Patient-state slot")
text(606, 616, 150, 42, "patient-state slots\n(output states)", 7.7, INK, True)

# Open grouping for reusable representation.
path([(530, 654), (530, 660), (538, 668), (774, 668), (785, 674)], PURPLE, 1.1, False, "Representation brace left")
path([(785, 674), (796, 668), (1031, 668), (1039, 660), (1039, 654)], PURPLE, 1.1, False, "Representation brace right")
text(566, 675, 455, 31, "Reusable current representation  Rₜ = (Dₜ¹:ᴸ, Sₜ)", 11.5, INK, True)
text(552, 702, 485, 20, "Rₜ is formed before any task query or forecast horizon is supplied.", 8.0, INK, False, True)

# Query decoder and outputs.
label_box(1114, 35, 101, 56, "Task query\n(q)", PURPLE, PURPLE_FILL, 8.7, "Task query")
label_box(1105, 124, 139, 108, "Read-only\nquery decoder\n(frozen)", PURPLE, PURPLE_FILL, 9.3, "Read-only query decoder")
text(1200, 132, 26, 30, "▣", 14, PURPLE)
text(1263, 125, 112, 54, "eₜ⁽ᑫ⁾\n(task latent)", 8.7, INK, True, font="Cambria Math")
token(1283, 178, 26, 27, PURPLE, PURPLE_FILL, name="Task latentRect")
label_box(1349, 179, 99, 75, "Lightweight\ntask head", PURPLE, PURPLE_FILL, 8.7, "Task head")
chart_icon(1484, 123)
text(1545, 123, 126, 62, "Global clinical\noutput\n(e.g., risk score)", 8.3, INK, True, align=PP_ALIGN.LEFT)
picture(scans["segment"], 1482, 228, 55, 63, "Dense output icon")
text(1544, 228, 128, 62, "Dense segmentation /\nlocalization output", 7.9, INK, True, align=PP_ALIGN.LEFT)

# Forecast modules and outputs.
label_box(1068, 397, 106, 52, "Transition query\n(τ)", ORANGE, ORANGE_FILL, 7.8, "Transition query")
label_box(1068, 486, 106, 77, "Horizon\nh\n(e.g., 3 months)", ORANGE, ORANGE_FILL, 7.8, "Forecast horizon")
label_box(1208, 455, 101, 72, "Conditional\nrepresentation", ORANGE, ORANGE_FILL, 6.5, "Conditional representation")
text(1233, 503, 54, 25, "Aₜ,ₕ", 9.5, ORANGE, True, True, font="Cambria Math")
label_box(1334, 482, 99, 83, "Latent\npredictor", ORANGE, ORANGE_FILL, 9.0, "Latent predictor")
text(1214, 565, 224, 56, "Inputs to predictor (exactly 4):\nDₜ¹:ᴸ,  Sₜ,  Aₜ,ₕ,  h", 7.6, INK, True, font="Cambria Math")

token(1487, 394, 25, 27, ORANGE, ORANGE_FILL, name="Predicted global latent")
text(1532, 383, 137, 61, "Predicted future\nGLOBAL latent\n(1 vector)", 8.3, ORANGE, True)
grid(1451, 483, 60, 36, ORANGE, 3, 5, ORANGE_FILL, "Predicted map header")
grid(1455, 527, 71, 37, ORANGE, 3, 5, ORANGE_FILL, "Predicted map 1/4")
grid(1465, 571, 62, 34, ORANGE, 3, 5, ORANGE_FILL, "Predicted map 1/8")
grid(1474, 612, 53, 31, ORANGE, 3, 5, ORANGE_FILL, "Predicted map 1/16")
text(1521, 452, 150, 64, "Predicted future\nREGIONAL multiscale\nlatents (dense)", 8.1, ORANGE, True)
text(1534, 524, 45, 30, "1/4", 8.5, INK, True)
text(1534, 569, 45, 30, "1/8", 8.5, INK, True)
text(1534, 610, 45, 30, "1/16", 8.5, INK, True)
text(1484, 649, 30, 39, "⋮", 15, INK, True)
text(1434, 466, 30, 187, "{", 58, ORANGE, False, font="Cambria Math")

# Training-only target pathway.
text(24, 729, 250, 50, "Training-only target pathway\n(not used online)", 9.0, GREEN, True, align=PP_ALIGN.LEFT)
box(243, 759, 251, 148, GREEN_FILL, GREEN, 1.0, name="Observed follow-up card")
text(255, 766, 226, 29, "Observed follow-up study Xₜ₊", 9.0, GREEN, True)
picture(scans["follow_ct"], 270, 804, 91, 83, "Follow-up CT")
picture(scans["follow_mri"], 384, 803, 73, 84, "Follow-up MRI")
label_box(547, 797, 179, 88, "EMA target encoder\n· stop-gradient", GREEN, GREEN_FILL, 9.0, "EMA target encoder")
text(703, 785, 27, 28, "▣", 14, GREEN)

text(786, 744, 103, 58, "Target future\nGLOBAL latent\n(1 vector)", 7.6, GREEN, True)
token(931, 757, 25, 27, GREEN, GREEN_FILL, name="Target global latent")
text(786, 816, 127, 66, "Target future\nREGIONAL multiscale\nlatents (dense)", 7.4, GREEN, True)
grid(922, 801, 81, 31, GREEN, 3, 5, GREEN_FILL, "Target map 1/4")
grid(927, 839, 74, 29, GREEN, 3, 5, GREEN_FILL, "Target map 1/8")
grid(934, 875, 66, 25, GREEN, 3, 5, GREEN_FILL, "Target map 1/16")
text(1018, 802, 47, 26, "1/4", 8.1, INK, True)
text(1018, 840, 47, 26, "1/8", 8.1, INK, True)
text(1018, 875, 47, 26, "1/16", 8.1, INK, True)
text(955, 902, 25, 30, "⋮", 13, INK, True)

oval(1251, 805, 111, 80, GREEN_FILL, GREEN, 1.0, name="Alignment loss")
text(1266, 820, 81, 49, "Alignment\nloss", 9.2, INK, True)
text(1547, 805, 115, 80, "future loss\ntrains the\nonline path", 8.7, GREEN, True)

OUTPUT.parent.mkdir(exist_ok=True)
prs.save(OUTPUT)
print(OUTPUT)
