"""Recreate fig1_v5 with editable PowerPoint shapes and tight vector exports."""
from io import BytesIO
from pathlib import Path
import subprocess
import tempfile

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "ppt/fig1_v5.pptx"
SOURCE = ROOT / "source/fig1_v5.png"
W, H, Y0, PPI = 1864, 843, 0, 120
INK, WHITE = "171717", "FFFFFF"
BLUE, PINK, GREEN, BEIGE = "9BCAE9", "E7BCE7", "C7E6A9", "EBD9C6"
PLUM, OLIVE, BROWN = "68256C", "557C2F", "8C6639"
CREAM, PALE_BLUE, SLOT, GRAY = "FEFAD1", "D0E7FA", "E89EBA", "BEBEBE"
FONT = "Comic Sans MS"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W / PPI), Inches(H / PPI)
prs.core_properties.title = "MedWorld-JEPA — fig1_v5"
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)


def u(px):
    return Inches(px / PPI)


def style(shape, fill, stroke=INK, width=1.0, name=""):
    shape.name = name or shape.name
    shape.shadow.inherit = False
    effect = shape._element.find("./" + qn("p:style") + "/" + qn("a:effectRef"))
    if effect is not None:
        effect.set("idx", "0")
    if hasattr(shape, "fill"):
        if fill:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            shape.fill.background()
    if stroke:
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(width)
    else:
        shape.line.fill.background()
    return shape


def box(x, y, w, h, fill=WHITE, stroke=INK, radius=0.08, name="", width=1.0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, u(x), u(y - Y0), u(w), u(h))
    if radius:
        s.adjustments[0] = radius
    return style(s, fill, stroke, width, name)


def text(x, y, w, h, value, size=20, bold=False, font=FONT,
         color=INK, align=PP_ALIGN.CENTER, italic=False, rotation=0):
    s = slide.shapes.add_textbox(u(x), u(y - Y0), u(w), u(h))
    s.name = "Label: " + value.replace("\n", " / ")
    s.rotation = rotation
    tf = s.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, label in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.04
        r = p.add_run()
        r.text = label
        r.font.name = font
        r.font.size = Pt(size * 72 / PPI)
        r.font.bold, r.font.italic = bold, italic
        r.font.color.rgb = RGBColor.from_string(color)
    return s


def line(points, arrow=False, color=INK, width=1.5, name="Connection"):
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      u(x1), u(y1 - Y0), u(x2), u(y2 - Y0))
        style(s, None, color, width, name)
        if arrow and i == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "sm", "len": "sm"}.items():
                end.set(key, value)
            s._element.spPr.get_or_add_ln().append(end)


def polygon(points, fill, stroke=INK, name="Freeform"):
    points = [(x, y - Y0) for x, y in points]
    builder = slide.shapes.build_freeform(*points[0], scale=u(1))
    builder.add_line_segments(points[1:], close=True)
    return style(builder.convert_to_shape(), fill, stroke, 1.0, name)


def trapezoid(x, y, w, h, fill, name, taper=0.24):
    # A native trapezoid rotated so the output edge is the short edge.
    s = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID,
                              u(x + (w - h) / 2), u(y - Y0 + (h - w) / 2),
                              u(h), u(w))
    s.rotation = 90
    s.adjustments[0] = taper
    return style(s, fill, INK, 1.1, name)


def token(x, y, color, w=27, h=28, name="Token"):
    return box(x, y, w, h, color, "787878", 0.06, name, 0.6)


def thumbnail(bounds, x, y, w, h, name):
    data = BytesIO()
    with Image.open(SOURCE) as im:
        im.crop(bounds).save(data, format="PNG")
    data.seek(0)
    s = slide.shapes.add_picture(data, u(x), u(y - Y0), u(w), u(h))
    s.name = name


def gradient(s, first, last):
    s.fill.gradient()
    s.fill.gradient_angle = 90
    for stop, color in zip(s.fill.gradient_stops, [first, last]):
        stop.color.rgb = RGBColor.from_string(color)
    return s


def math(x, y, w, h, runs, size=26):
    """Editable Times math: (text, baseline percentage) runs."""
    s = text(x, y, w, h, "", size, font="Times New Roman", italic=True)
    s.name = "Equation: " + "".join(v for v, _ in runs)
    p = s.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    for value, baseline in runs:
        r = p.add_run()
        r.text = value
        r.font.name = "Times New Roman"
        r.font.size = Pt(size * 72 / PPI)
        r.font.italic = True
        r.font.color.rgb = RGBColor.from_string(INK)
        if baseline:
            r._r.get_or_add_rPr().set("baseline", str(baseline * 1000))
    return s


def grid(x, y, step=22, side=18):
    for row in range(3):
        for col in range(3):
            gradient(box(x + col * step, y + row * step, side, side,
                         BLUE, "377CB6", 0, "Visual feature", 0.75),
                     "B4D9F8", "78B2E6")


def slots(x, y, count=3, step=23, w=17, h=25):
    for i in range(count):
        gradient(box(x + step * i, y, w, h, SLOT, "98346E", 0.15,
                     "Learned state slot", 0.9), "EDB5D7", "D68ABC")


def curve(points, name):
    """A native editable connector with a cubic path in slide coordinates."""
    x, y = points[0]
    coords = [(x, y)] + [pt for _, segment in points[1:] for pt in segment]
    left, top = min(p[0] for p in coords), min(p[1] for p in coords)
    w, h = max(p[0] for p in coords) - left, max(p[1] for p in coords) - top
    s = box(left, top, w, h, None, INK, 0, name, 1.5)
    sp = s._element.spPr
    sp.remove(sp.find(qn("a:prstGeom")))
    geom = OxmlElement("a:custGeom")
    for tag in ["avLst", "gdLst", "ahLst", "cxnLst"]:
        geom.append(OxmlElement("a:" + tag))
    paths = OxmlElement("a:pathLst")
    path = OxmlElement("a:path")
    for key, val in {"w": str(w), "h": str(h), "fill": "none"}.items():
        path.set(key, val)
    for tag, coords in [("moveTo", [(x, y)]), *points[1:]]:
        command = OxmlElement("a:" + tag)
        for xx, yy in coords:
            pt = OxmlElement("a:pt")
            pt.set("x", str(xx - left)); pt.set("y", str(yy - top))
            command.append(pt)
        path.append(command)
    paths.append(path); geom.append(paths); sp.append(geom)
    end = OxmlElement("a:tailEnd")
    end.set("type", "triangle"); end.set("w", "sm"); end.set("len", "sm")
    sp.get_or_add_ln().append(end)


def report(y, follow=False):
    gradient(box(28, y, 189, 94, "FFF4DF", "725A3A", 0.12,
                 "Follow-up report" if follow else "Current report", 1),
             "FFF9EF", "FFF0D9")
    text(46, y + 4, 167, 25, "Follow-up report" if follow else "Report",
         19, True, align=PP_ALIGN.LEFT)
    polygon([(48, y + 38), (71, y + 38), (81, y + 48),
             (81, y + 82), (48, y + 82)], "FFF9EB", name="Report page")
    line([(71, y + 38), (71, y + 48), (81, y + 48)], width=1)
    for yy in [54, 62, 70]:
        line([(54, y + yy), (74, y + yy)], width=1)
    text(101, y + 33, 110, 42, "Findings: …\nImpression: …", 16,
         font="Arial", align=PP_ALIGN.LEFT)


def vlm(y):
    gradient(box(561, y, 402, 207, PINK, "762381", 0.09, "VLM", 1.3),
             "F0CFEF", "E5BCE8")
    text(710, y + 6, 80, 35, "VLM", 27, True)
    text(895, y + 9, 59, 28, "LoRA", 19, font="Arial")
    box(572, y + 49, 380, 76, "FFFBEF", "963879", 0.14, "VLM input tokens", 1)
    for x, w, label in [(581, 99, "Visual tokens"), (695, 96, "Text tokens"),
                         (832, 108, "Learned slots")]:
        text(x, y + 56, w, 24, label, 15.5, True)
    for i in range(3):
        gradient(box(586 + i * 25, y + 89, 20, 23, BLUE, "285784", 0.08,
                     "VLM visual token", 1), "B4D9FA", "75B6EC")
        gradient(token(690 + i * 25, y + 88, GRAY, 20, 24, "VLM text token"),
                 "D1D1CF", "B9B9B9")
    text(661, y + 85, 25, 30, "…", 20)
    text(765, y + 85, 25, 30, "…", 20)
    slots(797, y + 87, 5, 20.5, 15, 24)
    text(899, y + 86, 23, 28, "…", 19)
    slots(924, y + 87, 1, w=15, h=24)
    for xx in [671, 704]:
        gradient(box(xx, y + 136, 20, 58, CREAM, INK, 0.15, "Transformer block", 1.2),
                 "FFFFEC", "F3F5BF")
    text(745, y + 148, 32, 34, "…", 27)
    text(787, y + 151, 123, 30, "Transformer", 19, True)


def state(x, y, future=False, predicted=False):
    w, h = (228, 244) if predicted else (187, 244)
    if future:
        h = 258
    box(x, y, w, h, "FFFFFF", "5A6170", 0.08, "Predicted future" if predicted else
        "Future target" if future else "Current state", 1)
    title = "Predicted future" if predicted else "Future target" if future else "Current state"
    tw = 157 if predicted else 132
    text(x + 10, y + 6, tw, 31, title, 18.5, True, align=PP_ALIGN.LEFT)
    suffix = "t,h" if predicted else "t+" if future else "t"
    rr = [("R̂" if predicted else "R", 0)]
    if future:
        rr += [("*", 40)]
    rr += [(suffix, -25)]
    math(x + tw + 14, y + 4, w - tw - 14, 36, rr, 26)
    by = y + (48 if predicted or future else 45)
    bh = 76 if predicted else 82 if future else 80
    gradient(box(x + 12, by, w - 24, bh, PALE_BLUE, "78AFF0", 0.14,
                 "Image state panel", 0.9), "E7F3FF", "CCE5FE")
    dx = x + (27 if predicted else 22)
    runs = [("D̂" if predicted else "D", 0)]
    if future:
        runs += [("*", 40)]
    math(dx, by + 9, 66, 40, runs + [(suffix, -25)], 29)
    grid(x + (89 if predicted else 73), by + 9)
    text(x + w - 51, by + 23, 34, 31, "…", 26)
    sy = by + bh + 9
    gradient(box(x + 12, sy, w - 24, 61 if not future else 64,
                 "F9E5F1", "D482BB", 0.15, "Semantic state panel", 0.9),
             "FDF0F8", "F5D7EB")
    ss = [("Ŝ", 0)]
    if future or predicted:
        ss += [("*", 40)]
    math(dx, sy + 5, 71, 42, ss + [(suffix, -25)], 29)
    slots(x + (88 if predicted else 65 if not future else 73), sy + 18)
    text(x + w - 51, sy + 16, 34, 31, "…", 26)
    if predicted:
        eq = [("R̂", 0), ("*", 40), ("t,h", -25), (" = [D", 0), ("*", 40),
              ("t+", -25), ("; Ŝ", 0), ("*", 40), ("t,h", -25), ("]", 0)]
    elif future:
        eq = [("R", 0), ("*", 40), ("t+", -25), (" = [D", 0), ("*", 40),
              ("t+", -25), ("; Ŝ", 0), ("*", 40), ("t+", -25), ("]", 0)]
    else:
        eq = [("R", 0), ("t", -25), (" = [D", 0), ("t", -25), ("; S", 0), ("t", -25), ("]", 0)]
    math(x + (22 if predicted else 10 if future else 32), y + h - 46, w - 15, 43,
         eq, 25 if not predicted else 26)


# Two copies of the shared encoding pipeline.
text(36, 1, 555, 47, "Shared multimodal state encoding", 31.5, True, align=PP_ALIGN.LEFT)
text(603, 13, 445, 33, "Image and/or report inputs supported", 21,
     font="Arial", align=PP_ALIGN.LEFT)
for future, y in [(False, 57), (True, 400)]:
    box(10, y, 1192, 315 if not future else 319, None, "64768B", 0.055,
        "Observed future encoding" if future else "Current observations encoding", 1)
    text(29, y + 1, 299 if not future else 258, 40,
         "Observed future, t⁺" if future else "Current observations, t", 25, True,
         align=PP_ALIGN.LEFT)
    if future:
        text(291, y + 6, 160, 31, "Training target", 18, align=PP_ALIGN.LEFT)
    iy = 477 if future else 130
    text(30 if future else 62, iy - 31, 170 if future else 100, 29,
         "Follow-up image" if future else "Image", 20, True, align=PP_ALIGN.LEFT)
    thumbnail((30, iy + 1, 164, iy + 113), 29, iy, 135, 113,
              "Follow-up chest radiograph" if future else "Current chest radiograph")
    ey = 459 if future else 113
    trapezoid(194, ey, 114, 130, "BBDDFA", "V-JEPA encoder", 0.30)
    text(203, ey + 35, 97, 67, "V-JEPA\nencoder", 22, True)
    gy = 502 if future else 156
    grid(342, gy, 19, 16)
    math(352 if not future else 343, gy - 40, 65, 37,
         [("D", 0)] + ([("*", 40), ("t+", -25)] if future else [("t", -25)]), 28)
    ay = 490 if future else 143
    gradient(box(426, ay, 105, 75, CREAM, "514D20", 0.16, "Visual adapter", 1.2),
             "FFFDE6", "F9F9CB")
    text(432, ay + 6, 93, 63, "Visual\nadapter", 20, True)
    ry = 606 if future else 259
    report(ry, future)
    ty = 620 if future else 273
    gradient(box(258, ty, 131, 67, CREAM, "514D20", 0.18, "Tokenizer + embedding", 1.2),
             "FFFDE6", "F9F9CB")
    text(266, ty + 6, 115, 54, "Tokenizer +\nembedding", 20, True)
    for i in range(3):
        gradient(token(424 + i * 26, ty + 20, GRAY, 20, 24, "Report text token"),
                 "D0D0CF", "BABABA")
    text(420, ty + 46, 102, 32, "Text tokens", 18, True, align=PP_ALIGN.LEFT)
    vy = 469 if future else 124
    vlm(vy)
    state(1001, 440 if future else 104, future)
    mid = ay + 37
    line([(164, mid), (194, mid)], True, name="Image to encoder")
    line([(308, mid), (339, mid)], True, name="Encoder to image features")
    line([(396, mid), (426, mid)], True, name="Image features to adapter")
    line([(531, mid), (561, mid)], True, name="Visual input to VLM")
    line([(217, ty + 33), (258, ty + 33)], True, name="Report to tokenizer")
    line([(389, ty + 33), (420, ty + 33)], True, name="Tokenizer to text tokens")
    line([(500, ty + 33), (530, ty + 33), (530, vy + 154), (561, vy + 154)], True,
         name="Text input to VLM")
    line([(963, vy + 106), (1001, vy + 106)], True, name="VLM to multimodal state")
text(408, 372, 480, 27, "Same encoding pipeline at both times", 20, True, italic=True,
     align=PP_ALIGN.LEFT)

# Current state branches to the task decoder and latent predictor.
curve([(1188, 230), ("lnTo", [(1208, 230)]),
       ("cubicBezTo", [(1245, 230), (1210, 198), (1251, 198)]),
       ("lnTo", [(1278, 198)])], "Current state to task decoder")
curve([(1188, 230), ("lnTo", [(1208, 230)]),
       ("cubicBezTo", [(1220, 230), (1227, 236), (1227, 251)]),
       ("lnTo", [(1227, 375)]),
       ("cubicBezTo", [(1227, 391), (1233, 396), (1249, 396)]),
       ("lnTo", [(1318, 396)])], "Current state to latent world model")
text(1245, 62, 119, 29, "Task query", 20, True)
gradient(box(1292, 94, 25, 26, BLUE, "234A7C", 0.12, "Task query", 1.2),
         "B6DDFF", "7BB8EE")
line([(1304, 120), (1304, 134)], True)
gradient(box(1278, 134, 201, 131, GREEN, "3E5922", 0.15, "Task decoder", 1.3),
         "D9ECC1", "C8DFAD")
text(1304, 139, 161, 36, "Task decoder", 26, True)
for xx in [1330, 1407]:
    gradient(box(xx, 180, 21, 54, CREAM, INK, 0.13, "Task transformer block", 1.2),
             "FFFFEC", "F7F7CF")
text(1359, 184, 39, 42, "…", 29)
text(1313, 233, 137, 29, "Transformer", 20, True)
text(1219, 420, 87, 33, "Horizon h", 20, True)
token(1244, 454, GRAY, 29, 28, "Horizon token")
line([(1273, 468), (1318, 468)], True)
gradient(box(1318, 351, 267, 164, BEIGE, "704824", 0.13, "Latent World Model", 1.3),
         "F9E7D1", "EDD4BD")
text(1337, 360, 230, 39, "Latent World Model", 27, True)
text(1526, 398, 45, 25, "LWM", 19, font="Arial")
for xx in [1392, 1488]:
    gradient(box(xx, 414, 23, 59, CREAM, INK, 0.15, "Predictor block", 1.2),
             "FFFFEC", "F8F7D3")
text(1425, 419, 53, 40, "…", 28)
text(1377, 477, 153, 31, "Predictor block", 21, True)
line([(1585, 421), (1619, 421)], True, name="LWM to predicted future")
state(1619, 304, predicted=True)

# Task loss: four editable task cards, with raster medical image details only.
text(1616, 1, 140, 38, "Task loss", 28, True)
box(1516, 41, 339, 235, None, "73A8E2", 0.05, "Task losses", 0.9)
for xx, yy, title in [(1523, 49, "Classification"), (1691, 49, "Segmentation"),
                       (1523, 168, "Disease recognition"), (1691, 168, "Super-resolution")]:
    gradient(box(xx, yy, 157, 107 if yy == 49 else 100, "EDF6FF", "7CAFEF", 0.1,
                 title + " card", 0.8), "F8FCFF", "EAF4FE")
    text(xx + 4, yy + 4, 149, 29, title, 18 if len(title) < 18 else 17, True)
line([(1479, 198), (1489, 198), (1489, 106), (1516, 106)], True,
     name="Decoder to top task losses")
line([(1489, 198), (1516, 198)], True, name="Decoder to bottom task losses")
line([(1561, 142), (1631, 142)], color="537CA4", width=0.9)
for xx, yy in [(1569, 120), (1588, 105), (1607, 86)]:
    gradient(box(xx, yy, 12, 142 - yy, BLUE, "4488C6", 0, "Classification bar", 0.7),
             "A9D3F9", "6DAEE8")
text(1558, 81, 18, 28, "⋮", 17, color="74ADE0")
thumbnail((1725, 77, 1818, 151), 1725, 77, 93, 74, "Segmentation radiograph")
box(1533, 207, 136, 43, "F2F9FF", "557CA6", 0.16, "Disease label", 1)
text(1543, 211, 118, 31, "Disease: …", 20, True)
thumbnail((1707, 205, 1753, 252), 1707, 205, 46, 47, "Low-resolution detail")
thumbnail((1787, 202, 1833, 250), 1787, 202, 46, 48, "High-resolution detail")
box(1706, 204, 48, 49, None, "81A24A", 0, "Low-resolution image border", 0.8)
box(1786, 201, 48, 50, None, "4D6475", 0, "High-resolution image border", 0.8)
line([(1755, 228), (1784, 228)], True, width=1.2)
line([(1758, 221), (1763, 214), (1783, 204)], True, width=0.65)
text(1747, 246, 49, 24, "MSE", 16, True, font="Arial", italic=True)

# Stop-gradient supervision and two training stages.
line([(1723, 548), (1723, 604)], True, name="Prediction to latent loss")
line([(1188, 650), (1558, 650)], True, name="Stop-gradient future target")
slide.shapes[-1].line.dash_style = MSO_LINE_DASH_STYLE.DASH
text(1324, 615, 129, 33, "stop-grad", 21, True, italic=True)
gradient(box(1558, 604, 297, 92, "FAD8C3", "743C29", 0.16, "Latent prediction loss", 1.2),
         "FFE4D5", "F9D2BA")
text(1573, 612, 267, 37, "Latent prediction loss", 25, True)
math(1619, 651, 224, 37, [("D̂ ↔ D", 0), ("*", 40), ("; Ŝ ↔ S", 0), ("*", 40)], 29)
for x, w, fill, stroke, title, subtitle in [
    (30, 873, "EAF4DE", "527A33", "Stage 1: Task-supervised state pretraining",
     "Single-time task data · VLM state tokens + task decoders"),
    (963, 874, "FFF2DA", "79532D", "Stage 2: Future-state prediction fine-tuning",
     "Longitudinal pairs · LWM + current-state VLM")]:
    gradient(box(x, 739, w, 83, fill, stroke, 0.13, title, 1), "FFFDF5" if x > 900 else "F1F8E8", fill)
    text(x + 10, 747, w - 20, 36, title, 26, True)
    text(x + 10, 782, w - 20, 30, subtitle, 23, font="Arial")
line([(903, 782), (963, 782)], True, name="Stage 1 to Stage 2")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
with tempfile.TemporaryDirectory(prefix="fig1_v5_lo_") as profile:
    subprocess.run(["libreoffice", f"-env:UserInstallation={Path(profile).as_uri()}",
                    "--headless", "--convert-to", "pdf:impress_pdf_Export",
                    "--outdir", str(OUT.parent), str(OUT)], check=True, capture_output=True)
with fitz.open(OUT.with_suffix(".pdf")) as pdf:
    page = pdf[0]
    OUT.with_suffix(".svg").write_text(page.get_svg_image(text_as_path=True))
    page.get_pixmap(matrix=fitz.Matrix(3, 3)).save(OUT.with_suffix(".png"))
print(OUT)
