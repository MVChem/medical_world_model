"""Rebuild fig2_v1 with native PowerPoint objects and render paper assets."""

import argparse
from pathlib import Path
import shutil
import subprocess
import tempfile

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "source" / "fig2_v1.png"
OUTPUT = ROOT / "ppt" / "fig2_v1.pptx"
PPI = 120
INK = "101010"


def build():
    with Image.open(SOURCE) as reference:
        width, height = reference.size
    # Coordinates follow the 2048 px wide on-screen reference; preserve its
    # actual 2061 x 763 source canvas, including surrounding white space.
    scale = width / 2048

    def u(value):
        return Inches(value * scale / PPI)

    prs = Presentation()
    prs.slide_width = Inches(width / PPI)
    prs.slide_height = Inches(height / PPI)
    prs.core_properties.title = "fig2_v1 — Task-supervised slot representation"
    prs.core_properties.subject = "Editable reconstruction of the supplied reference"
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    def box(name, x, y, w, h, fill, stroke, radius, line_width=3):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, u(x), u(y), u(w), u(h)
        )
        shape.name = name
        shape.adjustments[0] = radius
        shape.shadow.inherit = False
        for effect in shape._element.xpath("./p:style/a:effectRef"):
            effect.set("idx", "0")
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(line_width * scale * 72 / PPI)
        return shape

    def label(name, value, x, y, w, h, size=44, bold=False,
              font="Comic Sans MS", align=PP_ALIGN.CENTER, italic=False):
        shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
        shape.name = name
        tf = shape.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, line in enumerate(value.split("\n")):
            paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            paragraph.alignment = align
            paragraph.space_before = paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.0
            run = paragraph.add_run()
            run.text = line
            run.font.name = font
            run.font.size = Pt(size * scale * 72 / PPI)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = RGBColor.from_string(INK)
        return shape

    def route(name, points, arrow=True):
        if len(points) == 2:
            (x1, y1), (x2, y2) = points
            shape = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2)
            )
        else:
            path = slide.shapes.build_freeform(*points[0], scale=u(1))
            path.add_line_segments(points[1:], close=False)
            shape = path.convert_to_shape()
            shape.fill.background()
        shape.name = name
        shape.shadow.inherit = False
        for effect in shape._element.xpath("./p:style/a:effectRef"):
            effect.set("idx", "0")
        shape.line.color.rgb = RGBColor.from_string(INK)
        shape.line.width = Pt(3.4 * scale * 72 / PPI)
        if arrow:
            end = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "lg", "len": "lg"}.items():
                end.set(key, value)
            shape._element.spPr.get_or_add_ln().append(end)
        return shape

    box("Panel: vision-language", 97, 206, 654, 168,
        "FCEFFB", "8F329B", .14)
    box("Panel: vision encoder", 97, 402, 654, 169,
        "E6F3FF", "0872CA", .14)
    box("Module: lightweight adapter", 853, 299, 270, 168,
        "FFFADB", "99882C", .14)
    box("Module: task decoder", 1531, 302, 235, 163,
        "DDEFD5", "287F31", .14)

    for branch, y, fill, stroke in (
        ("vision-language", 286, "EFA9D2", "CD45A7"),
        ("vision encoder", 483, "9FD5F5", "05A6E1"),
    ):
        for i in range(12):
            active = i % 3 == 1
            box(f"Token: {branch} {i + 1:02}", 125 + 50.7 * i, y, 42, 64,
                fill if active else "D1D1D1", stroke if active else "A3A3A3", .11)

    for branch, y, fill, stroke in (
        ("vision-language", 309, "EFA9D2", "CD45A7"),
        ("vision encoder", 392, "9FD5F5", "05A6E1"),
    ):
        for i in range(4):
            box(f"Slot: {branch} {i + 1}", 1198 + 52.3 * i, y, 42, 63,
                fill, stroke, .11)

    route("Arrow: vision-language to adapter",
          [(751, 286), (799, 286), (799, 343), (848, 343)])
    route("Arrow: vision encoder to adapter",
          [(751, 483), (799, 483), (799, 422), (848, 422)])
    route("Arrow: adapter to language slots", [(1123, 343), (1184, 343)])
    route("Arrow: adapter to visual slots", [(1123, 422), (1184, 422)])
    route("Bracket: merge eight slots",
          [(1413, 340), (1460, 340), (1460, 425), (1413, 425)], arrow=False)
    route("Arrow: slots to task decoder", [(1460, 382), (1520, 382)])
    route("Arrow: task decoder to loss", [(1766, 382), (1828, 382)])

    label("Label: VLM", "VLM", 319, 115, 178, 80, size=64, bold=True)
    label("Label: vision-language", "Vision-language", 120, 211, 580, 68,
          size=46, align=PP_ALIGN.LEFT)
    label("Label: vision encoder", "Vision encoder", 120, 408, 580, 68,
          size=46, align=PP_ALIGN.LEFT)
    label("Label: lightweight adapter", "Lightweight\nadapter",
          864, 324, 248, 112, size=41)
    label("Label: eight slots", "8 slots", 1179, 232, 237, 63, size=45)
    label("Label: task decoder", "Task\ndecoder", 1550, 326, 197, 111, size=44)
    label("Label: joint optimization", "Joint optimization",
          847, 575, 443, 69, size=41)
    label("Equation: slot dimensions", "8 × d", 1232, 465, 130, 64,
          size=49, font="Times New Roman", italic=True)
    loss = label("Equation: task loss", "L", 1847, 329, 175, 89,
                 size=69, font="Times New Roman", italic=True,
                 align=PP_ALIGN.LEFT)
    run = loss.text_frame.paragraphs[0].add_run()
    run.text = "task"
    run.font.name = "Times New Roman"
    # Office applies automatic subscript scaling to this nominal size.
    run.font.size = Pt(65 * scale * 72 / PPI)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(INK)
    run._r.get_or_add_rPr().set("baseline", "-23000")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


def render():
    """Export the saved PPTX; the PNG is never drawn independently."""
    with tempfile.TemporaryDirectory(prefix="fig2_v1_render_") as directory:
        work = Path(directory)
        result = subprocess.run([
            "libreoffice", f"-env:UserInstallation={(work / 'profile').as_uri()}",
            "--headless", "--convert-to", "pdf:impress_pdf_Export",
            "--outdir", str(work), str(OUTPUT),
        ], check=True, capture_output=True, text=True, timeout=120)
        rendered = work / OUTPUT.with_suffix(".pdf").name
        if not rendered.is_file():
            raise RuntimeError(result.stdout + result.stderr)
        shutil.copy2(rendered, OUTPUT.with_suffix(".pdf"))
    with fitz.open(OUTPUT.with_suffix(".pdf")) as pdf:
        assert len(pdf) == 1
        page = pdf[0]
        page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5)).save(OUTPUT.with_suffix(".png"))
        clip = fitz.Rect()
        for drawing in page.get_drawings():
            if (drawing["fill"] == (1, 1, 1)
                    and drawing["rect"].get_area() > .99 * page.rect.get_area()):
                continue
            clip |= drawing["rect"]
        for block in page.get_text("blocks"):
            clip |= fitz.Rect(block[:4])
        clip = (clip + (-6, -6, 6, 6)) & page.rect
        with fitz.open() as cropped:
            target = cropped.new_page(width=clip.width, height=clip.height)
            target.show_pdf_page(target.rect, pdf, 0, clip=clip)
            cropped.save(OUTPUT.with_name("fig2_v1_paper.pdf"), garbage=4, deflate=True)
    print(OUTPUT.with_suffix(".png"))
    print(OUTPUT.with_name("fig2_v1_paper.pdf"))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--render-only", action="store_true",
                        help="Export the existing PPTX, preserving manual edits")
    args = parser.parse_args()
    if not args.render_only:
        build()
    render()
