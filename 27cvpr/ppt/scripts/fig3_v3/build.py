"""Rebuild Figure 3 as one editable PowerPoint slide, using CXR imagery.

Feature fields are deterministic illustrations, not learned attention/results.
Run: python build.py --render
"""
from __future__ import annotations

import argparse
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile

import numpy as np
from PIL import Image, ImageFilter, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / 'assets'
DEFAULT_OUTPUT = HERE.parents[1] / 'ppt' / 'fig3_v3.pptx'
W, H = 2018, 779
EMU = 9144  # reference pixels -> 1/100 inch
INK = '090909'
PINK, BLUE = 'B21777', '0759BD'


def xml(tag, **attrs):
    node = OxmlElement(tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def gradient(shape, stops, angle=90):
    sp = shape._element.spPr
    for node in list(sp):
        if node.tag.split('}')[-1] in ('solidFill', 'noFill', 'gradFill'):
            sp.remove(node)
    grad = xml('a:gradFill', rotWithShape='1')
    lst = xml('a:gsLst')
    for pos, color in stops:
        stop = xml('a:gs', pos=int(pos * 1000))
        stop.append(xml('a:srgbClr', val=color))
        lst.append(stop)
    grad.append(lst)
    grad.append(xml('a:lin', ang=int(angle * 60000), scaled='1'))
    # DrawingML fill belongs after geometry and before line/effects.
    geom = sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}prstGeom')
    if geom is None:
        geom = sp.find('{http://schemas.openxmlformats.org/drawingml/2006/main}custGeom')
    sp.insert(list(sp).index(geom) + 1, grad)


def style(shape, fill=None, stroke=None, width=1.5, dash=False):
    for node in list(shape._element):
        if node.tag.split('}')[-1] == 'style':
            shape._element.remove(node)
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        shape.fill.background()
    if stroke:
        shape.line.color.rgb = RGBColor.from_string(stroke)
        shape.line.width = Pt(width * .72)
        if dash:
            shape._element.spPr.get_or_add_ln().append(xml('a:prstDash', val='dash'))
    else:
        shape.line.fill.background()
    shape._element.spPr.append(xml('a:effectLst'))
    return shape


def rect(slide, name, x, y, w, h, fill=None, stroke=None, width=1.5,
         radius=0, dash=False, kind=None, target=None):
    kind = kind or (MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE)
    shape = (target if target is not None else slide.shapes).add_shape(
        kind, int(x * EMU), int(y * EMU), int(w * EMU), int(h * EMU))
    shape.name = name
    if radius:
        shape.adjustments[0] = radius / min(w, h)
    return style(shape, fill, stroke, width, dash)


def path(slide, name, pts, stroke=INK, width=1.8, fill=None, close=False,
         arrow=False, target=None):
    x0, y0 = min(x for x, y in pts), min(y for x, y in pts)
    w = max(1, max(x for x, y in pts) - x0)
    h = max(1, max(y for x, y in pts) - y0)
    shape = rect(slide, name, x0, y0, w, h, fill, stroke, width, target=target)
    sp = shape._element.spPr
    sp.remove(sp.prstGeom)
    geom = xml('a:custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geom.append(xml('a:' + tag))
    geom.append(xml('a:rect', l='0', t='0', r='r', b='b'))
    paths = xml('a:pathLst')
    p = xml('a:path', w=round(w * 1000), h=round(h * 1000))
    for i, (x, y) in enumerate(pts):
        node = xml('a:moveTo' if i == 0 else 'a:lnTo')
        node.append(xml('a:pt', x=round((x - x0) * 1000), y=round((y - y0) * 1000)))
        p.append(node)
    if close:
        p.append(xml('a:close'))
    paths.append(p)
    geom.append(paths)
    sp.insert(1, geom)
    if arrow:
        sp.get_or_add_ln().append(xml('a:tailEnd', type='triangle', w='med', len='med'))
    return shape


def line(slide, name, x1, y1, x2, y2, width=1.8, color=INK,
         arrow=False, dash=False, target=None):
    shape = (target if target is not None else slide.shapes).add_connector(
        MSO_CONNECTOR.STRAIGHT,
        round(x1 * EMU), round(y1 * EMU), round(x2 * EMU), round(y2 * EMU))
    shape.name = name
    for node in list(shape._element):
        if node.tag.split('}')[-1] == 'style':
            shape._element.remove(node)
    shape.line.color.rgb = RGBColor.from_string(color)
    shape.line.width = Pt(width * .72)
    ln = shape._element.spPr.get_or_add_ln()
    if dash:
        ln.append(xml('a:prstDash', val='dash'))
    if arrow:
        ln.append(xml('a:tailEnd', type='triangle', w='med', len='med'))
    shape._element.spPr.append(xml('a:effectLst'))
    return shape


def label(slide, name, content, x, y, w, h, size=20, bold=False,
          color=INK, font='Arial', align=PP_ALIGN.CENTER, leading=1.08):
    """Text stays at slide level, independently selectable from all artwork."""
    shape = slide.shapes.add_textbox(round(x * EMU), round(y * EMU),
                                     round(w * EMU), round(h * EMU))
    shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = content.split('\n') if isinstance(content, str) else [content]
    for i, row in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = Pt(size * leading * .72)
        spans = [{'text': row}] if isinstance(row, str) else row
        for span in spans:
            run = p.add_run()
            run.text = span['text']
            run.font.name = span.get('font', font)
            run.font.size = Pt(span.get('size', size) * .72)
            run.font.bold = span.get('bold', bold)
            run.font.italic = span.get('italic', False)
            run.font.color.rgb = RGBColor.from_string(span.get('color', color))
            if 'baseline' in span:
                run._r.get_or_add_rPr().set('baseline', str(span['baseline']))
    return shape


def math_label(slide, name, main, suffix, x, y, w, h, size=30, sub=False):
    spans = [{'text': main, 'font': 'Times New Roman', 'italic': True}]
    if suffix:
        spans.append({'text': suffix, 'font': 'Times New Roman',
                      'italic': not sub, 'size': size * .96,
                      'baseline': -24000 if sub else 33000})
    return label(slide, name, spans, x, y, w, h, size=size)


def picture(slide, name, im, x, y, w, h, target=None):
    stream = BytesIO()
    im.save(stream, format='PNG')
    stream.seek(0)
    shape = (target if target is not None else slide.shapes).add_picture(
        stream, round(x * EMU), round(y * EMU), round(w * EMU), round(h * EMU))
    shape.name = name
    return shape


def photo(slide, name, im, x, y, w, h):
    # Letterbox rather than stretching anatomy to the old MRI aspect ratio.
    framed = ImageOps.pad(im.convert('RGB'), (round(w * 6), round(h * 6)),
                          method=Image.Resampling.LANCZOS, color='black')
    return picture(slide, name, framed, x, y, w, h)


def token(slide, name, x, y, w, h, blue=False, faint=False):
    stroke = ('5899D5' if faint else BLUE) if blue else ('CD8BB5' if faint else PINK)
    colors = ('D3EDFD', '8ACCF3', 'B2DBF7') if blue else ('F8E6F2', 'E8BCD7', 'F0CDE3')
    shape = rect(slide, name, x, y, w, h, colors[0], stroke, 1.65, radius=3.4)
    gradient(shape, [(0, colors[0]), (22, colors[1]), (100, colors[2])], angle=0)
    return shape


def module(slide, name, text, x, y, w, h, size=20):
    shape = rect(slide, name + ' background', x, y, w, h, 'FFF3CD', 'E6B523', 1.3, radius=9)
    gradient(shape, [(0, 'FFF7DE'), (55, 'FFEFC0'), (100, 'FFF7DF')])
    label(slide, name + ' label', text, x + 3, y + 2, w - 6, h - 4, size, bold=True)


def feature_fields(cxr):
    """Small illustrative dense fields with CXR texture and smooth spatial peaks.

    The target and reconstruction stacks intentionally share structure. No model
    is run, and these arrays must not be used as experimental evidence.
    """
    gray = np.asarray(cxr.resize((224, 224)), dtype=float) / 255.
    blur = np.asarray(cxr.resize((224, 224)).filter(ImageFilter.GaussianBlur(9)), dtype=float) / 255.
    yy, xx = np.mgrid[0:1:224j, 0:1:224j]
    lungs = (np.exp(-((xx - .29) / .15) ** 2 - ((yy - .43) / .28) ** 2)
             + np.exp(-((xx - .72) / .15) ** 2 - ((yy - .43) / .28) ** 2))
    lung_dark = np.clip(1 - blur, 0, 1) * lungs
    # A short hand-specified viridis-like palette, preserving the reference hue.
    palette = np.array([[68, 1, 84], [65, 68, 135], [42, 120, 142],
                        [34, 168, 132], [122, 209, 81], [253, 231, 37]], dtype=float)
    target, reconstruction = [], []
    for idx, (cx, cy) in enumerate([(.29, .43), (.70, .40), (.55, .58)]):
        peak = np.exp(-((xx - cx) / .21) ** 2 - ((yy - cy) / .25) ** 2)
        field = .05 + .40 * lungs + .35 * peak + .18 * lung_dark + .62 * (gray - blur) + .16 * gray
        field = np.clip(field / 1.04, 0, 1)
        field = np.clip(field / np.quantile(field, .995), 0, 1)
        for dest, delta in [(target, 0), (reconstruction, .015)]:
            f = np.clip(field + delta * np.sin(11 * xx + 7 * yy + idx), 0, 1)
            rgb = np.stack([np.interp(f, np.linspace(0, 1, len(palette)), palette[:, c])
                            for c in range(3)], axis=-1)
            dest.append(Image.fromarray(np.uint8(np.clip(rgb, 0, 255))))
    return target, reconstruction


def feature_plane(slide, name, im, x, y, w=37, h=64, skew=27):
    """One compact editable component: bitmap texture + native edges/side."""
    group = slide.shapes.add_group_shape()
    group.name = name
    scale = 6
    src = im.resize((round(w * scale), round(h * scale)), Image.Resampling.LANCZOS).convert('RGBA')
    tilted = src.transform((round(w * scale), round((h + skew) * scale)),
                           Image.Transform.AFFINE,
                           (1, 0, 0, skew / w, 1, -skew * scale),
                           Image.Resampling.BICUBIC)
    path(slide, name + ' thickness', [(x - 5, y + skew + 2), (x, y + skew),
         (x, y + h + skew), (x - 5, y + h + skew - 2)], '83718D', .7,
         '9DA6B5', close=True, target=group.shapes)
    picture(slide, name + ' illustrative field', tilted, x, y, w, h + skew, target=group.shapes)
    path(slide, name + ' outline', [(x, y + skew), (x + w, y), (x + w, y + h),
         (x, y + h + skew)], '56718B', .65, close=True, target=group.shapes)


def coordinate_grid(slide):
    group = slide.shapes.add_group_shape()
    group.name = 'Coordinate query grid'
    for idx in range(8):
        t = idx / 7
        line(slide, f'Grid column {idx + 1}', 1471 + 49 * t, 216,
             1467 + 49 * t, 264, 1.15, '77BAF4', target=group.shapes)
        line(slide, f'Grid row {idx + 1}', 1467, 223 + 41 * t,
             1520, 216 + 41 * t, 1.15, '77BAF4', target=group.shapes)
    for idx in range(7):
        line(slide, f'Grid top tick {idx + 1}', 1471 + 7 * idx, 221,
             1476 + 7 * idx, 215, 1.0, '77BAF4', target=group.shapes)


def snowflake(slide, x, y):
    group = slide.shapes.add_group_shape()
    group.name = 'Frozen teacher snowflake'
    for i in range(6):
        a = np.pi * i / 3
        ex, ey = np.cos(a), np.sin(a)
        line(slide, f'Snowflake spoke {i + 1}', x, y, x + 10 * ex, y + 10 * ey,
             1.6, '0876C7', target=group.shapes)
        for sign in (-1, 1):
            bx, by = x + 6 * ex, y + 6 * ey
            line(slide, f'Snowflake branch {i + 1} {sign}', bx, by,
                 bx + 1.5 * ex - sign * 3 * ey, by + 1.5 * ey + sign * 3 * ex,
                 1.4, '0876C7', target=group.shapes)


def build(output):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W * EMU, H * EMU
    prs.core_properties.title = 'Multi-depth slot construction and visual slot spatial consistency'
    prs.core_properties.subject = 'Editable reconstruction of Figure 3, with CXR images'
    prs.core_properties.comments = 'Feature fields are Python-generated illustrations, not model results.'
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    cxr = Image.open(ASSETS / 'cxr.png').convert('L')
    target_maps, reconstructed_maps = feature_fields(cxr)

    # Panel backgrounds and titles.
    rect(slide, 'Panel A background', 110, 110, 1071, 559, 'F6FBFE', '75BFFC', 1.4, 13, True)
    rect(slide, 'Panel B background', 1188, 110, 723, 559, 'F4FAF6', '77CEA0', 1.4, 13, True)
    label(slide, 'Panel A title', '(a)  Multi-depth slot construction',
          125, 120, 1020, 37, 28, True, align=PP_ALIGN.LEFT)
    label(slide, 'Panel B title', [
        {'text': '(b)  Visual slot spatial consistency '},
        {'text': '(training only)', 'color': '787878'}],
        1203, 120, 698, 37, 28, True, align=PP_ALIGN.LEFT)

    # Input document, actual CXR, and report excerpt.
    rect(slide, 'Image and report card', 130, 205, 136, 328, 'FEFEFF', '686F7E', 1.5, 9)
    label(slide, 'Input heading', 'Image\n+ Report', 143, 215, 111, 43, 20)
    photo(slide, 'Input CXR', cxr, 145, 266, 106, 113)
    line(slide, 'Report to image arrow', 199, 401, 199, 384, 1.7, arrow=True)
    path(slide, 'Report paper outline', [(253, 427), (253, 517), (143, 517), (143, 404), (247, 404)],
         'BEC8D2', 1.7)
    for i, (x2, y) in enumerate([(239, 413), (239, 419), (239, 425)]):
        line(slide, f'Report decoration line {i + 1}', 150, y, x2, y, 2.2, 'BAC6D1')
    label(slide, 'CXR report excerpt', 'Findings:\nNo focal\nconsolidation.\nNo acute ...',
          149, 435, 103, 72, 16.7, align=PP_ALIGN.LEFT, leading=1.07)
    line(slide, 'Input branch stem', 267, 330, 283, 330, 1.8)
    path(slide, 'Input to VLM arrow', [(283, 330), (283, 277), (314, 277)], arrow=True)
    path(slide, 'Input to vision encoder arrow', [(283, 330), (283, 513), (314, 513)], arrow=True)

    # Encoder modules and their independently editable labels.
    vlm = path(slide, 'VLM trapezoid', [(318, 208), (474, 243), (474, 309), (318, 342)],
               PINK, 2, 'F1CBE2', close=True)
    gradient(vlm, [(0, 'F7DDEC'), (100, 'E8B4D3')])
    label(slide, 'VLM name', 'VLM', 334, 250, 122, 29, 24, True)
    label(slide, 'VLM fusion label', '(after fusion)', 328, 279, 137, 27, 20)
    vision = path(slide, 'Vision encoder trapezoid', [(318, 451), (474, 486), (474, 551), (318, 582)],
                  BLUE, 2, 'B4DFF8', close=True)
    gradient(vision, [(0, 'D0EBFC'), (100, '9DD2F3')])
    label(slide, 'Vision encoder name', 'Vision Encoder', 323, 494, 148, 28, 21, True)
    label(slide, 'Vision encoder fusion label', '(before fusion)', 324, 520, 146, 27, 20)
    line(slide, 'VLM to depth strip', 479, 246, 499, 246, 1.7, arrow=True)
    line(slide, 'Vision encoder to depth strip', 479, 494, 499, 494, 1.7, arrow=True)

    rect(slide, 'Fused depth strip', 502, 220, 432, 44, 'FCF0F8', 'EFDCEE', 1, 8)
    rect(slide, 'Visual depth strip', 502, 472, 432, 41, 'EFF8FF', 'D4EAFE', 1, 8)
    for i, x in enumerate([514, 537, 560, 634, 654, 674, 750, 787, 865, 902]):
        token(slide, f'Fused layer strip token {i + 1}', x, 222, 14, 40, faint=True)
        token(slide, f'Visual layer strip token {i + 1}', x, 474, 14, 37, blue=True, faint=True)
    for i, x in enumerate([590, 705, 820]):
        label(slide, f'Fused omitted layers {i + 1}', '…', x, 224, 27, 32, 26, font='Times New Roman')
        label(slide, f'Visual omitted layers {i + 1}', '…', x, 475, 27, 32, 26, font='Times New Roman')

    for i, x in enumerate([545, 661, 775, 890], 1):
        math_label(slide, f'Fused layer depth l{i}', 'ℓ', str(i), x - 18, 154, 38, 36, 29, True)
        line(slide, f'Fused depth pointer {i}', x, 191, x, 216, 1.8, arrow=True, dash=True)
        line(slide, f'Fused layer to readout {i}', x, 268, x, 289, 1.7, arrow=True, dash=True)
        module(slide, f'Readout {i}', 'Readout', x - 44, 291, 88, 40, 19)
        line(slide, f'Readout to fused slot {i}', x, 333, x, 349, 1.7, arrow=True)
        token(slide, f'Fused slot {i}', x - 11, 352, 23, 38)
        math_label(slide, f'Visual layer depth k{i}', 'k', str(i), x - 18, 413, 38, 33, 28, True)
        line(slide, f'Visual depth pointer {i}', x, 449, x, 469, 1.7, arrow=True, dash=True)
        line(slide, f'Visual layer to pooling {i}', x, 515, x, 529, 1.7, arrow=True, dash=True)
        module(slide, f'Query pooling {i}', 'Query\npooling', x - 44, 532, 88, 50, 18.5)
        line(slide, f'Pooling to visual slot {i}', x, 585, x, 604, 1.7, arrow=True, dash=True)
        token(slide, f'Visual slot {i}', x - 11, 607, 23, 38, blue=True)

    # Concatenation and state card.
    path(slide, 'State assembly bracket', [(946, 245), (967, 245), (967, 562), (946, 562)], width=2)
    line(slide, 'Assembled slots to state', 967, 390, 1011, 390, 3, arrow=True)
    rect(slide, 'State card background', 1019, 205, 140, 383, 'FEFEFF', 'BEC8D8', 1.6, 13)
    label(slide, 'State title', [{'text': 'State ', 'bold': True},
          {'text': 'S', 'font': 'Times New Roman', 'italic': True, 'size': 31}],
          1032, 212, 113, 33, 26)
    label(slide, 'State dimensions', [{'text': '8 × '}, {'text': 'd', 'italic': True}],
          1035, 245, 103, 28, 23, font='Times New Roman')
    for i in range(4):
        token(slide, f'State fused slot {i + 1}', 1038, 288 + 35.5 * i, 37, 24)
        token(slide, f'State visual slot {i + 1}', 1038, 444 + 35 * i, 37, 24, blue=True)
    rect(slide, 'Fused state brace', 1088, 289, 14, 130, stroke=INK, width=2.4, kind=MSO_SHAPE.RIGHT_BRACE)
    rect(slide, 'Visual state brace', 1088, 445, 14, 129, stroke=INK, width=2.4, kind=MSO_SHAPE.RIGHT_BRACE)
    math_label(slide, 'Fused state symbol', 'S', 'f', 1108, 333, 40, 45, 34)
    math_label(slide, 'Visual state symbol', 'S', 'v', 1108, 488, 40, 45, 34)
    rect(slide, 'Panel transition arrow', 1167, 390, 47, 27, '4E698A', kind=MSO_SHAPE.RIGHT_ARROW)

    # Spatial consistency reconstruction branch.
    label(slide, 'Visual slots heading', 'Visual slots', 1218, 240, 116, 31, 20, True, align=PP_ALIGN.LEFT)
    math_label(slide, 'Spatial visual state symbol', 'S', 'v', 1333, 236, 34, 34, 28)
    rect(slide, 'Visual slots container', 1254, 274, 57, 152, 'FFFFFF', 'ACC6E2', 1.6, 10)
    for i in range(4):
        token(slide, f'Spatial input visual slot {i + 1}', 1264, 285 + 35.2 * i, 36, 24, blue=True)
    line(slide, 'Slots to spatial reconstruction', 1320, 352, 1403, 352, 2, arrow=True)
    label(slide, 'Coordinate queries heading', 'Coordinate queries', 1425, 183, 207, 30, 20)
    coordinate_grid(slide)
    label(slide, 'Coordinate pair', [{'text': '(x, y)', 'italic': True}], 1536, 215, 64, 35,
          23, font='Times New Roman')
    line(slide, 'Coordinates to spatial reconstruction', 1495, 274, 1495, 310, 2, arrow=True)
    module(slide, 'Spatial reconstruction', 'Spatial\nreconstruction', 1407, 316, 179, 71, 21)
    line(slide, 'Reconstruction to feature field', 1592, 352, 1663, 352, 2, arrow=True)
    label(slide, 'Reconstructed feature field heading', 'Reconstructed\nfeature field',
          1637, 240, 186, 53, 21, True)
    for i, im in enumerate(reconstructed_maps):
        feature_plane(slide, f'Reconstructed plane {i + 1}', im, 1674 + 22 * i, 302 + 2 * i)
    line(slide, 'Reconstructed features to alignment', 1721, 404, 1721, 471, 2, arrow=True)

    # Frozen vision teacher branch, using the same CXR and two real crops.
    label(slide, 'Image and crops heading', 'Image + crops', 1211, 480, 141, 32, 20, True, align=PP_ALIGN.LEFT)
    photo(slide, 'Teacher input CXR', cxr, 1212, 519, 76, 96)
    cw, ch = cxr.size
    for i, box in enumerate([(.11, .15, .47, .55), (.55, .38, .91, .78)]):
        crop = cxr.crop(tuple(round(v * (cw if j % 2 == 0 else ch)) for j, v in enumerate(box)))
        photo(slide, f'CXR crop {i + 1}', crop, 1299, 520 + 54 * i, 38, 43)
    line(slide, 'CXR crops to teacher', 1344, 563, 1371, 563, 1.9, arrow=True)
    teacher = path(slide, 'Frozen teacher trapezoid', [(1376, 508), (1482, 534), (1482, 599), (1376, 622)],
                   '3B8BDD', 1.6, 'CFE8F9', close=True)
    gradient(teacher, [(0, 'E3F2FD'), (100, 'ACD9F5')])
    label(slide, 'Frozen teacher label', 'Frozen\nvision\nteacher', 1382, 532, 76, 65, 18, leading=1.1)
    snowflake(slide, 1461, 583)
    line(slide, 'Teacher to target features', 1488, 563, 1513, 563, 1.9, arrow=True)
    label(slide, 'Target features heading', 'Target features', 1500, 476, 149, 29, 20, True)
    for i, im in enumerate(target_maps):
        feature_plane(slide, f'Target plane {i + 1}', im, 1528 + 22 * i, 516 + 2 * i)
    label(slide, 'Omitted target features', '…', 1611, 541, 28, 32, 26, font='Times New Roman')
    path(slide, 'Target features to alignment', [(1643, 563), (1721, 563), (1721, 530)],
         width=1.8, arrow=True)
    align = rect(slide, 'Alignment background', 1663, 475, 118, 51, 'F9FAFC', '535A67', 1.6, 8)
    gradient(align, [(0, 'FFFFFF'), (100, 'F2F5F9')])
    label(slide, 'Alignment label', 'Align & match', 1668, 483, 108, 34, 16)
    line(slide, 'Alignment to consistency loss', 1787, 501, 1815, 501, 1.8, arrow=True)
    rect(slide, 'Consistency loss background', 1819, 472, 78, 65, 'FFF3F9', 'D95599', 1.5, 10, True)
    label(slide, 'Spatial consistency loss formula', [
        {'text': 'ℒ', 'font': 'Latin Modern Math', 'size': 30},
        {'text': 'VSSC', 'font': 'Times New Roman', 'size': 26, 'baseline': -24000}],
        1825, 484, 66, 39, 30)
    slide.notes_slide.notes_text_frame.text = (
        'Reconstructed from source/fig3_v3.png with MRI replaced by a real PA CXR. '
        'The report is abridged from the source CXR report. '
        'The six colored feature planes are deterministic Python illustrations '
        'based on the CXR texture and analytic spatial functions. They are not '
        'learned attention, true teacher embeddings, or measured model results. '
        'See scripts/fig3_v3/README.md and assets/provenance.json. '
        'Text, shapes, tokens, connectors, braces, and grid are native editable '
        'PowerPoint objects; CXR/crops and feature textures are bitmap pictures.')
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(output)
    return output


def render(pptx, preview):
    soffice, pdftoppm = shutil.which('libreoffice'), shutil.which('pdftoppm')
    if not soffice or not pdftoppm:
        raise RuntimeError('Rendering requires LibreOffice and pdftoppm.')
    with tempfile.TemporaryDirectory(prefix='fig3_render_') as tmp:
        tmp = Path(tmp)
        subprocess.run([soffice, '-env:UserInstallation=' + (tmp / 'profile').as_uri(),
                        '--headless', '--convert-to', 'pdf', '--outdir', str(tmp), str(pptx)],
                       check=True, timeout=90, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        subprocess.run([pdftoppm, '-png', '-singlefile', '-scale-to-x', str(W * 2),
                        '-scale-to-y', str(H * 2), str(tmp / (pptx.stem + '.pdf')),
                        str(preview.with_suffix(''))], check=True, timeout=60)
    print(preview)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument('--render', action='store_true', help='Render final PPT via LibreOffice to PNG.')
    args = parser.parse_args()
    output = build(args.output.resolve())
    if args.render:
        render(output, output.with_suffix('.png'))
