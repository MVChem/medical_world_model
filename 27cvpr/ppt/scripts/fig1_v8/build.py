"""Extend the archived editable v7 figure with the native forecast readout.

Keep v7 and its medical thumbnails unchanged. All new diagram components are
editable PowerPoint shapes; PDF/SVG exports preserve vector artwork.
"""
from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile

import fitz
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "ppt/ppt/fig1_v7.pptx"
OUT = ROOT / "ppt/ppt/fig1_v8.pptx"
INK, BLUE, PINK = "171717", "9BCAE9", "E89EBA"
prs = Presentation(SOURCE)
assert len(prs.slides) == 1
slide = prs.slides[0]
prs.slide_width, prs.slide_height = Inches(1890 / 120), Inches(1040 / 120)
prs.core_properties.title = "MedWorld-JEPA — native forecast readout (fig1_v8)"
original = list(slide.shapes)


def u(value):
    return Inches(value / 120)


def named(name):
    found = [s for s in slide.shapes if s.name == name]
    assert len(found) == 1, (name, len(found))
    return found[0]


def box(x, y, w, h, fill, stroke=INK, name="Module", width=1):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, u(x), u(y), u(w), u(h))
    shape.name = name
    shape.shadow.inherit = False
    shape.adjustments[0] = .08
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.color.rgb = RGBColor.from_string(stroke)
    shape.line.width = Pt(width)
    return shape


def text(x, y, w, h, value, size=20, bold=False, color=INK,
         font="Comic Sans MS", left=False, name=None):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    shape.name = name or "Label: " + value.replace("\n", " / ")
    tf = shape.text_frame
    tf.word_wrap, tf.auto_size = False, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(value.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = line
        run.font.name, run.font.size = font, Pt(size * 72 / 120)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
    return shape


def arrow(points, name, color=INK, width=1.2):
    result = []
    for index, (a, b) in enumerate(zip(points, points[1:])):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(a[0]), u(a[1]), u(b[0]), u(b[1]))
        shape.name = name
        shape.line.color.rgb = RGBColor.from_string(color)
        shape.line.width = Pt(width)
        if index == len(points) - 2:
            end = OxmlElement("a:tailEnd")
            for key, value in {"type": "triangle", "w": "sm", "len": "sm"}.items():
                end.set(key, value)
            shape._element.spPr.get_or_add_ln().append(end)
        result.append(shape)
    return result


def replace_label(shape, value, size=None):
    first = shape.text_frame.paragraphs[0].runs[0]
    first.text = value
    for p in list(shape.text_frame.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for run in list(shape.text_frame.paragraphs[0].runs)[1:]:
        run._r.getparent().remove(run._r)
    if size is not None:
        first.font.size = Pt(size * 72 / 120)


# Replace only the internals of the two obsolete final-layer-query diagrams.
for lower, upper in ((41, 68), (126, 153)):
    for shape in original[lower:upper + 1]:
        shape._element.getparent().remove(shape._element)
for shape in list(slide.shapes):
    if shape.name == "Report tokens to VLM":
        shape._element.getparent().remove(shape._element)
for top, future in ((85, False), (469, True)):
    title = "Frozen state encoder" if future else "4 + 4 state encoder"
    text(565, top + 6, 348, 35, title, 25, True)
    text(909, top + 12, 78, 25, "Fixed" if future else "LoRA", 17)
    box(569, top + 51, 281, 76, "FFFCF3", "A42A8C", "Fusion path, future" if future else "Fusion path, current", .8)
    text(575, top + 57, 269, 27, "Qwen language fusion", 19, True)
    text(575, top + 89, 269, 24, "JEPA tokens + available text", 17, font="Arial")
    text(850, top + 48, 133, 23, "4 fusion slots", 17, True)
    for i in range(4):
        box(862 + i * 30, top + 83, 24, 27, PINK, "CC398F", f"{'Target' if future else 'Current'} fusion slot {i + 1}", .7)
    box(610, top + 147, 240, 64, "EDF7FE", "1686D1", "Native vision path, future" if future else "Native vision path, current", .8)
    text(616, top + 149, 228, 26, "Qwen native vision", 19, True)
    text(616, top + 178, 228, 24, "Same image, four depths", 17, font="Arial")
    image_label = text(561, top + 153, 43, 27, "", 20, font="Times New Roman",
                       name="Native state image label, future" if future else "Native state image label, current")
    paragraph = image_label.text_frame.paragraphs[0]
    for value, baseline, size in (("X", 0, 20), ("t+" if future else "t", -25000, 15)):
        run = paragraph.add_run()
        run.text = value
        run.font.name, run.font.size = "Times New Roman", Pt(size * 72 / 120)
        run.font.color.rgb = RGBColor.from_string("1686D1")
        run._r.get_or_add_rPr().set("baseline", str(baseline))
    arrow([(570, top + 184), (610, top + 184)],
          "Future image to native target vision" if future else "Current image to native state vision", "1686D1")
    text(850, top + 143, 133, 23, "4 visual slots", 17, True)
    for i in range(4):
        box(862 + i * 30, top + 179, 24, 27, BLUE, "1686D1", f"{'Target' if future else 'Current'} visual slot {i + 1}", .7)
    arrow([(850, top + 96), (862, top + 96)], f"Fusion to {'target' if future else 'current'} slots")
    arrow([(850, top + 192), (862, top + 192)], f"Vision to {'target' if future else 'current'} slots")
    arrow([(976, top + 96), (987, top + 96), (987, top + 145), (994, top + 145)], "Fusion readout to state")
    arrow([(976, top + 192), (987, top + 192), (987, top + 145), (994, top + 145)], "Visual readout to state")
    start_y = 676 if future else 288
    arrow([(496, start_y), (538, start_y), (538, top + 120), (569, top + 120)], "Available text to fusion path")

replace_label(named("Label: Report"), "Report / EHR", 19)
replace_label(named("Label: Same encoding pipeline at both times"), "Same 4 + 4 design; fixed Stage 1 target", 18)
for shape in slide.shapes:
    if shape.name == "Label: Slot hidden / states":
        replace_label(shape, "4 + 4", 17)
    if "output token" in shape.name:
        position = int(shape.name.rsplit(" ", 1)[-1].split("/")[0])
        if position > 4:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor.from_string(BLUE)
            shape.line.color.rgb = RGBColor.from_string("1686D1")

# Append a readout panel; no new edge is connected to the latent predictor.
box(2, 771, 1856, 266, "F2F8F0", "78A45B", "Native forecast readout panel", 1)
text(24, 778, 602, 36, "Native forecast readout (Stage 2)", 27, True, left=True)
box(22, 823, 526, 172, "FFF6E8", "84501F", "True current forecast evidence", .9)
text(33, 829, 505, 27, "True current evidence", 23, True, left=True)
photo = named("Current chest radiograph")
thumb = slide.shapes.add_picture(BytesIO(photo.image.blob), u(37), u(865), u(105), u(105))
thumb.name = "Native decoder current radiograph (same image)"
text(158, 863, 376, 33, "Current image (512 px)", 20, True, left=True)
text(158, 903, 376, 59, "Current report + source EHR\nRequested horizon h", 20, left=True)
box(624, 823, 777, 172, "D8EBC1", "557F38", "Native Qwen forecast decoder", 1.1)
text(638, 829, 585, 28, "Native Qwen forecast decoder", 24, True)
box(642, 872, 223, 84, "EDF7FE", "1686D1", "Native forecast vision encoder", .9)
text(650, 880, 207, 65, "Native vision\nencoder", 23, True)
box(943, 872, 439, 99, "FFFDF1", "557F38", "Native forecast language model", .9)
text(954, 877, 415, 29, "Qwen language model + LoRA", 22, True)
text(954, 914, 415, 45, "Image + current text + h\n+ 8 projected future slots", 19)
arrow([(548, 882), (642, 882)], "Current image to native forecast vision")
text(554, 852, 70, 27, "Xₜ", 19, font="Times New Roman")
arrow([(865, 907), (943, 907)], "Native forecast vision tokens to language model")
text(868, 865, 72, 40, "Image\ntokens", 15, font="Arial")
arrow([(548, 971), (910, 971), (910, 944), (943, 944)], "Current report EHR horizon to native forecast language model")
text(636, 972, 260, 19, "Current report / EHR / h", 15, font="Arial")
box(946, 785, 394, 31, "FBE7F3", "CC398F", "Predicted future slot projection", .8)
text(952, 785, 382, 31, "A(predicted future slots)", 19, True)
arrow([(1320, 816), (1320, 872)], "Projected future slots to native forecast language model", "A42A8C", 1.25)
arrow([(1843, 502), (1875, 502), (1875, 754), (1143, 754), (1143, 785)], "Predicted future slots to native forecast readout", "A42A8C", 1.25)
box(1480, 841, 351, 135, "FFF6E8", "A36432", "Future report and finding outputs", .9)
text(1491, 850, 329, 34, "Future outputs", 24, True)
text(1491, 893, 329, 58, "Report (token CE)\nFindings (masked BCE)", 21)
arrow([(1401, 916), (1480, 916)], "Native forecast decoder to future outputs")
text(28, 1005, 1800, 22,
     "Future observations: detached latent targets and supervised losses only.  Image-based Stage 1 withholds current reports.",
     18, font="Arial", left=True)

# Verify the preserved predictor contract and basic editable figure contents.
assert sum(s.name == "Current state to latent world model" for s in slide.shapes) == 4
assert sum(s.name == "Horizon to latent world model" for s in slide.shapes) == 1
assert sum(s.name == "Observed future target to latent loss (stop-grad)" for s in slide.shapes) == 1
assert sum(s.name == "Predicted future slots to native forecast readout" for s in slide.shapes) == 4
assert sum(s.name == "Current image to native state vision" for s in slide.shapes) == 1
assert sum(s.name == "Future image to native target vision" for s in slide.shapes) == 1
assert sum(s.name == "Projected future slots to native forecast language model" for s in slide.shapes) == 1
assert all(s.left >= 0 and s.top >= 0 and s.left + s.width <= prs.slide_width + u(2)
           and s.top + s.height <= prs.slide_height + u(2) for s in slide.shapes)
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)

with tempfile.TemporaryDirectory(prefix="fig1_v8_export_") as directory:
    work = Path(directory)
    result = subprocess.run([
        "libreoffice", f"-env:UserInstallation={(work / 'profile').as_uri()}",
        "--headless", "--convert-to", "pdf:impress_pdf_Export",
        "--outdir", str(work), str(OUT),
    ], check=True, capture_output=True, text=True)
    with fitz.open(work / "fig1_v8.pdf") as source:
        assert len(source) == 1
        page = source[0]
        clip = fitz.Rect(0, 0, page.rect.width, page.rect.height)
        with fitz.open() as exported:
            target = exported.new_page(width=clip.width, height=clip.height)
            target.show_pdf_page(target.rect, source, 0, clip=clip)
            exported.save(OUT.with_suffix(".pdf"), garbage=4, deflate=True)
            OUT.with_suffix(".svg").write_text(target.get_svg_image(text_as_path=True))
            target.get_pixmap(matrix=fitz.Matrix(4200 / target.rect.width, 4200 / target.rect.width)).save(OUT.with_suffix(".png"))
shutil.copy2(OUT.with_suffix(".pdf"), ROOT / "imgs/fig1.pdf")
shutil.copy2(OUT.with_suffix(".png"), ROOT / "imgs/fig1.png")
print(OUT)
print(ROOT / "imgs/fig1.pdf")
