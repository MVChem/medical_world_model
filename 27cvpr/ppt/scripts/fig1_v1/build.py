from pathlib import Path
from copy import deepcopy
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

HERE = Path(__file__).resolve().parent
ASSETS = HERE / 'assets'
OUT = HERE.parents[1] / 'ppt'
SCALE = 914400 / 100
prs = Presentation()
prs.slide_width = int(2048*SCALE)
prs.slide_height = int(762*SCALE)
slide = prs.slides.add_slide(prs.slide_layouts[6])
INK, GRAY = '182234', '657183'

def el(tag, **attrs):
    e = OxmlElement(tag)
    for k,v in attrs.items(): e.set(k,str(v))
    return e

def shape(name, kind, x,y,w,h, fill, stroke=None, width=1.5, target=None):
    s=(target or slide.shapes).add_shape(kind,int(x*SCALE),int(y*SCALE),int(w*SCALE),int(h*SCALE))
    s.name=name
    for child in list(s._element):
        if child.tag.split('}')[-1]=='style': s._element.remove(child)
    s._element.spPr.append(el('a:effectLst'))
    if fill: s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(fill)
    else: s.fill.background()
    if stroke: s.line.color.rgb=RGBColor.from_string(stroke); s.line.width=Pt(width*.72)
    else: s.line.fill.background()
    return s

def text(name,content,x,y,w,h,size=25,color=INK,align=PP_ALIGN.CENTER,font='DejaVu Sans',italic=False):
    s=slide.shapes.add_textbox(int(x*SCALE),int(y*SCALE),int(w*SCALE),int(h*SCALE)); s.name=name
    tf=s.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=align
        p.space_before=Pt(0); p.space_after=Pt(0)
        r=p.add_run(); r.text=line; r.font.name=font; r.font.size=Pt(size*.72); r.font.italic=italic; r.font.color.rgb=RGBColor.from_string(color)
    return s

def path(name, cmds, color=None, width=2, dash=False, arrow=False, fill=None, target=None):
    xs=[v for _,c in cmds for v in c[0::2]]; ys=[v for _,c in cmds for v in c[1::2]]
    x0,y0=min(xs),min(ys); w,h=max(1,max(xs)-x0),max(1,max(ys)-y0)
    s=shape(name,MSO_SHAPE.RECTANGLE,x0,y0,w,h,fill,color,width,target)
    sp=s._element.spPr
    sp.remove(sp.prstGeom)
    geom=el('a:custGeom')
    for tag in ['avLst','gdLst','ahLst','cxnLst']: geom.append(el('a:'+tag))
    geom.append(el('a:rect',l='0',t='0',r='r',b='b'))
    pl=el('a:pathLst'); p=el('a:path',w=w,h=h)
    for cmd,coords in cmds:
        e=el('a:'+{'M':'moveTo','L':'lnTo','C':'cubicBezTo','Z':'close'}[cmd])
        for i in range(0,len(coords),2): e.append(el('a:pt',x=coords[i]-x0,y=coords[i+1]-y0))
        p.append(e)
    pl.append(p); geom.append(pl); sp.insert(1,geom)
    if dash: sp.get_or_add_ln().append(el('a:prstDash',val='dash'))
    if arrow: sp.get_or_add_ln().append(el('a:tailEnd',type='triangle',w='sm',len='sm'))
    return s

def curve(name,start,segments,color,width=2,dash=False,arrow=False):
    return path(name,[('M',start)]+[('C',s) for s in segments],color,width,dash,arrow)

def gradient(s,stops):
    sp=s._element.spPr
    for child in list(sp):
        if child.tag.split('}')[-1] in ['solidFill','noFill','gradFill']: sp.remove(child)
    g=el('a:gradFill',rotWithShape='1'); lst=el('a:gsLst')
    for pos,col in stops:
        stop=el('a:gs',pos=pos); stop.append(el('a:srgbClr',val=col)); lst.append(stop)
    g.append(lst); g.append(el('a:lin',ang=0,scaled='1')); sp.insert(2,g)

# Background and trajectories: editable native Bezier paths.
path('Latent patient space — cloud',[('M',[509,367]),('C',[504,291,592,260,701,232]),('C',[799,208,797,144,904,103]),('C',[985,70,1094,56,1170,78]),('C',[1245,91,1251,147,1334,174]),('C',[1392,194,1425,185,1474,234]),('C',[1521,283,1545,354,1545,426]),('C',[1548,526,1483,593,1382,613]),('C',[1302,636,1193,623,1100,625]),('C',[965,630,891,632,835,571]),('C',[789,518,775,510,717,496]),('C',[612,482,503,481,509,367]),('Z',[])],fill='F2F7FD')
trajectories=[('Upper blue',[630,317],[[683,290,716,296,770,297],[816,286,839,220,870,211],[921,198,969,267,1056,262],[1134,261,1170,173,1240,180],[1274,190,1308,246,1369,250]],'B8D2F2',[(630,317),(770,297),(870,211),(1056,262),(1240,180),(1369,250)]),('Upper branch',[952,205],[[972,183,977,174,990,172],[1051,173,1084,115,1138,115]],'C5D2E6',[(990,172),(1138,115)]),('Upper pink',[1205,304],[[1280,295,1286,334,1431,304]],'F2C6D9',[(1205,304),(1431,304)]),('Middle',[741,448],[[829,399,843,510,932,447],[978,426,1009,399,1046,409],[1113,452,1172,480,1234,477]],'B8D0EC',[(741,448),(932,447),(1046,409),(1234,477)]),('Middle pink',[1234,477],[[1283,470,1295,448,1334,444]],'F1C8DA',[(1334,444)]),('Bottom pink',[877,543],[[943,513,986,561,1036,561],[1084,561,1112,531,1151,523]],'F2C8DB',[(877,543),(1036,561)]),('Bottom gray',[1151,523],[[1219,550,1251,588,1348,558]],'CBD4DF',[(1151,523),(1348,558)])]
for name,start,segs,col,nodes in trajectories:
    curve(name+' trajectory',start,segs,col,2,True,name=='Upper pink')
    for i,(x,y) in enumerate(nodes): shape(f'{name} node {i+1}',MSO_SHAPE.OVAL,x-12,y-12,24,24,col)

ref=Image.open(ASSETS/'reference.png')
for name,box in [('xray',(71,177,279,355)),('future_xray',(1664,191,1771,294))]:
    ref.crop(box).save(ASSETS/(name+'.png'))
    x,y,x2,y2=box
    pic=slide.shapes.add_picture(str(ASSETS/(name+'.png')),int(x*SCALE),int(y*SCALE),int((x2-x)*SCALE),int((y2-y)*SCALE)); pic.name=name
text('X-ray label','X-ray',70,359,210,36,27)
shape('Clinical report card',MSO_SHAPE.ROUNDED_RECTANGLE,59,417,240,139,'FCF3F6','A7A8AC',2)
text('Clinical findings','Findings:\nMild cardiomegaly.\nNo focal consolidation.\n…',78,433,216,114,19,align=PP_ALIGN.LEFT)
text('Clinical report label','Clinical report',58,561,242,38,26)
curve('Image input',[286,262],[[310,262,324,262,330,262],[389,261,332,373,437,373],[466,373,481,373,505,373]],'3585C6',3,arrow=True)
curve('Report input',[299,483],[[368,483,353,478,367,417],[380,365,407,373,439,373]],'DA6697',3)
curve('Temporal evolution h',[633,378],[[856,315,1184,324,1413,378]],'606B7B',3.2,arrow=True)
text('Horizon h','h',986,286,54,55,44,'121212',font='STIX Two Math',italic=True)
for x,suffix in [(587,'t'),(1472,'t+h')]:
    s=shape('Patient state '+suffix,MSO_SHAPE.OVAL,x-33,350,66,64,'FFFFFF','718BBC',3)
    gradient(s,[(0,'83B8F1'),(49000,'E7F0FE'),(51000,'FCE6EE'),(100000,'E386AE')])
    t=text('State label '+suffix,'S',x-44,420,100,54,40,'111111',font='STIX Two Math',italic=True)
    r=t.text_frame.paragraphs[0].add_run(); r.text=suffix; r.font.name='STIX Two Math'; r.font.size=Pt(23); r.font.italic=True; r._r.get_or_add_rPr().set('baseline','-25000')
text('State modalities','(visual + clinical)',488,467,199,35,21,GRAY)
text('Latent space label','Latent patient space',904,635,399,51,33,GRAY)
for name,start,segs in [('Recognition',[517,439],[[456,455,438,495,432,535]]),('VQA',[580,494],[[579,511,577,521,577,537]]),('Spatial',[636,439],[[693,459,711,496,720,537]]),('Future findings',[1514,382],[[1590,380,1568,289,1651,258]]),('Disease progression',[1514,382],[[1584,382,1638,383,1688,383]]),('Future report',[1514,382],[[1590,392,1569,483,1651,517]])]:
    curve(name+' task arrow',start,segs,GRAY,1.6,True,True)

def tile(name,x,y,w=80,h=80):
    g=slide.shapes.add_group_shape(); g.name=name+' icon'
    shape(name+' background',MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,'EDF5FD','D4E0ED',1.8,g.shapes)
    return g

def lungs(name,x,y,heat=False):
    g=tile(name,x,y)
    for side,coords in [('left',[x+35,y+18,x+17,y+10,x+8,y+49,x+14,y+62,x+25,y+60,x+33,y+56,x+35,y+48]),('right',[x+45,y+18,x+61,y+10,x+70,y+49,x+65,y+62,x+54,y+60,x+46,y+56,x+45,y+48])]:
        path(name+' '+side,[('M',coords[:2]),('C',coords[2:8]),('C',coords[8:14]),('Z',[])],'7AA9E1',2,fill='A5CDF5',target=g.shapes)
    shape(name+' trachea',MSO_SHAPE.ROUNDED_RECTANGLE,x+38,y+11,5,22,'7AA9E1',target=g.shapes)
    if heat:
        for i,(w,h,col) in enumerate([(25,39,'B5DBA7'),(19,31,'F4E681'),(13,23,'F6BB73'),(7,15,'EF8C73')]):
            shape(name+' heat '+str(i),MSO_SHAPE.OVAL,x+55-w/2,y+43-h/2,w,h,col,target=g.shapes)
    return g
lungs('Recognition',391,546)
lungs('Spatial understanding',681,546,True)
g=tile('VQA',535,546)
shape('VQA bubble',MSO_SHAPE.OVAL,550,562,50,44,'FFFFFF','7AA9E1',3,g.shapes)
path('VQA bubble tail',[('M',[557,594]),('L',[550,609]),('L',[567,604])],'7AA9E1',3,fill='FFFFFF',target=g.shapes)
text('VQA question mark','?',558,565,35,44,33,'6D9EE3')
text('Recognition label','Recognition',351,636,161,39,25)
text('VQA label','VQA',532,636,89,39,25)
text('Spatial understanding label','Spatial\nunderstanding',636,636,171,70,25)
text('Future findings label','Future findings\n(e.g., increased opacity)',1786,211,260,70,20,align=PP_ALIGN.LEFT)
g=tile('Disease progression',1702,345,103,80)
curve_obj=path('Chart axes',[('M',[1711,355]),('L',[1711,416]),('L',[1795,416])],'BD829C',1.7,target=g.shapes)
pts=[(1719,406),(1736,390),(1753,395),(1769,379),(1788,363)]
path('Progression chart line',[('M',list(pts[0]))]+[('L',list(p)) for p in pts[1:]],'CB5885',2.6,target=g.shapes)
for i,(x,y) in enumerate(pts): shape('Progression marker '+str(i),MSO_SHAPE.OVAL,x-3,y-3,6,6,'CB5885',target=g.shapes)
text('Disease progression label','Disease progression',1823,370,223,43,23,align=PP_ALIGN.LEFT)
g=tile('Future report',1665,479,85,88)
path('Report page',[('M',[1687,496]),('L',[1714,496]),('L',[1727,509]),('L',[1727,551]),('L',[1687,551]),('Z',[])],'6A9FE7',3,fill='FFFFFF',target=g.shapes)
path('Report page fold',[('M',[1714,496]),('L',[1714,509]),('L',[1727,509])],'6A9FE7',3,target=g.shapes)
for i in range(3): path('Report text line '+str(i),[('M',[1694,517+9*i]),('L',[1720,517+9*i])],'6A9FE7',3,target=g.shapes)
text('Future report label','Future report\ngeneration',1774,499,256,66,24,align=PP_ALIGN.LEFT)
OUT.mkdir(parents=True,exist_ok=True)
prs.save(OUT/'fig1_v1.pptx')
print(OUT/'fig1_v1.pptx')
