from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parents[2]
SOURCE = BASE / "source" / "fig1_v1.png"
OUTPUT = BASE / "ppt" / "fig1_v1.pptx"

# The reference is 1536 x 1024. Coordinates below stay in source pixels.
PX_PER_INCH = 102.4
W, H = 1536, 1024

WHITE = "FFFFFF"
INK = "111111"
GRAY = "353535"
LIGHT_GRAY = "C8C8C8"
CARD = "FCFCFC"
BLUE = "174A91"
BLUE_DARK = "0C2A60"
BLUE_FILL = "F0F4FB"
BLUE_TOKEN = "9DBCEB"
PURPLE = "542482"
PURPLE_FILL = "F3EEF8"
PURPLE_TOKEN = "C7ADDD"
ORANGE = "C94308"
ORANGE_FILL = "FFF2EA"
ORANGE_TOKEN = "F4B18A"
GREEN = "216A21"
GREEN_DARK = "185918"
GREEN_FILL = "F1F7EF"
GREEN_TOKEN = "AFCDAA"
FONT = "Arial"


def emu(px):
    return Inches(px / PX_PER_INCH)


def rgb(value):
    return RGBColor.from_string(value)


def no_shadow(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def set_dash(shape, value="dash"):
    line = shape._element.spPr.get_or_add_ln()
    for child in list(line):
        if child.tag.endswith("prstDash"):
            line.remove(child)
    dash = OxmlElement("a:prstDash")
    dash.set("val", value)
    line.append(dash)


def set_end_arrow(shape, value="triangle"):
    line = shape._element.spPr.get_or_add_ln()
    for child in list(line):
        if child.tag.endswith("tailEnd"):
            line.remove(child)
    arrow = OxmlElement("a:tailEnd")
    arrow.set("type", value)
    arrow.set("w", "sm")
    arrow.set("len", "sm")
    line.append(arrow)


def box(slide, x, y, w, h, fill=WHITE, line=GRAY, width=1.1, rounded=True, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, emu(x), emu(y), emu(w), emu(h))
    if name:
        shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(width)
    if rounded:
        shape.adjustments[0] = 0.08
    return no_shadow(shape)


def text(
    slide,
    x,
    y,
    w,
    h,
    value,
    size=11.5,
    color=INK,
    bold=False,
    italic=False,
    align=PP_ALIGN.CENTER,
    valign=MSO_ANCHOR.MIDDLE,
    margin=1,
    name=None,
):
    shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(h))
    if name:
        shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = emu(margin)
    tf.margin_right = emu(margin)
    tf.margin_top = emu(margin)
    tf.margin_bottom = emu(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = 0.95
    run = p.add_run()
    run.text = value
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    return no_shadow(shape)


def line(slide, x1, y1, x2, y2, color=GRAY, width=1.25, arrow=False, dashed=False, name=None):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, emu(x1), emu(y1), emu(x2), emu(y2)
    )
    if name:
        shape.name = name
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        set_end_arrow(shape)
    if dashed:
        set_dash(shape)
    return no_shadow(shape)


def path(slide, points, color=GRAY, width=1.25, arrow=True, dashed=False, name="path"):
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        line(
            slide,
            x1,
            y1,
            x2,
            y2,
            color,
            width,
            arrow=arrow and i == len(points) - 2,
            dashed=dashed,
            name=f"{name} {i + 1}",
        )


def polygon(slide, points, fill, line_color, width=0.8, name=None):
    min_x = min(p[0] for p in points)
    min_y = min(p[1] for p in points)
    local = [(p[0] - min_x, p[1] - min_y) for p in points]
    builder = slide.shapes.build_freeform(
        local[0][0], local[0][1], scale=Inches(1) / PX_PER_INCH
    )
    builder.add_line_segments(local[1:], close=True)
    shape = builder.convert_to_shape(emu(min_x), emu(min_y))
    if name:
        shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(width)
    return no_shadow(shape)


def ellipse(slide, x, y, w, h, fill=None, line_color=GRAY, width=1.0, dashed=False, name=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, emu(x), emu(y), emu(w), emu(h))
    if name:
        shape.name = name
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line_color)
    shape.line.width = Pt(width)
    if dashed:
        set_dash(shape)
    return no_shadow(shape)


def crop_blob(image, crop):
    stream = BytesIO()
    image.crop(crop).save(stream, format="PNG")
    return stream.getvalue()


def picture(slide, blob, x, y, w, h, name):
    shape = slide.shapes.add_picture(BytesIO(blob), emu(x), emu(y), emu(w), emu(h))
    shape.name = name
    shape.line.color.rgb = rgb("E3E3E3")
    shape.line.width = Pt(0.45)
    return no_shadow(shape)


def token_bar(
    slide,
    x,
    y,
    w,
    h,
    color,
    token_fill,
    count=5,
    ellipsis=True,
    fill=WHITE,
    name="tokens",
):
    box(slide, x, y, w, h, fill, color, 1.0, True, name)
    size = min(h - 14, 18)
    gap = 8
    start = x + 9
    for idx in range(count):
        sx = start + idx * (size + gap)
        if sx + size > x + w - (30 if ellipsis else 8):
            break
        box(
            slide,
            sx,
            y + (h - size) / 2,
            size,
            size,
            token_fill,
            color,
            0.7,
            False,
            f"{name} token {idx + 1}",
        )
    if ellipsis:
        text(slide, x + w - 31, y + 1, 25, h - 2, "···", 12, color, True, name=f"{name} ellipsis")


def cube(slide, x, y, size, depth, front, top, side, edge, name):
    polygon(
        slide,
        [(x, y), (x + depth, y - depth), (x + size + depth, y - depth), (x + size, y)],
        top,
        edge,
        0.55,
        f"{name} top",
    )
    polygon(
        slide,
        [(x + size, y), (x + size + depth, y - depth), (x + size + depth, y + size - depth), (x + size, y + size)],
        side,
        edge,
        0.55,
        f"{name} side",
    )
    box(slide, x, y, size, size, front, edge, 0.55, False, f"{name} front")


def volume(slide, x, y, scale=1.0, green=False, name="volume"):
    if green:
        colors = ("679E64", "8DB78A", "3D793D", "EAF4E8")
    else:
        colors = ("4C79C6", "769DE0", "27539C", "EAF1FD")
    front, top, side, edge = colors
    size, depth = 25 * scale, 8 * scale
    spots = [(44, 8), (70, 16), (22, 31), (48, 37), (76, 32), (0, 57), (29, 56), (58, 60), (87, 51)]
    for idx, (dx, dy) in enumerate(spots):
        cube(slide, x + dx * scale, y + dy * scale, size, depth, front, top, side, edge, f"{name} cube {idx + 1}")


def map_stack(slide, x, y, w, h, color, fills, name):
    layer_h = max(11, h * 0.25)
    slant = w * 0.20
    for idx in range(3):
        yy = y + h - layer_h - idx * (h * 0.23)
        xx = x + idx * (w * 0.05)
        polygon(
            slide,
            [(xx, yy + layer_h), (xx + slant, yy), (xx + w, yy), (xx + w - slant, yy + layer_h)],
            fills[min(idx, len(fills) - 1)],
            color,
            0.8,
            f"{name} layer {idx + 1}",
        )
        line(slide, xx + slant + 5, yy + layer_h * 0.25, xx + w - 7, yy + layer_h * 0.25, color, 0.45, name=f"{name} detail {idx + 1}")


def wire_cube(slide, x, y, size, depth, color, name):
    box(slide, x, y, size, size, None, color, 0.85, False, f"{name} front")
    box(slide, x + depth, y - depth, size, size, None, color, 0.75, False, f"{name} back")
    for idx, (x1, y1, x2, y2) in enumerate(
        [
            (x, y, x + depth, y - depth),
            (x + size, y, x + size + depth, y - depth),
            (x, y + size, x + depth, y + size - depth),
            (x + size, y + size, x + size + depth, y + size - depth),
        ]
    ):
        line(slide, x1, y1, x2, y2, color, 0.7, name=f"{name} corner {idx + 1}")
    for idx in (1, 2):
        off = size * idx / 3
        line(slide, x + off, y, x + off, y + size, color, 0.45, name=f"{name} vertical {idx}")
        line(slide, x, y + off, x + size, y + off, color, 0.45, name=f"{name} horizontal {idx}")
        line(slide, x + depth + off, y - depth, x + depth + off, y + size - depth, color, 0.4, name=f"{name} back vertical {idx}")
        line(slide, x + depth, y - depth + off, x + size + depth, y - depth + off, color, 0.4, name=f"{name} back horizontal {idx}")


def tokenizer_icon(slide, x, y, name):
    size = 9
    for row in range(3):
        for col in range(3):
            box(slide, x + col * 10, y + row * 10, size, size, "315E9E", WHITE, 0.35, False, f"{name} 2D {row}-{col}")
            box(slide, x + 51 + col * 10, y + row * 10, size, size, "315E9E", WHITE, 0.35, False, f"{name} 3D front {row}-{col}")
    for idx in range(3):
        line(slide, x + 51 + idx * 10, y, x + 58 + idx * 10, y - 6, BLUE, 0.6, name=f"{name} depth top {idx}")
        line(slide, x + 80, y + idx * 10, x + 87, y - 6 + idx * 10, BLUE, 0.6, name=f"{name} depth side {idx}")


def report_icons(slide, x, y, name):
    box(slide, x, y, 33, 48, WHITE, GRAY, 0.9, False, f"{name} document")
    for idx in range(4):
        line(slide, x + 7, y + 10 + idx * 7, x + 27, y + 10 + idx * 7, GRAY, 0.7, name=f"{name} document line {idx + 1}")
    box(slide, x + 58, y + 10, 52, 29, WHITE, GRAY, 0.8, False, f"{name} waveform frame")
    pts = [(x + 63, y + 26), (x + 68, y + 25), (x + 72, y + 18), (x + 77, y + 32), (x + 82, y + 15), (x + 87, y + 27), (x + 93, y + 20), (x + 100, y + 25), (x + 105, y + 24)]
    path(slide, pts, GRAY, 0.75, False, name=f"{name} waveform")
    box(slide, x + 110, y + 10, 31, 29, WHITE, GRAY, 0.8, False, f"{name} context grid")
    line(slide, x + 125, y + 10, x + 125, y + 39, GRAY, 0.6, name=f"{name} context vertical")
    line(slide, x + 110, y + 24, x + 141, y + 24, GRAY, 0.6, name=f"{name} context horizontal")
    ellipse(slide, x + 115, y + 16, 3, 3, GRAY, GRAY, 0.3, name=f"{name} dot 1")
    ellipse(slide, x + 130, y + 29, 3, 3, GRAY, GRAY, 0.3, name=f"{name} dot 2")


def dynamics_icon(slide, x, y, name):
    ellipse(slide, x, y, 104, 104, None, ORANGE, 1.0, True, f"{name} outer orbit")
    ellipse(slide, x + 15, y + 19, 75, 64, None, "E9A67C", 0.7, name=f"{name} inner orbit")
    cx, cy = x + 52, y + 52
    nodes = [(cx, cy), (cx, y + 18), (x + 19, y + 50), (x + 86, y + 50), (cx, y + 85)]
    for idx, (nx, ny) in enumerate(nodes[1:]):
        line(slide, cx, cy, nx, ny, ORANGE, 0.75, name=f"{name} spoke {idx + 1}")
    for idx, (nx, ny) in enumerate(nodes):
        ellipse(slide, nx - 8, ny - 8, 16, 16, "E98651" if idx else "C85A1B", ORANGE, 0.65, name=f"{name} node {idx + 1}")


def gmm_fan(slide, x, y, w, h, name):
    box(slide, x, y, w, h, WHITE, ORANGE, 1.0, True, name)
    origin = (x + 18, y + h / 2)
    ellipse(slide, origin[0] - 4, origin[1] - 4, 8, 8, "E88650", ORANGE, 0.5, name=f"{name} origin")
    row_y = [y + 15, y + 43, y + 71]
    for ridx, yy in enumerate(row_y):
        line(slide, origin[0], origin[1], x + 48, yy + 9, ORANGE, 0.65, dashed=True, name=f"{name} branch {ridx + 1}")
        for idx in range(4 if ridx < 2 else 3):
            box(slide, x + 57 + idx * 29, yy, 19, 18, ORANGE_FILL, ORANGE, 0.65, False, f"{name} row {ridx + 1} token {idx + 1}")
        text(slide, x + 177, yy - 3, 21, 24, "···", 12, ORANGE, True, name=f"{name} row {ridx + 1} ellipsis")
    text(slide, x + w - 31, y + 10, 22, 23, "A", 12, INK, True, italic=True, name=f"{name} A")
    text(slide, x + w - 31, y + 38, 22, 23, "G", 12, INK, True, italic=True, name=f"{name} G")
    text(slide, x + w - 31, y + 66, 22, 23, "M", 12, INK, True, italic=True, name=f"{name} M")


def loss_label(slide, x, y, w, symbol, label, name):
    shape = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(32))
    shape.name = name
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = emu(1)
    tf.margin_top = tf.margin_bottom = emu(1)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.space_before = p.space_after = Pt(0)
    r = p.add_run()
    r.text = "L"
    r.font.name = "Times New Roman"
    r.font.size = Pt(16.5)
    r.font.italic = True
    r.font.color.rgb = rgb(INK)
    sub = p.add_run()
    sub.text = symbol
    sub.font.name = "Times New Roman"
    sub.font.size = Pt(10.5)
    sub.font.italic = True
    sub.font.color.rgb = rgb(INK)
    sub._r.get_or_add_rPr().set("baseline", "-25000")
    rest = p.add_run()
    rest.text = f"  ·  {label}"
    rest.font.name = FONT
    rest.font.size = Pt(11.2)
    rest.font.bold = True
    rest.font.color.rgb = rgb(INK)
    return no_shadow(shape)


def draw_top_connectors(slide):
    line(slide, 292, 161, 336, 161, GRAY, 1.35, True, name="Studies to tokenizers")
    line(slide, 454, 161, 491, 161, GRAY, 1.35, True, name="Tokenizers to online JEPA")
    path(slide, [(686, 171), (706, 171), (706, 108), (741, 108)], BLUE, 1.35, True, name="JEPA to dense")
    line(slide, 686, 171, 741, 171, BLUE, 1.35, True, name="JEPA to anatomy")
    path(slide, [(686, 171), (706, 171), (706, 252), (733, 252)], BLUE, 1.35, True, name="JEPA to global")
    line(slide, 677, 257, 677, 324, BLUE, 1.35, True, name="JEPA to fusion")
    line(slide, 784, 272, 784, 324, BLUE, 1.35, True, name="Global state to fusion")

    path(slide, [(199, 340), (220, 340), (220, 308), (249, 308)], PURPLE, 1.45, True, name="Report to text encoder")
    path(slide, [(199, 340), (220, 340), (220, 377), (606, 377)], PURPLE, 1.45, True, name="Clinical context to fusion")
    line(slide, 371, 309, 427, 309, PURPLE, 1.45, True, name="Text encoder to text tokens")
    path(slide, [(562, 312), (580, 312), (580, 346), (606, 346)], PURPLE, 1.45, True, name="Text tokens to fusion")
    line(slide, 810, 359, 850, 359, PURPLE, 1.45, True, name="Fusion to multimodal tokens")
    line(slide, 1012, 359, 1054, 359, PURPLE, 1.45, True, name="Current state to dynamics")

    line(slide, 1106, 136, 1106, 183, ORANGE, 1.35, True, name="Horizon to dynamics")
    line(slide, 1207, 136, 1207, 183, ORANGE, 1.35, True, name="Intervention to dynamics")
    line(slide, 1256, 205, 1294, 205, ORANGE, 1.35, True, name="Dynamics to transition")
    line(slide, 1256, 322, 1294, 322, ORANGE, 1.35, True, name="Dynamics to future distribution")


def draw_top(slide, scans):
    line(slide, 1035, 21, 1035, 461, LIGHT_GRAY, 0.9, name="Panel divider")
    draw_top_connectors(slide)

    # Panel (a), visual lane.
    box(slide, 18, 85, 274, 140, WHITE, GRAY, 1.1, True, "Current studies card")
    text(slide, 34, 98, 242, 28, "Current study · XR · CT · MRI", 11.2, INK, True, name="Current studies label")
    picture(slide, scans[0], 29, 134, 66, 74, "Current XR")
    picture(slide, scans[1], 106, 134, 69, 74, "Current CT")
    picture(slide, scans[2], 184, 134, 69, 74, "Current MRI")
    text(slide, 259, 145, 26, 45, "···", 13.5, INK, True, name="Current studies ellipsis")

    box(slide, 338, 105, 116, 121, BLUE_FILL, BLUE, 1.1, True, "Tokenizer module")
    text(slide, 348, 119, 96, 44, "2D / 3D\ntokenizers", 11.5, INK, True, name="Tokenizer label")
    tokenizer_icon(slide, 354, 177, "Tokenizer icon")

    box(slide, 495, 87, 191, 171, BLUE_FILL, BLUE, 1.25, True, "Online volume JEPA")
    text(slide, 509, 99, 162, 44, "TRAINABLE ·\nOnline volume JEPA", 11.5, INK, True, name="Online volume JEPA label")
    volume(slide, 525, 149, 1.0, False, "Online volume")

    box(slide, 742, 74, 99, 68, BLUE_FILL, BLUE, 1.0, True, "Dense-state glyph card")
    map_stack(slide, 752, 84, 78, 49, BLUE, ["A9C3ED", "6F97D7", "3D6DBD"], "Dense-state maps")
    box(slide, 862, 82, 139, 40, BLUE_FILL, BLUE, 1.0, True, "Dense-state label card")
    text(slide, 869, 83, 125, 38, "Dₜ · Dense", 12.3, INK, True, name="Dense-state label")

    box(slide, 743, 152, 99, 70, BLUE_FILL, BLUE, 1.0, True, "Anatomy-state glyph card")
    wire_cube(slide, 762, 169, 44, 11, BLUE_DARK, "Anatomy-state grid")
    box(slide, 863, 156, 138, 40, BLUE_FILL, BLUE, 1.0, True, "Anatomy-state label card")
    text(slide, 870, 157, 124, 38, "Aₜ · Anatomy", 12.3, INK, True, name="Anatomy-state label")

    token_bar(slide, 735, 233, 114, 39, BLUE, BLUE_TOKEN, 4, True, WHITE, "Global-state tokens")
    box(slide, 870, 232, 125, 40, BLUE_FILL, BLUE, 1.0, True, "Global-state label card")
    text(slide, 876, 233, 113, 38, "Gₜ · Global", 12.3, INK, True, name="Global-state label")

    # Panel (a), language lane and multimodal state.
    box(slide, 15, 282, 184, 121, WHITE, GRAY, 1.0, True, "Clinical context card")
    text(slide, 29, 291, 157, 46, "Report + history /\nclinical context", 11.3, INK, True, name="Clinical context label")
    report_icons(slide, 39, 342, "Clinical context icons")
    box(slide, 250, 278, 121, 63, PURPLE_FILL, PURPLE, 1.1, True, "Medical text encoder")
    text(slide, 262, 283, 97, 52, "Medical\ntext encoder", 11.2, INK, True, name="Medical text encoder label")
    token_bar(slide, 428, 298, 134, 29, PURPLE, PURPLE_TOKEN, 4, True, WHITE, "Report tokens")
    box(slide, 607, 325, 203, 69, PURPLE_FILL, PURPLE, 1.2, True, "Cross-modal fusion")
    text(slide, 619, 336, 180, 46, "Cross-modal fusion", 11.4, INK, True, name="Cross-modal fusion label")
    token_bar(slide, 851, 342, 161, 37, PURPLE, PURPLE_TOKEN, 5, True, WHITE, "Multimodal tokens")
    text(slide, 861, 380, 151, 31, "Mₜ · Multimodal", 12.2, INK, True, name="Multimodal-state label")

    # Underbrace, approximated with one clean native polyline made of segments.
    path(slide, [(209, 415), (209, 421), (216, 428), (522, 428), (532, 435)], INK, 1.2, False, name="Current-state brace left")
    path(slide, [(532, 435), (542, 428), (834, 428), (841, 421), (841, 415)], INK, 1.2, False, name="Current-state brace right")
    text(slide, 430, 439, 314, 28, "HORIZON-INDEPENDENT Sₜ", 11.3, INK, False, name="Horizon-independent state label")

    # Panel (b).
    box(slide, 1055, 82, 102, 54, ORANGE_FILL, ORANGE, 1.0, True, "Horizon card")
    text(slide, 1065, 89, 82, 40, "Horizon h", 11.3, INK, False, name="Horizon label")
    box(slide, 1176, 82, 163, 54, ORANGE_FILL, ORANGE, 1.0, True, "Recorded intervention card")
    text(slide, 1187, 87, 141, 44, "Optional recorded\nintervention aₜ:ₜ₊ₕ", 10.6, INK, False, name="Recorded intervention label")
    box(slide, 1055, 185, 201, 191, ORANGE_FILL, ORANGE, 1.2, True, "Distributional dynamics")
    text(slide, 1065, 196, 181, 48, "TRAINABLE ·\nDistributional dynamics", 10.6, INK, True, name="Distributional dynamics label")
    dynamics_icon(slide, 1104, 254, "Distributional dynamics icon")
    token_bar(slide, 1295, 185, 222, 42, ORANGE, ORANGE_TOKEN, 6, True, WHITE, "Transition tokens")
    text(slide, 1314, 232, 181, 30, "Δₜ,ₕ · Transition", 12.0, INK, True, name="Transition label")
    gmm_fan(slide, 1295, 272, 220, 103, "Future-state GMM")
    text(slide, 1307, 380, 194, 42, "K-component GMM ·\npredicted A / G / M", 11.3, INK, True, name="Future-state GMM label")
    path(
        slide,
        [(1178, 385), (1179, 399), (1178, 411), (1173, 423), (1165, 431), (1154, 435), (1143, 434), (1133, 429), (1126, 420), (1122, 409), (1121, 398), (1121, 385)],
        ORANGE,
        1.55,
        True,
        name="Multi-step rollout arrow",
    )
    text(slide, 1087, 441, 161, 27, "multi-step rollout", 11.3, INK, True, name="Multi-step rollout label")

    text(slide, 20, 15, 320, 35, "(a) Structured current state", 17.2, BLUE_DARK, True, align=PP_ALIGN.LEFT, name="Panel A title")
    text(slide, 1054, 15, 344, 35, "(b) Predictive world model", 17.2, ORANGE, True, align=PP_ALIGN.LEFT, name="Panel B title")


def draw_training_connectors(slide):
    line(slide, 250, 620, 301, 620, GREEN, 1.35, True, name="Target studies to EMA encoder")
    path(slide, [(470, 571), (490, 571), (490, 565), (519, 565)], GREEN, 1.25, True, name="EMA encoder to dense targets")
    line(slide, 470, 622, 519, 622, GREEN, 1.25, True, name="EMA encoder to anatomy targets")
    path(slide, [(470, 672), (490, 672), (490, 690), (519, 690)], GREEN, 1.25, True, name="EMA encoder to global targets")
    path(slide, [(740, 561), (756, 561), (756, 654), (787, 654)], GREEN, 1.15, True, name="Dense target to target fusion")
    path(slide, [(740, 626), (756, 626), (756, 654)], GREEN, 1.15, False, name="Anatomy target join")
    path(slide, [(740, 690), (756, 690), (756, 654)], GREEN, 1.15, False, name="Global target join")
    line(slide, 881, 591, 881, 620, GREEN, 1.25, True, name="Future report to target fusion")
    line(slide, 971, 654, 1018, 654, GREEN, 1.25, True, name="Target fusion to future tokens")
    line(slide, 1144, 647, 1168, 647, GREEN, 1.25, True, name="Future targets to stop-gradient card")

    # EMA parameter update and supervision routes.
    line(slide, 384, 541, 384, 480, "737373", 1.0, True, True, "EMA update arrow")
    path(slide, [(1232, 686), (1232, 741), (224, 741), (224, 773)], GREEN, 1.1, True, True, "Supervision to volume loss")
    line(slide, 555, 741, 555, 773, GREEN, 1.1, True, True, "Supervision to VL loss")
    line(slide, 913, 741, 913, 773, GREEN, 1.1, True, True, "Supervision to future loss")
    path(slide, [(1391, 686), (1391, 741), (1320, 741), (1320, 773)], GREEN, 1.1, True, True, "Supervision to rollout loss")


def draw_training(slide, scans):
    frame = box(slide, 8, 479, 1512, 372, "FEFFFD", GREEN, 1.25, True, "Training-only frame")
    frame.adjustments[0] = 0.015
    set_dash(frame)
    draw_training_connectors(slide)

    box(slide, 20, 541, 230, 161, WHITE, GREEN, 1.0, True, "Target studies card")
    text(slide, 38, 552, 194, 48, "Target studies\nxₜ  and  xₜ₊ₕ", 11.5, INK, True, name="Target studies label")
    picture(slide, scans[0], 31, 613, 60, 67, "Target XR")
    picture(slide, scans[1], 99, 613, 61, 67, "Target CT")
    picture(slide, scans[2], 169, 613, 60, 67, "Target MRI")
    text(slide, 224, 624, 23, 36, "···", 12.5, INK, True, name="Target studies ellipsis")

    box(slide, 303, 541, 167, 162, GREEN_FILL, GREEN, 1.0, True, "EMA visual target encoder")
    text(slide, 317, 551, 139, 48, "EMA visual\ntarget encoder", 11.5, INK, True, name="EMA visual target encoder label")
    volume(slide, 336, 603, 0.75, True, "EMA target volume")

    box(slide, 520, 537, 93, 56, GREEN_FILL, GREEN, 0.95, True, "Target dense glyph card")
    map_stack(slide, 531, 546, 70, 42, GREEN, ["BBD8B7", "8DBA88", "5B9658"], "Target dense maps")
    token_bar(slide, 621, 546, 119, 29, GREEN, GREEN_TOKEN, 4, True, WHITE, "Target dense tokens")
    box(slide, 520, 605, 93, 57, GREEN_FILL, GREEN, 0.95, True, "Target anatomy glyph card")
    wire_cube(slide, 539, 619, 40, 10, GREEN_DARK, "Target anatomy grid")
    token_bar(slide, 621, 613, 119, 29, GREEN, GREEN_TOKEN, 4, True, WHITE, "Target anatomy tokens")
    token_bar(slide, 520, 668, 220, 43, GREEN, GREEN_TOKEN, 7, True, WHITE, "Target global tokens")

    box(slide, 788, 538, 185, 53, GREEN_FILL, GREEN, 0.95, True, "Optional future report")
    text(slide, 797, 543, 167, 43, "Optional future report ·\nTARGET SIDE ONLY", 10.1, INK, False, name="Optional future report label")
    box(slide, 788, 622, 183, 65, GREEN_FILL, GREEN, 1.0, True, "Target-side cross-modal fusion")
    text(slide, 800, 628, 159, 52, "Target-side\ncross-modal fusion", 11.2, INK, True, name="Target-side fusion label")
    token_bar(slide, 1019, 628, 125, 38, GREEN, GREEN_TOKEN, 4, True, WHITE, "Future multimodal target tokens")

    box(slide, 1169, 543, 320, 143, GREEN_FILL, GREEN, 1.05, True, "Stop-gradient targets")
    text(slide, 1196, 551, 266, 43, "STOP-GRADIENT\nDENSE + FUTURE TARGETS", 10.7, INK, False, name="Stop-gradient targets label")
    map_stack(slide, 1186, 600, 80, 60, GREEN, ["BBD8B7", "8DBA88", "5B9658"], "Stop-gradient dense maps")
    wire_cube(slide, 1289, 612, 48, 12, GREEN_DARK, "Stop-gradient anatomy grid")
    token_bar(slide, 1360, 620, 121, 39, GREEN, GREEN_TOKEN, 4, True, WHITE, "Stop-gradient future tokens")

    for x, w, symbol, label, name in [
        (87, 273, "vol", "masked volume", "Masked-volume loss"),
        (406, 286, "VL", "image–text", "Image-text loss"),
        (782, 307, "future", "future-state NLL", "Future-state loss"),
        (1155, 293, "roll", "rollout", "Rollout loss"),
    ]:
        box(slide, x, 774, w, 51, GREEN_FILL, GREEN, 0.95, True, name)
        loss_label(slide, x + 8, 784, w - 16, symbol, label, f"{name} label")

    text(slide, 397, 496, 104, 28, "EMA update", 10.4, INK, True, align=PP_ALIGN.LEFT, name="EMA update label")
    text(slide, 515, 494, 506, 30, "TRAINING ONLY · EMA TARGETS AND JOINT SUPERVISION", 11.3, GREEN_DARK, True, name="Training-only title")


def small_gmm(slide, x, y, name):
    origin = (x + 4, y + 30)
    ellipse(slide, origin[0] - 4, origin[1] - 4, 8, 8, "E88650", ORANGE, 0.5, name=f"{name} origin")
    for ridx, yy in enumerate((y + 2, y + 24, y + 46)):
        line(slide, origin[0], origin[1], x + 28, yy + 8, ORANGE, 0.55, dashed=True, name=f"{name} branch {ridx + 1}")
        for idx in range(4 if ridx < 2 else 3):
            box(slide, x + 33 + idx * 24, yy, 18, 16, ORANGE_FILL, ORANGE, 0.55, False, f"{name} token {ridx + 1}-{idx + 1}")
        text(slide, x + 132, yy - 3, 21, 22, "···", 11, ORANGE, True, name=f"{name} ellipsis {ridx + 1}")


def draw_bottom(slide):
    text(slide, 20, 863, 460, 35, "(c) Frozen state · lightweight readouts", 16.5, GRAY, True, align=PP_ALIGN.LEFT, name="Panel C title")

    cards = [(21, 911, 319, 95), (381, 911, 359, 95), (781, 911, 303, 95), (1124, 911, 388, 95)]
    for idx, (x, y, w, h) in enumerate(cards):
        box(slide, x, y, w, h, CARD, "666666", 0.9, True, f"Readout card {idx + 1}")

    map_stack(slide, 35, 923, 71, 54, BLUE, ["A9C3ED", "6F97D7", "3D6DBD"], "Readout dense maps")
    wire_cube(slide, 119, 937, 43, 10, BLUE_DARK, "Readout anatomy grid")
    text(slide, 195, 918, 133, 77, "Dₜ / Aₜ  →\nSegmentation ·\nLocalization", 11.6, INK, True, align=PP_ALIGN.LEFT, name="Dense anatomy readout label")

    token_bar(slide, 396, 942, 163, 34, PURPLE, PURPLE_TOKEN, 5, True, WHITE, "Readout multimodal tokens")
    text(slide, 582, 918, 147, 77, "Gₜ / Mₜ  →\nClassification ·\nRetrieval · VQA", 11.5, INK, True, align=PP_ALIGN.LEFT, name="Global multimodal readout label")

    token_bar(slide, 795, 942, 133, 34, ORANGE, ORANGE_TOKEN, 4, True, WHITE, "Readout transition tokens")
    text(slide, 960, 923, 112, 69, "Δₜ,ₕ  →\nProgression", 11.8, INK, True, align=PP_ALIGN.LEFT, name="Transition readout label")

    small_gmm(slide, 1139, 927, "Readout future-state GMM")
    text(slide, 1327, 921, 174, 72, "Predicted future state  →\nForecasting", 11.5, INK, True, align=PP_ALIGN.LEFT, name="Forecasting readout label")


def build():
    with Image.open(SOURCE) as image:
        scans = [
            crop_blob(image, (29, 134, 96, 209)),
            crop_blob(image, (106, 134, 176, 209)),
            crop_blob(image, (184, 134, 254, 209)),
        ]

    prs = Presentation()
    prs.slide_width = Inches(15)
    prs.slide_height = Inches(10)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    draw_top(slide, scans)
    draw_training(slide, scans)
    draw_bottom(slide)

    prs.core_properties.title = "MedWorld-JEPA Figure 1"
    prs.core_properties.subject = "Editable recreation of fig1_v1"
    prs.core_properties.author = "Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
