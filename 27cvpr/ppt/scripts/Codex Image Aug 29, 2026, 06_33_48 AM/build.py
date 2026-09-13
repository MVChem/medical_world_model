from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
NAME = "Codex Image Aug 29, 2026, 06_33_48 AM"
SOURCE = ROOT / "source" / f"{NAME}.png"
OUTPUT = ROOT / "ppt" / f"{NAME}.pptx"

W, H, PX_PER_INCH = 1536, 1024, 100
BG, WHITE, INK = "FCFBF7", "FFFDFC", "101721"
GRAY, MID_GRAY, PALE_GRAY = "6E7476", "A9AEAD", "E8E8E3"
BLUE, BLUE_MID, BLUE_PALE = "104477", "6F98BE", "EAF1F6"
TEAL, TEAL_MID, TEAL_PALE = "087477", "55A8A4", "E9F5F2"
PURPLE, PURPLE_MID, PURPLE_PALE = "5A3478", "8D78AE", "F2EDF5"
ORANGE, ORANGE_MID, ORANGE_PALE = "D34B22", "E89872", "FFF1E8"
GREEN, GREEN_MID, GREEN_PALE = "3D7C42", "86B689", "EFF6EB"
FONT, MATH = "Arial", "Times New Roman"


def u(px):
    return Inches(px / PX_PER_INCH)


def rgb(value):
    return RGBColor.from_string(value)


def clean(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def dash(shape, value="dash"):
    ln = shape._element.spPr.get_or_add_ln()
    for child in list(ln):
        if child.tag.endswith("prstDash"):
            ln.remove(child)
    node = OxmlElement("a:prstDash")
    node.set("val", value)
    ln.append(node)


def arrow(shape):
    ln = shape._element.spPr.get_or_add_ln()
    node = OxmlElement("a:tailEnd")
    node.set("type", "triangle")
    node.set("w", "sm")
    node.set("len", "sm")
    ln.append(node)


def box(slide, x, y, w, h, fill=WHITE, line=GRAY, width=1, radius=True, name="Box"):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, u(x), u(y), u(w), u(h))
    s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(width)
    if radius:
        s.adjustments[0] = 0.08
    return clean(s)


def text(slide, x, y, w, h, value, size=12, color=INK, bold=False, italic=False,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, font=FONT, name=None):
    s = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    s.name = name or f"Text: {value[:28]}"
    tf = s.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = u(1)
    tf.margin_top = tf.margin_bottom = u(1)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 0.94
    r = p.add_run()
    r.text = value
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    return clean(s)


def line(slide, x1, y1, x2, y2, color=GRAY, width=1.1, end=False, dashed=False, name="Connector"):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2))
    s.name = name
    s.line.color.rgb = rgb(color)
    s.line.width = Pt(width)
    if dashed:
        dash(s)
    if end:
        arrow(s)
    return clean(s)


def path(slide, points, color=GRAY, width=1.1, end=True, dashed=False, name="Path"):
    for i, (a, b) in enumerate(zip(points, points[1:])):
        line(slide, *a, *b, color, width, end and i == len(points) - 2, dashed, f"{name} {i + 1}")


def poly(slide, points, fill, edge, width=0.8, name="Polygon"):
    x0, y0 = min(x for x, _ in points), min(y for _, y in points)
    local = [(x - x0, y - y0) for x, y in points]
    b = slide.shapes.build_freeform(local[0][0], local[0][1], scale=Inches(1) / PX_PER_INCH)
    b.add_line_segments(local[1:], close=True)
    s = b.convert_to_shape(u(x0), u(y0))
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(edge)
    s.line.width = Pt(width)
    return clean(s)


def oval(slide, x, y, w, h, fill=None, edge=GRAY, width=1, name="Oval"):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(w), u(h))
    s.name = name
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    if edge is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = rgb(edge)
        s.line.width = Pt(width)
    return clean(s)


def crop_blob(image, crop):
    stream = BytesIO()
    image.crop(crop).save(stream, format="PNG")
    stream.seek(0)
    return stream


def picture(slide, blob, x, y, w, h, name):
    blob.seek(0)
    s = slide.shapes.add_picture(blob, u(x), u(y), u(w), u(h))
    s.name = name
    return clean(s)


def grid(slide, x, y, w, h, cols, rows, edge, fill, name, line_width=0.45):
    box(slide, x, y, w, h, fill, edge, 0.8, False, name)
    for i in range(1, cols):
        line(slide, x + w * i / cols, y, x + w * i / cols, y + h, edge, line_width,
             name=f"{name} vertical {i}")
    for i in range(1, rows):
        line(slide, x, y + h * i / rows, x + w, y + h * i / rows, edge, line_width,
             name=f"{name} horizontal {i}")


def heat_grid(slide, x, y, w, h, cols, rows, palette, edge, name, gap=0):
    cw, ch = w / cols, h / rows
    for row in range(rows):
        for col in range(cols):
            idx = (row * 5 + col * 3 + row * col) % len(palette)
            box(slide, x + col * cw + gap / 2, y + row * ch + gap / 2,
                cw - gap, ch - gap, palette[idx], edge, 0.28, False,
                f"{name} cell {row + 1}-{col + 1}")
    box(slide, x, y, w, h, None, edge, 0.8, False, f"{name} border")


def token(slide, x, y, color, fill, name, w=15, h=18, radius=2):
    s = box(slide, x, y, w, h, fill, color, 0.65, True, name)
    s.adjustments[0] = min(0.18, radius / max(1, min(w, h)))
    return s


def cube(slide, x, y, size, color, light, dark, name, cells=3):
    depth = size * 0.24
    front_y = y + depth
    grid(slide, x, front_y, size, size, cells, cells, color, light, f"{name} front", 0.35)
    poly(slide, [(x, front_y), (x + depth, y), (x + size + depth, y), (x + size, front_y)],
         light, color, 0.65, f"{name} top")
    poly(slide, [(x + size, front_y), (x + size + depth, y),
                 (x + size + depth, y + size), (x + size, front_y + size)],
         dark, color, 0.65, f"{name} side")
    for i in range(1, cells):
        t = i / cells
        line(slide, x + depth + size * t, y, x + size * t, front_y, color, 0.3,
             name=f"{name} top grid {i}")
        line(slide, x + size, front_y + size * t,
             x + size + depth, y + size * t, color, 0.3, name=f"{name} side grid {i}")


def lock_icon(slide, x, y, color, name, scale=1):
    oval(slide, x + 3 * scale, y, 12 * scale, 13 * scale, None, color, 1, f"{name} shackle")
    box(slide, x, y + 7 * scale, 18 * scale, 17 * scale, color, color, 0.6, True, f"{name} body")
    oval(slide, x + 7.2 * scale, y + 11 * scale, 3.6 * scale, 4 * scale, WHITE, WHITE, 0.2,
         f"{name} keyhole")
    line(slide, x + 9 * scale, y + 14 * scale, x + 9 * scale, y + 18 * scale, WHITE, 0.8,
         name=f"{name} key stem")


def flame(slide, x, y, name, scale=1):
    pts = [(x + 8 * scale, y), (x + 13 * scale, y + 9 * scale),
           (x + 12 * scale, y + 18 * scale), (x + 7 * scale, y + 22 * scale),
           (x + 1 * scale, y + 18 * scale), (x, y + 10 * scale),
           (x + 5 * scale, y + 4 * scale), (x + 5 * scale, y + 12 * scale)]
    poly(slide, pts, ORANGE, "A93C1D", 0.55, name)
    poly(slide, [(x + 7 * scale, y + 9 * scale), (x + 10 * scale, y + 15 * scale),
                 (x + 7 * scale, y + 19 * scale), (x + 4 * scale, y + 15 * scale)],
         WHITE, WHITE, 0.3, f"{name} inner")


def shield(slide, x, y, color, name, scale=1):
    pts = [(x + 14 * scale, y), (x + 26 * scale, y + 5 * scale),
           (x + 24 * scale, y + 19 * scale), (x + 14 * scale, y + 29 * scale),
           (x + 4 * scale, y + 19 * scale), (x + 2 * scale, y + 5 * scale)]
    poly(slide, pts, color, color, 0.7, name)
    poly(slide, [(x + 14 * scale, y + 4 * scale), (x + 22 * scale, y + 7 * scale),
                 (x + 20 * scale, y + 17 * scale), (x + 14 * scale, y + 23 * scale),
                 (x + 8 * scale, y + 17 * scale), (x + 6 * scale, y + 7 * scale)],
         WHITE, WHITE, 0.4, f"{name} inner")
    line(slide, x + 14 * scale, y + 7 * scale, x + 14 * scale, y + 19 * scale, color, 1,
         name=f"{name} vertical")
    line(slide, x + 10 * scale, y + 13 * scale, x + 18 * scale, y + 13 * scale, color, 1,
         name=f"{name} horizontal")


def bar_chart(slide, x, y, color, name):
    line(slide, x, y + 28, x + 42, y + 28, color, 0.7, name=f"{name} axis")
    for i, (height, shade) in enumerate(((13, "8BC0B9"), (28, color), (17, "59A39B"))):
        box(slide, x + 7 + i * 10, y + 28 - height, 7, height, shade, color, 0.45, False,
            f"{name} bar {i + 1}")


def heart_pulse(slide, x, y, color, name):
    h = slide.shapes.add_shape(MSO_SHAPE.HEART, u(x), u(y), u(38), u(34))
    h.name = name
    h.fill.solid()
    h.fill.fore_color.rgb = rgb(color)
    h.line.color.rgb = rgb(color)
    clean(h)
    path(slide, [(x + 4, y + 17), (x + 12, y + 17), (x + 16, y + 9),
                 (x + 21, y + 25), (x + 25, y + 16), (x + 34, y + 16)],
         WHITE, 1, False, False, f"{name} pulse")


def network(slide, x, y, w, h, color, name):
    left = [(x, y + h * v) for v in (0.18, 0.5, 0.82)]
    mid = [(x + w * 0.52, y + h * v) for v in (0.05, 0.35, 0.65, 0.95)]
    right = [(x + w, y + h * v) for v in (0.22, 0.52, 0.82)]
    for i, a in enumerate(left):
        for j, b in enumerate(mid):
            line(slide, *a, *b, color, 0.45, name=f"{name} edge L{i}M{j}")
    for i, a in enumerate(mid):
        for j, b in enumerate(right):
            line(slide, *a, *b, color, 0.45, name=f"{name} edge M{i}R{j}")
    for i, (nx, ny) in enumerate(left + mid + right):
        oval(slide, nx - 6, ny - 6, 12, 12, ORANGE_MID if color == ORANGE else GREEN_MID,
             color, 0.65, f"{name} node {i + 1}")


def draw_routes(slide):
    line(slide, 314, 269, 339, 269, BLUE, 1.25, True, name="Evidence to encoder")
    path(slide, [(464, 270), (496, 270), (496, 390), (522, 390)], BLUE, 1.3, True,
         name="Encoder to VLM")
    line(slide, 399, 411, 399, 532, BLUE, 1.15, True, name="Encoder to resampler")
    for y, color in ((576, BLUE_MID), (616, BLUE), (691, BLUE)):
        line(slide, 307, y, 338, y, color, 1.05, True, name=f"Context route {y}")
    line(slide, 307, 765, 338, 765, ORANGE, 1.05, True, name="Report to tokenizer")
    line(slide, 508, 625, 559, 625, BLUE, 1.2, True, name="Resampler to representation")
    for x in (569, 703, 822):
        line(slide, x, 214, x, 292, BLUE, 0.95, True, True, f"Dense memory read {x}")
    line(slide, 716, 548, 716, 589, PURPLE, 1.2, True, name="VLM to representation")
    path(slide, [(873, 627), (943, 627), (943, 715)], PURPLE, 1.25, True,
         name="Representation to predictor")
    line(slide, 905, 421, 938, 421, PURPLE, 1.05, True, name="Slots to query panel")
    path(slide, [(942, 324), (942, 136), (1087, 136)], TEAL, 1, True, True,
         name="Read route to downstream reuse")
    path(slide, [(1057, 422), (1118, 422), (1118, 294)], TEAL, 1.25, True,
         name="Task query to downstream reuse")
    path(slide, [(1057, 521), (1094, 521), (1094, 618)], ORANGE, 1.2, True,
         name="Forecast query to horizon")
    line(slide, 1094, 678, 1094, 722, ORANGE, 1.1, True, name="Horizon to transition query")
    line(slide, 1058, 817, 1077, 817, ORANGE, 1.2, True, name="Predictor to transition")
    line(slide, 1192, 817, 1212, 817, ORANGE, 1.2, True, name="Transition to prediction")
    line(slide, 1354, 627, 1354, 649, GREEN, 1.05, True, True, "EMA target to loss")
    path(slide, [(1136, 648), (1225, 648), (1248, 657), (1307, 681)], ORANGE, 1,
         True, True, "Horizon loss signal")
    line(slide, 1354, 763, 1354, 735, ORANGE, 1.05, True, True, "Prediction to loss")
    path(slide, [(1212, 899), (1150, 925), (602, 925), (537, 910), (537, 886),
                 (621, 821), (676, 760), (696, 663)], ORANGE, 1.05, True, True,
         "Future loss shapes representation")


def draw_evidence(slide, scans):
    box(slide, 10, 18, 305, 786, WHITE, BLUE_MID, 1.3, True, "Evidence panel")
    text(slide, 66, 34, 195, 36, "Evidence ≤ t", 16, BLUE, True, name="Evidence heading")

    sections = [
        (23, 75, 280, 133, BLUE_MID, "2D studies"),
        (23, 226, 280, 146, BLUE_MID, "3D studies"),
        (22, 390, 281, 145, ORANGE_MID, "Report (text)"),
        (22, 553, 281, 228, PURPLE_MID, "History & context"),
    ]
    for i, (x, y, w, h, color, label) in enumerate(sections):
        box(slide, x, y, w, h, WHITE, color, 0.85, True, f"Evidence section {i + 1}")
        text(slide, x + 13, y + 7, w - 26, 24, label, 11.5, color if i > 1 else BLUE,
             i < 2, align=PP_ALIGN.LEFT, name=f"{label} label")

    frames = [(32, 111, 75, 80), (111, 111, 71, 80), (187, 111, 73, 80),
              (32, 261, 75, 93), (111, 261, 75, 93), (191, 261, 78, 93)]
    for i, ((x, y, w, h), scan) in enumerate(zip(frames, scans)):
        picture(slide, scan, x, y, w, h, f"Medical evidence thumbnail {i + 1}")
    text(slide, 267, 132, 28, 36, "…", 15, BLUE, True, name="2D studies ellipsis")
    text(slide, 270, 286, 26, 36, "…", 15, BLUE, True, name="3D studies ellipsis")

    box(slide, 33, 425, 223, 100, "FFFCF8", ORANGE_MID, 0.65, True, "Report text card")
    report = "Findings: Bilateral ground-glass\nopacities in the lungs  …\nImpression: Infectious or\ninflammatory etiology  …"
    text(slide, 43, 432, 203, 87, report, 9.2, INK, False, False, PP_ALIGN.LEFT,
         MSO_ANCHOR.TOP, name="Editable report text")
    box(slide, 266, 479, 25, 34, WHITE, ORANGE_MID, 0.8, False, "Report document icon")
    for i, length in enumerate((11, 14, 10)):
        line(slide, 272, 488 + i * 6, 272 + length, 488 + i * 6, ORANGE, 0.7,
             name=f"Report document line {i + 1}")
    text(slide, 264, 513, 29, 12, "▰▰▰", 4.5, ORANGE_MID, name="Report document dots")

    # History icons and labels.
    oval(slide, 52, 600, 12, 12, PURPLE_MID, PURPLE, 0.55, "Person head")
    poly(slide, [(48, 624), (51, 615), (58, 612), (65, 615), (68, 624)],
         PURPLE_MID, PURPLE, 0.55, "Person shoulders")
    text(slide, 78, 596, 129, 26, "Age / Sex", 9.8, align=PP_ALIGN.LEFT, name="Age sex label")
    box(slide, 208, 597, 22, 31, WHITE, PURPLE, 0.7, False, "Context clipboard")
    box(slide, 214, 594, 9, 5, ORANGE_MID, PURPLE, 0.5, True, "Clipboard clip")
    for r in range(3):
        for c in range(2):
            box(slide, 213 + c * 7, 605 + r * 6, 4, 4, "E9E4F0", PURPLE, 0.35, False,
                f"Clipboard cell {r}-{c}")
    text(slide, 253, 602, 26, 23, "…", 14, BLUE, True, name="Context ellipsis")

    box(slide, 49, 641, 18, 25, WHITE, PURPLE, 0.7, False, "Diagnosis clipboard")
    box(slide, 54, 638, 8, 5, ORANGE_MID, PURPLE, 0.5, True, "Diagnosis clip")
    for i in range(3):
        line(slide, 54, 648 + i * 5, 63, 648 + i * 5, PURPLE, 0.55,
             name=f"Diagnosis line {i + 1}")
    text(slide, 78, 638, 151, 27, "Prior diagnoses", 9.8, align=PP_ALIGN.LEFT,
         name="Prior diagnoses label")

    box(slide, 53, 682, 10, 23, "E5DFF0", PURPLE, 0.7, True, "Lab tube")
    line(slide, 51, 681, 65, 681, PURPLE, 0.8, name="Lab tube lip")
    box(slide, 56, 695, 4, 7, PURPLE_MID, None, 0, False, "Lab tube liquid")
    text(slide, 78, 678, 150, 28, "Labs", 9.8, align=PP_ALIGN.LEFT, name="Labs label")

    text(slide, 35, 721, 51, 23, "Timing", 9.7, align=PP_ALIGN.LEFT, name="Timing label")
    line(slide, 92, 732, 270, 732, MID_GRAY, 0.75, True, name="Timing axis")
    for i, (x, fill, edge) in enumerate(((125, "D7D6D2", MID_GRAY),
                                         (172, "D7D6D2", MID_GRAY),
                                         (219, BLUE_MID, BLUE))):
        oval(slide, x - 7, 725, 14, 14, fill, edge, 0.55, f"Timing point {i + 1}")
    text(slide, 106, 742, 39, 24, "t-2", 9, italic=True, font=MATH, name="Timing t-2")
    text(slide, 153, 742, 39, 24, "t-1", 9, italic=True, font=MATH, name="Timing t-1")
    text(slide, 202, 742, 35, 24, "t", 9, italic=True, font=MATH, name="Timing t")
    text(slide, 273, 719, 24, 25, "…", 13, BLUE, True, name="Timing ellipsis")

    # Token key beneath evidence panel.
    token(slide, 74, 835, BLUE, BLUE_MID, "Visual token key 1", 13, 16)
    token(slide, 91, 835, BLUE, BLUE_MID, "Visual token key 2", 13, 16)
    text(slide, 121, 832, 167, 24, "Visual tokens (studies)", 9.2, align=PP_ALIGN.LEFT,
         name="Visual token key label")
    token(slide, 74, 867, ORANGE, "E8967C", "Text token key 1", 13, 16)
    token(slide, 91, 867, ORANGE, "E8967C", "Text token key 2", 13, 16)
    text(slide, 121, 864, 176, 24, "Text / context tokens", 9.2, align=PP_ALIGN.LEFT,
         name="Text token key label")
    token(slide, 74, 899, PURPLE, PURPLE_MID, "Special token key 1", 13, 16)
    token(slide, 91, 899, PURPLE, PURPLE_MID, "Special token key 2", 13, 16)
    text(slide, 121, 896, 165, 24, "Special tokens", 9.2, align=PP_ALIGN.LEFT,
         name="Special token key label")


def draw_encoder_and_resampler(slide):
    poly(slide, [(339, 110), (464, 146), (464, 365), (339, 412)],
         "EEF3F6", BLUE_MID, 1.05, "2D 3D encoder panel")
    text(slide, 354, 190, 93, 66, "2D / 3D\nencoder", 15.5, BLUE, True, name="Encoder label")
    cube(slide, 359, 264, 54, BLUE, "AFC7DB", "7298BA", "Encoder cube", 3)

    box(slide, 339, 533, 169, 180, "F8FBFC", BLUE_MID, 1, True, "Resampler projector")
    text(slide, 352, 543, 145, 47, "Resampler\n+ projector", 12.3, BLUE, True,
         align=PP_ALIGN.LEFT, name="Resampler projector label")
    poly(slide, [(351, 606), (388, 606), (370, 662)], "80A4C1", BLUE, 0.85, "Resampler funnel")
    for yy in (605, 613, 621):
        line(slide, 353, yy, 385, yy, BLUE, 0.35, name=f"Funnel band {yy}")
    for row in range(3):
        text(slide, 398, 605 + row * 35, 20, 22, "→" if row < 2 else "…", 11, BLUE, True,
             name=f"Resampler row arrow {row + 1}")
        for col in range(4):
            token(slide, 426 + col * 18, 605 + row * 35, BLUE, "A9C2D7",
                  f"Resampler token {row + 1}-{col + 1}", 15, 17)

    box(slide, 338, 735, 170, 61, "FFF9F4", ORANGE_MID, 0.9, True, "Text context tokenizer")
    text(slide, 352, 742, 144, 46, "Text & context\ntokenizer", 11.7, "A8230E", True,
         align=PP_ALIGN.LEFT, name="Text context tokenizer label")


def draw_dense_memory(slide):
    box(slide, 503, 39, 384, 195, "FBFCFD", BLUE_MID, 1.05, True, "Dense memory panel")
    text(slide, 520, 47, 350, 35, "Dense memory  Dₜ¹:ᴸ  (multiscale)", 13.2, BLUE, False,
         name="Dense memory heading")
    text(slide, 526, 83, 96, 27, "L (coarse)", 10.5, INK, False, False, font=MATH,
         name="Coarse memory label")
    text(slide, 675, 84, 69, 27, "ℓ", 12.5, INK, False, True, font=MATH,
         name="Intermediate memory label")
    text(slide, 789, 84, 77, 27, "1 (fine)", 10.5, INK, False, False, font=MATH,
         name="Fine memory label")
    blue_palette = ["D9E5EE", "BCD0DF", "91AEC7", "6F94B6", "EAF0F4"]
    heat_grid(slide, 518, 113, 108, 103, 6, 6, blue_palette, BLUE, "Coarse dense memory")
    heat_grid(slide, 665, 121, 86, 84, 5, 5, blue_palette, BLUE, "Intermediate dense memory")
    heat_grid(slide, 787, 121, 84, 84, 10, 8, blue_palette, BLUE, "Fine dense memory")
    text(slide, 632, 141, 26, 32, "…", 15, INK, True, name="Dense memory ellipsis 1")
    text(slide, 756, 141, 26, 32, "…", 15, INK, True, name="Dense memory ellipsis 2")


def draw_vlm(slide):
    box(slide, 522, 297, 383, 251, "FCF9FC", PURPLE_MID, 1.1, True, "VLM LoRA panel")
    text(slide, 595, 309, 236, 39, "VLM + LoRA", 16, "392252", True, name="VLM LoRA heading")
    flame(slide, 866, 305, "VLM trainable flame", 1.0)

    box(slide, 528, 369, 188, 59, "FFFCFF", "B8A7C7", 0.6, True, "VLM input context")
    box(slide, 722, 369, 176, 59, "FFFCFF", "B8A7C7", 0.6, True, "VLM slots region")
    for i in range(4):
        token(slide, 535 + i * 21, 386, BLUE, "82A7C7", f"VLM visual token {i + 1}", 16, 18)
    text(slide, 619, 383, 25, 25, "…", 13, INK, True, name="VLM input ellipsis")
    for i in range(3):
        token(slide, 652 + i * 21, 386, ORANGE, "E5957B", f"VLM context token {i + 1}", 16, 18)
    line(slide, 716, 373, 716, 424, "B8A7C7", 0.6, dashed=True, name="VLM token divider")
    text(slide, 745, 346, 134, 27, "8 generic slots  Sₜ", 10.5, "392252", name="Generic slots label")
    for i in range(8):
        token(slide, 730 + i * 21, 386, PURPLE, PURPLE_MID, f"Generic slot {i + 1}", 16, 18)
        text(slide, 728 + i * 21, 408, 20, 20, str(i + 1), 7.5, INK, name=f"Slot number {i + 1}")

    xs = [531 + i * 21 for i in range(17)]
    rows = [441, 474, 508]
    for row, yy in enumerate(rows):
        for i, xx in enumerate(xs):
            fill = "F0EAF3" if (i + row) % 4 else "E5DCEB"
            box(slide, xx, yy, 19, 18, fill, "C8BBD2", 0.35, True,
                f"Transformer block {row + 1}-{i + 1}")
            if row < 2:
                line(slide, xx + 9.5, yy + 18, xx + 9.5, rows[row + 1], "C8BBD2", 0.35,
                     name=f"Transformer vertical {row + 1}-{i + 1}")
        if row < 2:
            for i in range(16):
                if i % 2 == row:
                    line(slide, xs[i] + 10, yy + 9, xs[i + 1] + 9, yy + 9, "D5CBDD", 0.32,
                         name=f"Transformer lateral {row + 1}-{i + 1}")
    for i in range(7):
        line(slide, 543 + i * 21, 404, 543 + i * 21, 441,
             BLUE if i < 4 else ORANGE, 0.45, name=f"Token projection {i + 1}")
    for i in range(8):
        line(slide, 738 + i * 21, 404, 738 + i * 21, 441, PURPLE, 0.45,
             name=f"Slot projection {i + 1}")
    box(slide, 535, 523, 178, 12, "EEE7F1", "C8BBD2", 0.35, True, "Transformer readout left")
    box(slide, 725, 523, 166, 12, "EEE7F1", "C8BBD2", 0.35, True, "Transformer readout right")
    oval(slide, 918, 414, 12, 12, PURPLE, PURPLE, 0.5, "VLM query output")

    box(slide, 559, 590, 314, 73, "FBF8FC", PURPLE, 1.05, True, "Reusable representation")
    text(slide, 586, 594, 260, 39, "Rₜ = (Dₜ¹:ᴸ, Sₜ)", 16.5, "2E1749", False, True,
         font=MATH, name="Reusable representation formula")
    text(slide, 596, 629, 241, 26, "Reusable current representation", 10.2, INK,
         name="Reusable representation subtitle")


def draw_queries(slide):
    box(slide, 938, 324, 119, 280, "FAFAF8", MID_GRAY, 1, True, "Read-only query panel")
    text(slide, 951, 335, 91, 38, "Read-only\nqueries Q", 9.5, INK, name="Read-only queries heading")
    lock_icon(slide, 1026, 337, GRAY, "Read-only query lock", 0.65)
    box(slide, 947, 386, 100, 87, "F2FAF8", TEAL_MID, 0.85, True, "Task query panel")
    text(slide, 955, 391, 84, 25, "Task query", 9.1, "114B5B", name="Task query label")
    text(slide, 961, 412, 72, 25, "q", 12.5, TEAL, False, True, font=MATH, name="Task query symbol")
    for i in range(4):
        token(slide, 960 + i * 19, 443, TEAL, "64B0AC", f"Task query token {i + 1}", 15, 17)

    box(slide, 947, 498, 100, 90, "FFF8F2", ORANGE_MID, 0.85, True, "Forecast query panel")
    text(slide, 951, 504, 92, 24, "Forecast query", 8.8, "B32612", name="Forecast query label")
    text(slide, 958, 526, 77, 27, "(A, h)", 12, ORANGE, False, True, font=MATH,
         name="Forecast query symbol")
    for i in range(4):
        token(slide, 960 + i * 19, 559, ORANGE, "F0B08F", f"Forecast query token {i + 1}", 15, 17)

    box(slide, 1050, 618, 86, 60, "FFF9F1", "EAB575", 0.8, True, "Horizon cue")
    text(slide, 1071, 621, 45, 26, "h", 14, INK, False, True, font=MATH, name="Horizon symbol")
    text(slide, 1057, 645, 72, 25, "(horizon cue)", 8.2, INK, name="Horizon cue label")

    box(slide, 1077, 722, 115, 100, "FFF9F3", "E8A56D", 0.85, True, "Transition query")
    text(slide, 1091, 727, 87, 29, "Aₜ,ₕ", 13, INK, False, True, font=MATH,
         name="Transition query symbol")
    text(slide, 1084, 753, 100, 38, "(transition\nquery)", 8.7, INK, name="Transition query label")
    for i in range(5):
        token(slide, 1088 + i * 19, 795, ORANGE, "FFE1CD", f"Transition token {i + 1}", 15, 17)


def draw_downstream(slide, reuse_scans):
    box(slide, 1087, 15, 432, 279, "F4F8F4", "6B9B89", 1.15, True, "Downstream task reuse")
    box(slide, 1092, 19, 422, 270, None, "B1C7B6", 0.45, True, "Downstream inner outline")
    text(slide, 1108, 24, 390, 31, "Downstream task reuse (read Sₜ, Dₜ¹:ᴸ)", 12.3, "07585D", True,
         name="Downstream task reuse heading")
    box(slide, 1133, 59, 342, 97, "FFFEFC", TEAL_MID, 0.85, True, "Global clinical output")
    text(slide, 1144, 66, 200, 24, "Global clinical output", 10.5, "07585D", True,
         align=PP_ALIGN.LEFT, name="Global clinical output heading")
    bar_chart(slide, 1164, 97, TEAL, "Risk chart")
    shield(slide, 1252, 93, TEAL, "Survival shield", 1)
    heart_pulse(slide, 1337, 94, TEAL, "Outcome heart")
    for x in (1226, 1310):
        line(slide, x, 92, x, 146, "C8DAD6", 0.55, name=f"Clinical output divider {x}")
    text(slide, 1158, 130, 55, 22, "Risk", 9.2, name="Risk label")
    text(slide, 1236, 130, 67, 22, "Survival", 9.2, name="Survival label")
    text(slide, 1327, 130, 69, 22, "Outcome", 9.2, name="Outcome label")
    text(slide, 1410, 95, 34, 31, "…", 14, INK, True, name="Global output ellipsis")

    box(slide, 1133, 164, 342, 118, "FFFEFC", TEAL_MID, 0.85, True,
        "Dense localization segmentation")
    text(slide, 1145, 169, 257, 25, "Dense localization / segmentation", 10.4, "07585D", True,
         align=PP_ALIGN.LEFT, name="Dense localization heading")
    placements = [(1147, 197, 91, 55), (1252, 197, 80, 55), (1344, 194, 94, 65)]
    for i, (scan, (x, y, w, h)) in enumerate(zip(reuse_scans, placements)):
        picture(slide, scan, x, y, w, h, f"Downstream medical thumbnail {i + 1}")
    text(slide, 1442, 211, 28, 31, "…", 15, INK, True, name="Dense output ellipsis")


def draw_future_target(slide, future_scan):
    box(slide, 1171, 342, 343, 159, "F7FAF3", GREEN_MID, 0.95, True, "Future EMA target")
    text(slide, 1220, 348, 240, 31, "Future Xₜ⁺  ·  EMA target", 11.8, "176524", True,
         name="Future EMA heading")
    shield(slide, 1475, 352, GREEN, "EMA shield", 0.8)
    box(slide, 1188, 384, 51, 104, "FFFEFA", "C9D8C5", 0.55, True, "Future scan card")
    picture(slide, future_scan, 1191, 387, 46, 48, "Future target scan")
    line(slide, 1193, 452, 1233, 452, "D8E1D5", 0.5, name="Future card line")
    text(slide, 1197, 458, 31, 21, "…", 11, INK, True, name="Future card ellipsis")
    line(slide, 1242, 436, 1267, 436, GREEN, 1.05, True, name="Future scan to EMA encoder")
    box(slide, 1268, 389, 114, 98, "F5FAF2", GREEN_MID, 0.75, True, "EMA target encoder")
    text(slide, 1280, 394, 90, 35, "EMA target\nencoder", 9.3, "174D20", name="EMA encoder label")
    network(slide, 1288, 437, 72, 38, GREEN, "EMA encoder network")
    line(slide, 1383, 436, 1407, 436, GREEN, 1.05, True, name="EMA encoder to pooling")
    box(slide, 1408, 389, 94, 98, "F5FAF2", GREEN_MID, 0.75, True, "Multiscale pooling")
    text(slide, 1416, 395, 78, 35, "Pooling\n(multiscale)", 9.1, "174D20", name="Pooling label")
    grid(slide, 1427, 446, 55, 41, 4, 3, GREEN, "C9DFC5", "Pooling grid", 0.35)
    box(slide, 1428, 446, 27, 13.5, "A5CAA0", None, 0, False, "Pooling highlight 1")
    box(slide, 1441.5, 459.5, 13.5, 27, "91BD8C", None, 0, False, "Pooling highlight 2")

    box(slide, 1178, 502, 336, 125, "F9FBF5", GREEN_MID, 0.85, True,
        "Stop-gradient target hierarchy")
    text(slide, 1244, 506, 205, 24, "Stop-gradient target hierarchy", 9.2, "174D20",
         name="Target hierarchy heading")
    text(slide, 1195, 534, 55, 31, "Global\nlatent", 8, INK, name="Target global latent label")
    cube(slide, 1198, 555, 31, GREEN, "9BC994", "5B9A5B", "Target global cube", 3)
    text(slide, 1250, 570, 24, 28, "+", 18, "174D20", True, name="Target plus")
    text(slide, 1275, 535, 180, 24, "Regional latents (multiscale)", 8.4, INK,
         name="Target regional latents label")
    heat_grid(slide, 1283, 558, 45, 57, 4, 5, ["DCEAD8", "B6D3B0", "82B27F"], GREEN,
              "Target regional grid coarse")
    heat_grid(slide, 1360, 558, 45, 57, 5, 5, ["E0ECD9", "BAD5B5", "79AE75"], GREEN,
              "Target regional grid medium")
    heat_grid(slide, 1434, 566, 35, 42, 5, 5, ["E6EFE1", "BCD7B7", "7EB37A"], GREEN,
              "Target regional grid fine")
    text(slide, 1480, 570, 25, 25, "…", 13, INK, True, name="Target hierarchy ellipsis")


def draw_predictor_and_loss(slide):
    labels = [(741, 706, "Dₜ¹:ᴸ", BLUE), (746, 757, "Sₜ", PURPLE),
              (735, 804, "Aₜ,ₕ", ORANGE), (747, 849, "h", ORANGE)]
    for x, y, value, color in labels:
        text(slide, x, y, 44, 29, value, 12.5, color, False, True, font=MATH,
             name=f"Predictor input {value}")
    heat_grid(slide, 781, 715, 61, 27, 8, 3, ["D4E2ED", "A8C1D7", "7C9FBE"], BLUE,
              "Predictor dense input")
    text(slide, 846, 715, 29, 28, "↦", 13, PURPLE, name="Dense input map arrow")
    for i in range(4):
        token(slide, 780 + i * 18, 758, PURPLE, "9A84BA", f"Predictor slot input {i + 1}", 14, 17)
        token(slide, 780 + i * 18, 803, ORANGE, "EDA17E", f"Predictor action input {i + 1}", 14, 17)
    token(slide, 780, 847, "E8A765", "FFF0DB", "Predictor horizon input", 14, 17)
    for yy, color in ((729, BLUE), (767, PURPLE), (812, ORANGE), (856, ORANGE)):
        line(slide, 850 if yy == 729 else 852, yy, 884, yy, color, 1, True,
             name=f"Predictor input route {yy}")

    box(slide, 884, 715, 174, 176, "FFF7F2", ORANGE_MID, 0.95, True, "Latent predictor")
    text(slide, 909, 731, 125, 31, "Latent predictor", 11.5, "A6270F", True,
         name="Latent predictor heading")
    flame(slide, 1032, 726, "Predictor trainable flame", 0.9)
    network(slide, 912, 777, 112, 91, ORANGE, "Latent predictor network")

    oval(slide, 1307, 649, 92, 87, "FFF7F1", ORANGE, 0.9, "Latent loss")
    text(slide, 1335, 659, 38, 30, "ℒ", 15.5, INK, False, True, font=MATH, name="Loss symbol")
    text(slide, 1322, 690, 62, 26, "(latent loss)", 8.5, INK, name="Latent loss label")

    box(slide, 1212, 763, 303, 136, "FFF9F3", ORANGE_MID, 0.9, True, "Prediction panel")
    text(slide, 1283, 770, 177, 31, "Prediction Ŷₜ,ₕ", 13, ORANGE, False,
         name="Prediction heading")
    lock_icon(slide, 1476, 775, "B49570", "Prediction read-only lock", 0.78)
    text(slide, 1224, 809, 54, 34, "Global\nlatent", 8.3, INK, name="Prediction global label")
    cube(slide, 1228, 841, 35, ORANGE, "F1AF8E", "DA6D3D", "Prediction global cube", 3)
    text(slide, 1280, 849, 28, 29, "+", 18, "8F2A16", True, name="Prediction plus")
    text(slide, 1304, 808, 168, 25, "Regional latents (multiscale)", 8.1, INK,
         name="Prediction regional label")
    heat_grid(slide, 1310, 832, 53, 54, 5, 5, ["FFE0CE", "F4B18F", "E78153"], ORANGE,
              "Prediction regional coarse")
    heat_grid(slide, 1385, 837, 42, 49, 4, 5, ["FFE5D3", "F4B18F", "E77A48"], ORANGE,
              "Prediction regional medium")
    heat_grid(slide, 1456, 848, 31, 32, 5, 5, ["FFE8D8", "F1AD89", "DC6D3D"], ORANGE,
              "Prediction regional fine")
    text(slide, 1493, 844, 19, 29, "…", 12, INK, True, name="Prediction ellipsis")

    text(slide, 598, 897, 443, 24, "Future loss shapes Rₜ (representation at the present)",
         11.2, "C52912", False, name="Future loss shapes label")


def draw_legend(slide):
    box(slide, 142, 961, 1264, 49, "FFFEFC", "B8B8B3", 0.75, True, "Flow legend")
    line(slide, 166, 985, 205, 985, BLUE, 1.2, True, name="Legend online flow")
    text(slide, 216, 972, 157, 27, "Data flow (online)", 8.5, align=PP_ALIGN.LEFT,
         name="Legend online flow label")
    line(slide, 390, 985, 437, 985, BLUE, 1, True, True, "Legend use read")
    text(slide, 448, 972, 125, 27, "Use / read", 8.5, align=PP_ALIGN.LEFT,
         name="Legend use read label")
    line(slide, 585, 985, 645, 985, ORANGE, 1, True, True, "Legend learning signal")
    text(slide, 657, 972, 132, 27, "Learning signal", 8.5, align=PP_ALIGN.LEFT,
         name="Legend learning signal label")
    lock_icon(slide, 835, 972, GRAY, "Legend read-only lock", 0.72)
    text(slide, 867, 972, 102, 27, "Read-only", 8.5, align=PP_ALIGN.LEFT,
         name="Legend read-only label")
    flame(slide, 992, 971, "Legend trainable flame", 0.75)
    text(slide, 1023, 972, 194, 27, "Trainable (LoRA / predictor)", 8.5, align=PP_ALIGN.LEFT,
         name="Legend trainable label")
    lock_icon(slide, 1238, 972, GREEN, "Legend EMA lock", 0.72)
    text(slide, 1270, 972, 125, 27, "EMA stop-gradient", 8.5, align=PP_ALIGN.LEFT,
         name="Legend EMA label")


def build():
    with Image.open(SOURCE) as image:
        scans = [
            crop_blob(image, (32, 111, 107, 191)),
            crop_blob(image, (111, 111, 182, 191)),
            crop_blob(image, (187, 111, 260, 191)),
            crop_blob(image, (32, 261, 107, 354)),
            crop_blob(image, (111, 261, 186, 354)),
            crop_blob(image, (191, 261, 269, 354)),
        ]
        reuse_scans = [
            crop_blob(image, (1147, 197, 1238, 252)),
            crop_blob(image, (1252, 197, 1332, 252)),
            crop_blob(image, (1344, 194, 1438, 259)),
        ]
        future_scan = crop_blob(image, (1191, 387, 1237, 435))

    prs = Presentation()
    prs.slide_width = u(W)
    prs.slide_height = u(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    draw_routes(slide)
    draw_evidence(slide, scans)
    draw_encoder_and_resampler(slide)
    draw_dense_memory(slide)
    draw_vlm(slide)
    draw_queries(slide)
    draw_downstream(slide, reuse_scans)
    draw_future_target(slide, future_scan)
    draw_predictor_and_loss(slide)
    draw_legend(slide)

    prs.core_properties.title = "MedWorld-JEPA editable architecture figure"
    prs.core_properties.subject = f"Editable recreation of {SOURCE.name}"
    prs.core_properties.author = "Codex"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
