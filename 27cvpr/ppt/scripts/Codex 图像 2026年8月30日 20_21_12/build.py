from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


NAME = "Codex 图像 2026年8月30日 20_21_12"
ROOT = Path(__file__).resolve().parents[2]
ASSETS = ROOT / "source/mimic_appendix_cases"
OUTPUT = ROOT / "ppt" / f"{NAME}.pptx"

W, H, PPI = 1055, 1491, 100
FONT = "Roboto Condensed"
BG = "F9F8F4"
NAVY = "092D52"
BRIDGE = "526F8D"
RAIL = "7890A5"
BODY = "242321"
RULE = "D4D2CB"
DOT = "CBC8C1"
BORDER = "333331"


def u(px):
    return Inches(px / PPI)


def rgb(value):
    return RGBColor.from_string(value)


def clean(shape):
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:
        effect = style.find(qn("a:effectRef"))
        if effect is not None:
            effect.set("idx", "0")
    return shape


def text(slide, x, y, w, h, value, size=12, color=BODY, bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, spacing=1.0, name=None,
         font=FONT):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    shape.name = name or f"Text: {value[:24]}"
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = u(1)
    frame.margin_top = frame.margin_bottom = u(0.5)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = spacing
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return clean(shape)


def summary_text(slide, x, y, w, h, value, size, after, name):
    shape = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = u(1)
    frame.margin_top = frame.margin_bottom = u(0.5)
    frame.vertical_anchor = MSO_ANCHOR.TOP
    values = value.splitlines()
    for index, value_line in enumerate(values):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(after if index < len(values) - 1 else 0)
        run = paragraph.add_run()
        run.text = value_line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = rgb(BODY)
    return clean(shape)


def line(slide, x1, y1, x2, y2, color=NAVY, width=0.8,
         dashed=False, name="Line"):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2)
    )
    shape.name = name
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    xml_line = shape._element.spPr.get_or_add_ln()
    if dashed:
        dash = OxmlElement("a:prstDash")
        dash.set("val", "dash")
        xml_line.append(dash)
    return clean(shape)


def picture(slide, path, x, y, w, h, name):
    pic = slide.shapes.add_picture(str(path), u(x), u(y), u(w), u(h))
    pic.name = name
    clean(pic)
    border = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, u(x), u(y), u(w), u(h)
    )
    border.name = f"{name} border"
    border.fill.background()
    border.line.color.rgb = rgb(BORDER)
    border.line.width = Pt(1.0)
    border.adjustments[0] = 0.018
    clean(border)


def clock(slide, x, y, name):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(27), u(27))
    shape.name = f"{name} clock"
    shape.fill.background()
    shape.line.color.rgb = rgb(BRIDGE)
    shape.line.width = Pt(0.9)
    clean(shape)
    line(slide, x + 13.5, y + 5, x + 13.5, y + 14, BRIDGE, 0.9,
         name=f"{name} clock minute hand")
    line(slide, x + 13.5, y + 14, x + 19.5, y + 18, BRIDGE, 0.9,
         name=f"{name} clock hour hand")


def trajectory_arrow(slide, y, name):
    line(slide, 578, y, 630, y, BRIDGE, 0.8, name=f"{name} shaft")
    line(slide, 622, y - 8, 630, y, BRIDGE, 0.8, name=f"{name} upper head")
    line(slide, 630, y, 622, y + 8, BRIDGE, 0.8, name=f"{name} lower head")


CASES = [
    {
        "case": "A", "status": "Clear change",
        "meta": "20.1 h · 0–24 h · SICU",
        "summary": "No reported edema or\npleural effusion →\nmarked bilateral\nopacity and pulmonary\nedema; small right\npleural effusion.",
        "time": "20.1 h", "heading_y": 24, "image_y": 55, "image_h": 434,
        "current_x": 214, "current_w": 345, "follow_x": 650, "follow_w": 377,
        "case_y": 46, "rule_top": 66, "rule_bottom": 180,
        "status_y": 130, "meta_y": 155, "short_rule_y": 207,
        "summary_y": 234, "summary_h": 182, "arrow_y": 232,
        "clock_y": 258, "time_y": 290, "row_rule_y": 502,
        "current": "case_a_current_cxr.jpg",
        "follow": "case_a_observed_followup_cxr.jpg",
    },
    {
        "case": "B", "status": "Stable follow-up",
        "meta": "24.0 h · 0–24 h · SICU",
        "summary": "Persistent bibasal\nopacification, pleural\neffusions, and\ncompressive\natelectasis with little\ninterval change.",
        "time": "24.0 h", "heading_y": 518, "image_y": 549, "image_h": 424,
        "current_x": 211, "current_w": 348, "follow_x": 650, "follow_w": 378,
        "case_y": 537, "rule_top": 557, "rule_bottom": 673,
        "status_y": 621, "meta_y": 647, "short_rule_y": 698,
        "summary_y": 725, "summary_h": 182, "arrow_y": 724,
        "clock_y": 750, "time_y": 782, "row_rule_y": 985,
        "current": "case_b_current_cxr.jpg",
        "follow": "case_b_observed_followup_cxr.jpg",
    },
    {
        "case": "C", "status": "Clear change",
        "meta": "50.5 h · 24–72 h ·\nTrauma SICU",
        "summary": "No focal opacity or\npleural effusion →\nextensive bilateral\nparenchymal opacities,\ngreatest in the right\nupper lung; probable\nmultifocal pneumonia,\npossible effusion.",
        "time": "50.5 h", "heading_y": 997, "image_y": 1028, "image_h": 432,
        "current_x": 210, "current_w": 349, "follow_x": 650, "follow_w": 378,
        "case_y": 1017, "rule_top": 1037, "rule_bottom": 1174,
        "status_y": 1101, "meta_y": 1128, "short_rule_y": 1188,
        "summary_y": 1213, "summary_h": 205, "arrow_y": 1199,
        "clock_y": 1225, "time_y": 1257, "row_rule_y": None,
        "current": "case_c_current_cxr.jpg",
        "follow": "case_c_observed_followup_cxr.jpg",
    },
]


def add_case(slide, data):
    text(slide, data["current_x"] - 6, data["heading_y"] - 2,
         data["current_w"], 27, "Current", 13, NAVY, True,
         PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
         name=f"Case {data['case']} current heading", font="Arial")
    text(slide, data["follow_x"], data["heading_y"] - 2, data["follow_w"], 27,
         "Observed follow-up", 13, NAVY, True, PP_ALIGN.CENTER,
         MSO_ANCHOR.MIDDLE, name=f"Case {data['case']} follow-up heading",
         font="Arial")

    picture(slide, ASSETS / data["current"], data["current_x"], data["image_y"],
            data["current_w"], data["image_h"], f"Case {data['case']} current radiograph")
    picture(slide, ASSETS / data["follow"], data["follow_x"], data["image_y"],
            data["follow_w"], data["image_h"], f"Case {data['case']} observed follow-up")

    line(slide, 23, data["rule_top"], 23, data["rule_top"] + 36, NAVY, 1.0,
         name=f"Case {data['case']} upper accent")
    line(slide, 23, data["rule_top"] + 36, 23, data["rule_bottom"], RAIL, 1.0,
         name=f"Case {data['case']} lower accent")
    text(slide, 37, data["case_y"], 55, 66, data["case"], 48, NAVY, True,
         name=f"Case {data['case']} label")
    text(slide, 32, data["status_y"], 155, 25, data["status"], 13,
         name=f"Case {data['case']} status")
    text(slide, 32, data["meta_y"], 165, 46, data["meta"], 12.6,
         spacing=1.05, name=f"Case {data['case']} metadata")
    line(slide, 23, data["short_rule_y"], 184, data["short_rule_y"], DOT,
         0.7, True, name=f"Case {data['case']} summary divider")
    summary_size = 12.6 if data["case"] == "C" else 13.5
    paragraph_gap = {"A": 4.4, "B": 5.0, "C": 1.4}[data["case"]]
    summary_text(slide, 22, data["summary_y"], 186, data["summary_h"],
                 data["summary"], summary_size, paragraph_gap,
                 f"Case {data['case']} summary")

    trajectory_arrow(slide, data["arrow_y"], f"Case {data['case']} trajectory arrow")
    clock(slide, 591, data["clock_y"], f"Case {data['case']}")
    text(slide, 570, data["time_y"], 67, 27, data["time"], 13, BRIDGE,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE,
         name=f"Case {data['case']} elapsed time")

    if data["row_rule_y"]:
        line(slide, 23, data["row_rule_y"], 1031, data["row_rule_y"], RULE,
             0.7, name=f"Case {data['case']} row divider")


def build():
    needed = [ASSETS / data[key] for data in CASES for key in ("current", "follow")]
    missing = [str(path) for path in needed if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing radiograph assets:\n" + "\n".join(missing))

    presentation = Presentation()
    presentation.slide_width = u(W)
    presentation.slide_height = u(H)
    presentation.core_properties.title = "Representative longitudinal MIMIC trajectories"
    presentation.core_properties.subject = "Editable portrait appendix figure"
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)

    for data in CASES:
        add_case(slide, data)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
