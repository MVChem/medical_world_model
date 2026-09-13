"""Rebuild the reference as native PowerPoint objects and export paper vectors."""
from io import BytesIO
from pathlib import Path
import subprocess
import tempfile
import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
SOURCE, OUT = ROOT / 'source/fig1_v6.png', ROOT / 'ppt/fig1_v6.pptx'
W, H, PPI = 1860, 810, 120
INK, BLUE, PINK, CREAM = '171717', '9BCAE9', 'E89EBA', 'FEFAD1'
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W / PPI), Inches(H / PPI)
prs.core_properties.title = 'MedWorld-JEPA — fig1_v6'
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')


def u(v):
    return Inches(v / PPI)


def style(s, fill, stroke=INK, width=1):
    s.shadow.inherit = False
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = RGBColor.from_string(fill)
    else:
        s.fill.background()
    if stroke:
        s.line.color.rgb = RGBColor.from_string(stroke)
        s.line.width = Pt(width)
    else:
        s.line.fill.background()
    return s


def box(x, y, w, h, fill, stroke=INK, radius=.1, name='', width=1):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else
                               MSO_SHAPE.RECTANGLE, u(x), u(y), u(w), u(h))
    if radius:
        s.adjustments[0] = radius
    s.name = name or 'Module'
    return style(s, fill, stroke, width)


def text(x, y, w, h, value, size=20, bold=False, font='Comic Sans MS',
         color=INK, left=False, italic=False):
    s = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    s.name = 'Label: ' + value.replace('\n', ' / ')
    tf = s.text_frame
    tf.word_wrap, tf.auto_size = False, MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, label in enumerate(value.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT if left else PP_ALIGN.CENTER
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.0
        r = p.add_run()
        r.text = label
        r.font.name, r.font.size = font, Pt(size * 72 / PPI)
        r.font.bold, r.font.italic = bold, italic
        r.font.color.rgb = RGBColor.from_string(color)
    return s


def math(x, y, w, h, runs, size=27, color=INK):
    s = text(x, y, w, h, '', size, font='Times New Roman', left=True)
    s.name = 'Equation: ' + ''.join(v for v, _ in runs)
    for value, baseline in runs:
        r = s.text_frame.paragraphs[0].add_run()
        r.text = value
        r.font.name, r.font.italic = 'Times New Roman', True
        r.font.size = Pt(size * 72 / PPI)
        r.font.color.rgb = RGBColor.from_string(color)
        if baseline:
            r._r.get_or_add_rPr().set('baseline', str(baseline * 1000))
    return s


def line(points, arrow=False, color=INK, width=1.2, name='Connection', dashed=False):
    for i, ((x1, y1), (x2, y2)) in enumerate(zip(points, points[1:])):
        s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2))
        s.name = name
        s.line.color.rgb, s.line.width = RGBColor.from_string(color), Pt(width)
        if dashed:
            s.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow and i == len(points) - 2:
            end = OxmlElement('a:tailEnd')
            for k, v in {'type': 'triangle', 'w': 'sm', 'len': 'sm'}.items():
                end.set(k, v)
            s._element.spPr.get_or_add_ln().append(end)


def polygon(points, fill, name):
    b = slide.shapes.build_freeform(*points[0], scale=u(1))
    b.add_line_segments(points[1:], close=True)
    s = b.convert_to_shape()
    s.name = name
    return style(s, fill)


def gradient(s, top, bottom):
    s.fill.gradient()
    s.fill.gradient_angle = 90
    for stop, color in zip(s.fill.gradient_stops, [top, bottom]):
        stop.color.rgb = RGBColor.from_string(color)
    return s


def token(x, y, w, h, kind, name):
    fill, stroke, light = {'blue': (BLUE, '1686D1', 'BDDDF8'),
                            'pink': (PINK, 'CC398F', 'F2C5E4'),
                            'gray': ('BEBEBE', '888888', 'D0D0D0')}[kind]
    return gradient(box(x, y, w, h, fill, stroke, .08, name, .8), light, fill)


def slots(x, y, step, w, h, name):
    for i in range(8):
        token(x + i * step, y, w, h, 'pink', f'{name} {i + 1}/8')


def photo(bounds, x, y, w, h, name):
    data = BytesIO()
    with Image.open(SOURCE) as im:
        im.crop(bounds).save(data, format='PNG')
    data.seek(0)
    s = slide.shapes.add_picture(data, u(x), u(y), u(w), u(h))
    s.name = name


def report(y, future):
    box(21, y, 189, 98, 'FFF6E8', '84501F', .11, 'Follow-up report' if future else 'Report', .9)
    text(35, y + 4, 168, 28, 'Follow-up report' if future else 'Report', 20, True, left=True)
    polygon([(39, y+39), (62, y+39), (73, y+50), (73, y+87), (39, y+87)],
            'FFFBF0', 'Report document')
    line([(62, y+39), (62, y+50), (73, y+50)], width=.9)
    for yy in [57, 65, 74]:
        line([(46, y+yy), (66, y+yy)], width=.8)
    text(93, y+37, 111, 47, 'Findings: ...\nImpression: ...', 15.5, font='Arial', left=True)


def vlm(y, name):
    gradient(box(556, y, 438, 234, 'E7BCE7', '762381', .075, name, 1.3), 'EDD0EE', 'E5BEE9')
    text(714, y+6, 114, 40, 'VLM', 32, True)
    text(924, y+12, 59, 28, 'LoRA', 20)
    box(565, y+54, 419, 91, 'FFFCF3', 'A42A8C', .14, name+' input strip', .9)
    for x, w, label in [(578, 102, 'Visual tokens'), (702, 105, 'Text tokens'),
                        (835, 141, 'Learned slots U')]:
        text(x, y+61, w, 24, label, 16, True, font='Arial')
    for i in range(3):
        token(579+i*28, y+94, 23, 28, 'blue', name+' visual token')
        token(699+i*28, y+94, 23, 28, 'gray', name+' text token')
    text(668, y+89, 22, 30, '...', 20)
    text(786, y+89, 21, 30, '...', 20)
    slots(816, y+94, 20, 17, 28, name+' input slot')
    math(868, y+121, 90, 25, [('8 × d', 0)], 22)
    line([(775, y+145), (775, y+166)], True, width=1)
    for xx in [674, 710]:
        box(xx, y+161, 23, 61, CREAM, INK, .14, name+' transformer block', 1.1)
    text(757, y+177, 32, 32, '...', 24)
    text(805, y+181, 149, 29, 'Transformer', 21, True)


def state(x, y, future=False, predicted=False):
    w, h = (185, 132) if predicted else (180 if not future else 176, 120)
    name = 'Predicted future' if predicted else 'Future target' if future else 'Current state'
    gradient(box(x, y, w, h, 'FCE9F4', 'DD42A4', .1, name, .9), 'FFF4FA', 'FADDED')
    if predicted:
        text(x+8, y+5, w-16, 28, name, 19, True)
        math(x+78, y+29, 84, 37, [('Ŝ', 0), ('t,h', -25)], 29)
        ty = y+73
    else:
        text(x+8, y+12, 140, 30, name, 18, True, left=True)
        math(x+143 if not future else x+133, y+8, 52, 38,
             [('S', 0)] + ([('*', 40), ('t⁺', -25)] if future else [('t', -25)]), 26)
        ty = y+53 if not future else y+60
    slots(x+12, ty, 20, 18, 30 if not predicted else 26, name+' output token')
    math(x+66, ty+34 if not predicted else ty+27, 91, 25, [('8 × d', 0)], 22)


# Backgrounds and training-stage labels follow the v6 reference image.
box(2, 67, 1856, 354, 'F0F9EF', '398B31', .045, 'Current observations lane', .9)
box(2, 473, 1242, 328, 'EDF7FE', '65AEF6', .055, 'Observed future lane', .9)
box(1250, 477, 608, 321, 'FDEAE0', 'FFB2AA', .055, 'Future prediction stage', .9)
box(1305, 4, 551, 49, 'E1F3D9', 'D1EDC1', .25, 'Stage 1 heading', .8)
text(1317, 8, 530, 37, 'Stage 1: Task-supervised state pretraining', 25, True)
box(1341, 426, 517, 45, 'FEE6DD', 'FFB8B0', .23, 'Stage 2 heading', .8)
text(1350, 429, 500, 36, 'Stage 2: Future-state prediction fine-tuning', 22.5, True, color='6D200B')
text(727, 431, 374, 32, 'Same encoding pipeline at both times', 19, color='606773')

for future in [False, True]:
    dy = 390 if future else 0
    text(22, 76 if not future else 478, 322 if not future else 250, 35,
         'Observed future, t⁺' if future else 'Current observations, t', 26, True, left=True)
    if future:
        text(283, 480, 170, 28, '(Training target)', 19, left=True)
    text(24 if future else 54, 522 if future else 122, 171 if future else 96, 28,
         'Follow-up image' if future else 'Image', 20, True, left=True)
    photo((23,554,158,667) if future else (21,155,160,281),
          23 if future else 21, 554 if future else 155, 135 if future else 139,
          113 if future else 126, 'Follow-up chest radiograph' if future else 'Current chest radiograph')
    polygon([(186,141+dy), (301,179+dy), (301,250+dy), (186,285+dy)],
            'BFDFF9', 'Future V-JEPA encoder' if future else 'Current V-JEPA encoder')
    text(195, 188+dy, 100, 61, 'V-JEPA\nencoder', 24, True)
    for row in range(3):
        for col in range(3):
            token(330+col*20, 188+dy+row*20, 17, 17, 'blue', 'Future D feature' if future else 'Current D feature')
    math(339, 143+dy, 70, 40, [('D', 0)]+([('*', 40), ('t⁺', -25)] if future else [('t', -25)]), 29)
    box(416, 184+dy, 106, 73, CREAM, '665D29', .17, 'Visual adapter', .9)
    text(430, 194+dy, 79, 52, 'Visual\nadapter', 21, True)
    ry, ty, vy = (686,708,533) if future else (300,320,149)
    report(ry, future)
    box(247, ty, 136, 65, CREAM, '665D29', .17, 'Tokenizer and embedding', .9)
    text(256, ty+5, 117, 53, 'Tokenizer +\nembedding', 20, True)
    for i in range(3):
        token(417+i*28, ty+18, 23, 26, 'gray', 'Report text token')
    text(411, ty+47, 111, 28, 'Text tokens', 21, True, left=True)
    vlm(vy, 'Future VLM' if future else 'Current VLM')
    state(1064, 612 if future else 219, future)
    cy = 219+dy
    line([(158 if future else 160, cy-3), (186, cy-3)], True, name='Image to encoder')
    line([(301, cy), (328, cy)], True, name='Encoder to spatial features')
    line([(387, cy), (416, cy)], True, name='Spatial features to visual adapter')
    line([(522, cy), (556, cy)], True, name='Visual adapter to VLM')
    line([(210, ty+32), (247, ty+32)], True, name='Report to tokenizer')
    line([(383, ty+32), (415, ty+32)], True, name='Tokenizer to report tokens')
    line([(496, ty+32), (527, ty+32), (527, vy+168), (556, vy+168)], True, name='Report tokens to VLM')
    line([(994, 678 if future else 294), (1064, 678 if future else 294)], True, name='Slot hidden states to state readout')
    text(997, 246 if not future else 630, 67, 41, 'Slot hidden\nstates', 14, font='Arial')

# The spatial branch and query feed the task decoder separately from S_t.
line([(387, 201), (394, 201), (394, 112), (1277, 112), (1277, 188)], True,
     color='0086DC', width=1.4, name='Current spatial features to spatial task decoder')
math(617, 76, 44, 32, [('D', 0), ('t', -25)], 24, '0086DC')
text(646, 78, 435, 31, ': spatial tasks (for segmentation / super-resolution)',
     18, True, color='0086DC', left=True)
box(1268, 189, 215, 153, 'CFE7B5', '3E6527', .12, 'Task decoder', 1.1)
text(1286, 202, 180, 35, 'Task decoder', 28, True)
for xx in [1320,1406]:
    box(xx, 242, 23, 60, CREAM, INK, .14, 'Task transformer block', 1.1)
text(1358, 250, 31, 37, '...', 27)
text(1307, 306, 151, 31, 'Transformer', 22, True)
text(1322, 88, 110, 31, 'Task query', 22, True, left=True)
math(1437, 88, 30, 32, [('q',0)], 26)
token(1371, 121, 28, 29, 'blue', 'Task query token')
line([(1385,150), (1385,188)], True, name='Task query to task decoder')
line([(1244,279), (1268,279)], True, name='Current state to task decoder')
# Explicit junction: S_t also supplies the world model.
line([(1253,279), (1253,448), (1332,448), (1332,508), (1369,508)], True,
     name='Current state to latent world model')

# Four task cards; only the medical image areas are raster.
box(1517,64,341,294,'E8F5FF','59ACFF',.055,'Task loss panel',.9)
text(1559,70,257,37,'Task loss (Stage 1)',27,True)
for x,y,w,h,title in [(1526,111,162,119,'Classification'),(1694,111,153,119,'Segmentation'),
                       (1526,242,162,106,'Disease recognition'),(1694,242,153,106,'Super-resolution')]:
    box(x,y,w,h,'F3FAFE','67B1FC',.095,title+' card',.75)
    text(x+3,y+5,w-6,29,title,18 if len(title)<18 else 17,True)
line([(1483,252),(1497,252),(1497,172),(1517,172)],True,name='Task decoder to top task outputs')
line([(1483,274),(1517,274)],True,name='Task decoder to lower task outputs')
line([(1570,155),(1570,214),(1645,214)],color='278FDD',width=.9,name='Classification axes')
for x,y in [(1576,191),(1596,175),(1616,155)]:
    box(x,y,13,214-y,BLUE,'278FDD',0,'Classification bar',.6)
photo((1717,152,1821,223),1717,152,104,71,'Segmentation chest radiograph')
box(1540,284,134,43,'F8FCFF','456D94',.16,'Disease text box',.8)
text(1546,288,125,33,'Disease: ...',22,True)
photo((1708,283,1752,324),1708,283,44,41,'Low-resolution image detail')
photo((1786,280,1831,323),1786,280,45,43,'High-resolution image detail')
box(1708,282,44,43,None,'A5C568',0,'Low-resolution outline',.8)
box(1785,279,47,45,None,'587B94',0,'High-resolution outline',.8)
line([(1755,304),(1784,304)],True,width=.9,name='Super-resolution comparison')
line([(1757,296),(1765,287),(1782,279)],True,width=.8)
text(1744,324,51,23,'MSE',20,True,font='Arial',italic=True)

box(1369,487,242,146,'F3D8BF','79491F',.13,'Latent World Model',1.2)
text(1377,498,226,35,'Latent World Model',23,True)
for xx in [1425,1529]:
    box(xx,541,23,57,CREAM,INK,.14,'Predictor block',1.1)
text(1471,550,36,32,'...',25)
text(1408,600,173,28,'Predictor block',20,True)
text(1262,525,112,31,'Horizon h',21,True,left=True)
token(1290,558,31,30,'gray','Horizon token')
line([(1321,573),(1369,573)],True,name='Horizon to latent world model')
state(1658,500,predicted=True)
line([(1611,564),(1658,564)],True,name='Latent world model to predicted future state')
box(1593,671,253,97,'FAD9C8','874727',.13,'Latent prediction loss',1)
text(1602,681,235,33,'Latent prediction loss',22,True)
math(1632,718,213,41,[('(',0),('Ŝ',0),('t,h',-25),(' ↔ S',0),('*',40),('t⁺',-25),(')',0)],31)
line([(1746,632),(1746,671)],True,name='Predicted future state to latent loss')
line([(1240,715),(1593,715)],True,name='Observed future target to latent loss (stop-grad)',dashed=True,width=1)
text(1357,682,108,29,'stop-grad',20,True,italic=True)

# Disable theme shadows so every native shape stays flat.
for s in slide.shapes:
    for effect in s._element.findall('.//' + qn('a:effectRef')):
        effect.set('idx', '0')
    sp = s._element.find(qn('p:spPr'))
    if sp is not None and sp.find(qn('a:effectLst')) is None:
        sp.append(OxmlElement('a:effectLst'))

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
with tempfile.TemporaryDirectory(prefix='fig1_v6_') as work:
    subprocess.run(['libreoffice', f'-env:UserInstallation={Path(work).as_uri()}',
                    '--headless','--convert-to','pdf:impress_pdf_Export',
                    '--outdir',str(OUT.parent),str(OUT)], check=True, capture_output=True)
with fitz.open(OUT.with_suffix('.pdf')) as pdf:
    page = pdf[0]
    OUT.with_suffix('.svg').write_text(page.get_svg_image(text_as_path=True))
    page.get_pixmap(matrix=fitz.Matrix(4200/page.rect.width,4200/page.rect.width)).save(OUT.with_suffix('.png'))
print(OUT)
